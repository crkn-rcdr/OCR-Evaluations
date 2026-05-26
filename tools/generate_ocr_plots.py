#!/usr/bin/env python3
"""Generate Plotly dashboards for the OCR evaluation workbooks.

Default usage:

  python tools/generate_ocr_plots.py

This reads:

- test-results/paddle_vs_abbyy_errors_artifacts.xlsx
- test-results/first_article_ocr_comparison.xlsx

and writes an HTML dashboard to:

- test-results/plots/ocr_evaluation_dashboard.html

It also writes PNG files for README embedding under:

- test-results/plots/png/
"""

from __future__ import annotations

import argparse
import html
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from plotly import graph_objects as go
from plotly.io import to_html, write_image
from plotly.subplots import make_subplots


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PADDLE_VS_ABBYY = REPO_ROOT / "test-results" / "paddle_vs_abbyy_errors_artifacts.xlsx"
DEFAULT_FIRST_ARTICLE = REPO_ROOT / "test-results" / "first_article_ocr_comparison.xlsx"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "test-results" / "plots"
DEFAULT_OUTPUT_HTML = DEFAULT_OUTPUT_DIR / "ocr_evaluation_dashboard.html"
DEFAULT_OUTPUT_PNG_DIR = DEFAULT_OUTPUT_DIR / "png"


STRINGS = {
    "en": {
        "line_bar": "Paddle line count vs ABBYY",
        "baseline_100": "ABBYY baseline = 100%",
        "line_axis": "Percent of ABBYY line count",
        "word_bar": "Paddle word count vs ABBYY",
        "word_axis": "Percent of ABBYY word count",
        "abbyy_only": "ABBYY",
        "human_reference": "Human transcription",
        "paddle_group": "Paddle workflows",
        "word_count_axis": "Whole-page word count",
        "word_count_hover": "Word count",
        "article_word_count_axis": "First-article word count",
        "article_word_count_hover": "Word count",
        "artifact_bar": "Paddle workflow",
        "artifact_line": "ABBYY baseline",
        "artifact_axis": "Artifact-entry count",
        "artifact_percent_axis": "Percent of ABBYY",
        "artifact_hover_pct": "Percent of ABBYY",
        "artifact_hover_paddle": "Paddle",
        "artifact_hover_abbyy": "ABBYY",
        "artifact_percent_legend": "Paddle as % of ABBYY",
        "wer_axis": "WER (%)",
        "cer_axis": "CER (%)",
        "word_accuracy": "Word accuracy",
        "word_edits": "Word edits",
        "edit_sub": "Substitutions",
        "edit_ins": "Insertions",
        "edit_del": "Deletions",
        "edit_axis": "Word edits",
        "dashboard_title": "OCR Evaluation Dashboard",
        "dashboard_desc": "Generated from the workbook outputs used in this repo's testing section.",
        "generated": "Generated",
        "inputs": "Inputs",
        "english_section": "English Charts",
        "french_section": "Version française",
    },
    "fr": {
        "line_bar": "Nombre de lignes Paddle vs ABBYY",
        "baseline_100": "Référence ABBYY = 100 %",
        "line_axis": "Pourcentage du nombre de lignes d’ABBYY",
        "word_bar": "Nombre de mots Paddle vs ABBYY",
        "word_axis": "Pourcentage du nombre de mots d’ABBYY",
        "abbyy_only": "ABBYY",
        "human_reference": "Transcription humaine",
        "paddle_group": "Flux Paddle",
        "word_count_axis": "Nombre de mots sur la page entière",
        "word_count_hover": "Nombre de mots",
        "article_word_count_axis": "Nombre de mots du premier article",
        "article_word_count_hover": "Nombre de mots",
        "artifact_bar": "Flux Paddle",
        "artifact_line": "Référence ABBYY",
        "artifact_axis": "Nombre d’artéfacts",
        "artifact_percent_axis": "Pourcentage d’ABBYY",
        "artifact_hover_pct": "Pourcentage d’ABBYY",
        "artifact_hover_paddle": "Paddle",
        "artifact_hover_abbyy": "ABBYY",
        "artifact_percent_legend": "Paddle en % d’ABBYY",
        "wer_axis": "Taux d’erreur de mots (%)",
        "cer_axis": "Taux d’erreur de caractères (%)",
        "word_accuracy": "Exactitude des mots",
        "word_edits": "Modifications de mots",
        "edit_sub": "Substitutions",
        "edit_ins": "Insertions",
        "edit_del": "Suppressions",
        "edit_axis": "Modifications de mots",
        "dashboard_title": "Tableau de bord d’évaluation OCR",
        "dashboard_desc": "Généré à partir des classeurs utilisés dans la section de test de ce dépôt.",
        "generated": "Généré",
        "inputs": "Entrées",
        "english_section": "Charts in English",
        "french_section": "Graphiques en français",
    },
}


@dataclass(frozen=True)
class SheetSummary:
    label: str
    paddle_source: str
    abbyy_source: str
    paddle_line_count: int
    abbyy_line_count: int
    paddle_word_count: int
    abbyy_word_count: int
    paddle_artifact_entries: int
    abbyy_artifact_entries: int
    paddle_error_counts: dict[str, int]
    abbyy_error_counts: dict[str, int]


@dataclass(frozen=True)
class HumanSummary:
    rank: int
    label: str
    source_path: str
    word_accuracy: float
    wer: float
    cer: float
    word_edits: int
    substitutions: int
    insertions: int
    deletions: int
    character_edits: int
    human_words: int
    ocr_words: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate Plotly visualizations for OCR evaluation workbooks."
    )
    parser.add_argument(
        "--paddle-vs-abbyy",
        type=Path,
        default=DEFAULT_PADDLE_VS_ABBYY,
        help="Path to paddle_vs_abbyy_errors_artifacts.xlsx",
    )
    parser.add_argument(
        "--first-article",
        type=Path,
        default=DEFAULT_FIRST_ARTICLE,
        help="Path to first_article_ocr_comparison.xlsx",
    )
    parser.add_argument(
        "--output-html",
        type=Path,
        default=DEFAULT_OUTPUT_HTML,
        help="Output HTML dashboard path.",
    )
    parser.add_argument(
        "--output-png-dir",
        type=Path,
        default=DEFAULT_OUTPUT_PNG_DIR,
        help="Output directory for PNG chart exports.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    paddle_vs_abbyy = args.paddle_vs_abbyy.resolve()
    first_article = args.first_article.resolve()
    output_html = args.output_html.resolve()
    output_png_dir = args.output_png_dir.resolve()
    output_html.parent.mkdir(parents=True, exist_ok=True)
    output_png_dir.mkdir(parents=True, exist_ok=True)

    whole_page = load_whole_page_summaries(paddle_vs_abbyy)
    human_rows = load_human_summaries(first_article)

    figures_en = [
        (
            "line_count_recovery",
            "Whole-Page Line Recovery",
            (
                "How many lines each workflow recovered relative to ABBYY. "
                "Values above 100% mean Paddle produced more line segments than "
                "ABBYY, which can include over-segmentation or duplicates."
            ),
            build_line_recovery_figure(whole_page, "en"),
        ),
        (
            "word_count_recovery",
            "Whole-Page Word Recovery",
            (
                "How many words each workflow recovered relative to ABBYY. "
                "This makes the preprocess-only under-extraction problem obvious "
                "and shows why `PaddleOCR (No VL)` became the base pipeline."
            ),
            build_word_recovery_figure(whole_page, "en"),
        ),
        (
            "whole_page_word_counts",
            "Whole-Page Word Counts",
            (
                "Absolute whole-page word counts for ABBYY and each Paddle-based "
                "workflow."
            ),
            build_absolute_word_count_figure(whole_page, "en"),
        ),
        (
            "human_word_counts",
            "First-Article Word Counts",
            (
                "Absolute word counts for the human first-article transcription, "
                "ABBYY, and each Paddle-based workflow. Bars are sorted from low "
                "to high."
            ),
            build_human_word_count_figure(human_rows, "en"),
        ),
        (
            "artifact_totals",
            "Whole-Page Artifact Totals",
            (
                "Overall artifact-entry counts from the side-by-side workbook. "
                "Lower is cleaner, but these totals should be read together with "
                "the coverage chart above."
            ),
            build_artifact_totals_figure(whole_page, "en"),
        ),
        (
            "artifact_categories",
            "Artifact Categories vs ABBYY",
            (
                "Each subplot shows one Paddle workflow. Category values are "
                "expressed as a percent of ABBYY's count for the same category, "
                "so 100% means parity with ABBYY, above 100% means more tagged "
                "artifacts than ABBYY, and below 100% means fewer."
            ),
            build_artifact_category_vs_abbyy_figure(whole_page, "en"),
        ),
        (
            "human_tradeoffs",
            "Human-Transcription Tradeoffs",
            (
                "WER and CER plotted together for the comparison of human "
                "transcription of the first article on the page to the AI "
                "transcription. Points closer to the lower-left are better; "
                "marker size tracks word accuracy."
            ),
            build_human_scatter_figure(human_rows, "en"),
        ),
        (
            "human_edit_breakdown",
            "Human-Transcription Edit Breakdown",
            (
                "Word-level substitutions, insertions, and deletions for the "
                "comparison of human transcription of the first article on the "
                "page to the AI transcription."
            ),
            build_edit_breakdown_figure(human_rows, "en"),
        ),
    ]
    figures_fr = [
        (
            "line_count_recovery",
            "Récupération des lignes sur la page entière",
            (
                "Montre combien de lignes chaque flux de travail a récupérées "
                "par rapport à ABBYY. Au-dessus de 100 %, Paddle a produit plus "
                "de segments de lignes qu’ABBYY, ce qui peut inclure de la "
                "sur-segmentation ou des doublons."
            ),
            build_line_recovery_figure(whole_page, "fr"),
        ),
        (
            "word_count_recovery",
            "Récupération des mots sur la page entière",
            (
                "Montre combien de mots chaque flux de travail a récupérés par "
                "rapport à ABBYY. Cela met en évidence la sous-extraction du "
                "prétraitement seul et explique pourquoi `PaddleOCR (No VL)` est "
                "devenu le pipeline de base."
            ),
            build_word_recovery_figure(whole_page, "fr"),
        ),
        (
            "whole_page_word_counts",
            "Nombre total de mots sur la page entière",
            (
                "Montre le nombre absolu de mots pour ABBYY et pour chaque flux "
                "de travail basé sur Paddle."
            ),
            build_absolute_word_count_figure(whole_page, "fr"),
        ),
        (
            "human_word_counts",
            "Nombre de mots du premier article",
            (
                "Montre le nombre absolu de mots pour la transcription humaine "
                "du premier article, ABBYY et chaque flux de travail basé sur "
                "Paddle. Les barres sont triées du plus faible au plus élevé."
            ),
            build_human_word_count_figure(human_rows, "fr"),
        ),
        (
            "artifact_totals",
            "Total des artéfacts sur la page entière",
            (
                "Montre le nombre total d’entrées d’artéfacts dans le classeur "
                "comparatif. Une valeur plus faible est préférable, mais ce "
                "résultat doit être lu avec les graphiques de couverture."
            ),
            build_artifact_totals_figure(whole_page, "fr"),
        ),
        (
            "artifact_categories",
            "Catégories d’artéfacts par rapport à ABBYY",
            (
                "Chaque sous-graphe montre un flux Paddle. Les valeurs sont "
                "exprimées en pourcentage du nombre d’artéfacts d’ABBYY pour la "
                "même catégorie : 100 % signifie une parité avec ABBYY, au-dessus "
                "de 100 % signifie plus d’artéfacts qu’ABBYY, et en dessous de "
                "100 % signifie moins d’artéfacts."
            ),
            build_artifact_category_vs_abbyy_figure(whole_page, "fr"),
        ),
        (
            "human_tradeoffs",
            "Compromis par rapport à la transcription humaine",
            (
                "Le WER et le CER sont tracés ensemble pour comparer la "
                "transcription humaine du premier article de la page à la "
                "transcription de l’IA. Les points plus proches du coin inférieur "
                "gauche sont meilleurs; la taille du marqueur suit l’exactitude "
                "des mots."
            ),
            build_human_scatter_figure(human_rows, "fr"),
        ),
        (
            "human_edit_breakdown",
            "Répartition des modifications par rapport à la transcription humaine",
            (
                "Montre les substitutions, insertions et suppressions au niveau "
                "des mots pour la comparaison entre la transcription humaine du "
                "premier article et la transcription de l’IA."
            ),
            build_edit_breakdown_figure(human_rows, "fr"),
        ),
    ]

    html_text = build_dashboard_html(figures_en, figures_fr, paddle_vs_abbyy, first_article)
    output_html.write_text(html_text, encoding="utf-8")
    write_png_exports(figures_en, output_png_dir)
    print(f"Wrote {output_html}")
    print(f"Wrote PNG charts to {output_png_dir}")


def load_whole_page_summaries(path: Path) -> list[SheetSummary]:
    workbook = load_workbook(path, data_only=True)
    summaries: list[SheetSummary] = []
    for worksheet in workbook.worksheets:
        header = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
        error_column_map = parse_error_columns(header)
        summary_values = {
            str(worksheet.cell(row, 1).value or "").strip().lower(): worksheet.cell(row, 2).value
            for row in range(1, min(8, worksheet.max_row) + 1)
        }
        abbyy_summary_values = {
            str(worksheet.cell(row, 1).value or "").strip().lower(): worksheet.cell(row, 3).value
            for row in range(1, min(8, worksheet.max_row) + 1)
        }

        paddle_error_counts = count_nonempty_error_cells(worksheet, error_column_map, "paddle")
        abbyy_error_counts = count_nonempty_error_cells(worksheet, error_column_map, "abbyy")

        summaries.append(
            SheetSummary(
                label=friendly_whole_page_label(worksheet.title),
                paddle_source=str(summary_values.get("source file") or ""),
                abbyy_source=str(abbyy_summary_values.get("source file") or ""),
                paddle_line_count=parse_int(summary_values.get("line count")),
                abbyy_line_count=parse_int(abbyy_summary_values.get("line count")),
                paddle_word_count=parse_int(summary_values.get("word count")),
                abbyy_word_count=parse_int(abbyy_summary_values.get("word count")),
                paddle_artifact_entries=parse_int(summary_values.get("error/artifact entries")),
                abbyy_artifact_entries=parse_int(abbyy_summary_values.get("error/artifact entries")),
                paddle_error_counts=paddle_error_counts,
                abbyy_error_counts=abbyy_error_counts,
            )
        )
    return summaries


def parse_error_columns(header: list[object]) -> dict[str, dict[str, int]]:
    columns: dict[str, dict[str, int]] = {}
    for index, value in enumerate(header, start=1):
        text = str(value or "").strip()
        if " - paddle" in text:
            error_type = text.replace(" - paddle", "")
            columns.setdefault(error_type, {})["paddle"] = index
        elif " - abbyy" in text:
            error_type = text.replace(" - abbyy", "")
            columns.setdefault(error_type, {})["abbyy"] = index
    return columns


def count_nonempty_error_cells(
    worksheet,
    error_column_map: dict[str, dict[str, int]],
    side: str,
) -> dict[str, int]:
    counts: dict[str, int] = {}
    for error_type, columns in error_column_map.items():
        column = columns.get(side)
        if column is None:
            counts[error_type] = 0
            continue
        count = 0
        for row in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row, column).value
            if value is None:
                continue
            if str(value).strip():
                count += 1
        counts[error_type] = count
    return counts


def load_human_summaries(path: Path) -> list[HumanSummary]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["Summary"]
    rows: list[HumanSummary] = []
    for row in range(2, worksheet.max_row + 1):
        label = friendly_human_label(str(worksheet.cell(row, 2).value or ""))
        rows.append(
            HumanSummary(
                rank=parse_int(worksheet.cell(row, 1).value),
                label=label,
                source_path=str(worksheet.cell(row, 3).value or ""),
                word_accuracy=parse_float(worksheet.cell(row, 4).value),
                wer=parse_float(worksheet.cell(row, 5).value),
                cer=parse_float(worksheet.cell(row, 6).value),
                word_edits=parse_int(worksheet.cell(row, 7).value),
                substitutions=parse_int(worksheet.cell(row, 8).value),
                insertions=parse_int(worksheet.cell(row, 9).value),
                deletions=parse_int(worksheet.cell(row, 10).value),
                character_edits=parse_int(worksheet.cell(row, 11).value),
                human_words=parse_int(worksheet.cell(row, 12).value),
                ocr_words=parse_int(worksheet.cell(row, 13).value),
            )
        )
    return rows


def t(lang: str, key: str) -> str:
    return STRINGS[lang][key]


def build_line_recovery_figure(rows: list[SheetSummary], lang: str) -> go.Figure:
    labels = [row.label for row in rows]
    line_pct = [percentage(row.paddle_line_count, row.abbyy_line_count) for row in rows]
    hover_line = [
        f"Paddle {row.paddle_line_count:,}<br>ABBYY {row.abbyy_line_count:,}"
        for row in rows
    ]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name=t(lang, "line_bar"),
            x=labels,
            y=line_pct,
            text=[f"{value:.1f}%" for value in line_pct],
            textposition="outside",
            hovertemplate="%{x}<br>%{customdata}<extra></extra>",
            customdata=hover_line,
            marker_color="#33658A",
        )
    )
    figure.add_trace(
        go.Scatter(
            name=t(lang, "baseline_100"),
            x=labels,
            y=[100.0] * len(labels),
            mode="lines",
            line=dict(color="#666666", width=2, dash="dash"),
            hovertemplate=f"{t(lang, 'baseline_100')}<extra></extra>",
        )
    )
    figure.update_layout(
        yaxis_title=t(lang, "line_axis"),
        template="plotly_white",
        margin=dict(l=50, r=30, t=50, b=100),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    figure.update_yaxes(range=[0, max(line_pct) * 1.18])
    return figure


def build_word_recovery_figure(rows: list[SheetSummary], lang: str) -> go.Figure:
    labels = [row.label for row in rows]
    word_pct = [percentage(row.paddle_word_count, row.abbyy_word_count) for row in rows]
    hover_word = [
        f"Paddle {row.paddle_word_count:,}<br>ABBYY {row.abbyy_word_count:,}"
        for row in rows
    ]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name=t(lang, "word_bar"),
            x=labels,
            y=word_pct,
            text=[f"{value:.1f}%" for value in word_pct],
            textposition="outside",
            hovertemplate="%{x}<br>%{customdata}<extra></extra>",
            customdata=hover_word,
            marker_color="#55A630",
        )
    )
    figure.add_trace(
        go.Scatter(
            name=t(lang, "baseline_100"),
            x=labels,
            y=[100.0] * len(labels),
            mode="lines",
            line=dict(color="#666666", width=2, dash="dash"),
            hovertemplate=f"{t(lang, 'baseline_100')}<extra></extra>",
        )
    )
    figure.update_layout(
        yaxis_title=t(lang, "word_axis"),
        template="plotly_white",
        margin=dict(l=50, r=30, t=50, b=100),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    figure.update_yaxes(range=[0, max(word_pct) * 1.18])
    return figure


def build_absolute_word_count_figure(rows: list[SheetSummary], lang: str) -> go.Figure:
    entries = [("ABBYY", rows[0].abbyy_word_count, "abbyy")]
    entries.extend((row.label, row.paddle_word_count, "paddle") for row in rows)
    ordered = sorted(entries, key=lambda item: (item[1], item[0]))
    ordered_labels = [label for label, _value, _kind in ordered]
    values = [value for _label, value, _kind in ordered]
    abbyy_labels = [label for label, _value, kind in ordered if kind == "abbyy"]
    abbyy_values = [value for _label, value, kind in ordered if kind == "abbyy"]
    paddle_labels = [label for label, _value, kind in ordered if kind == "paddle"]
    paddle_values = [value for _label, value, kind in ordered if kind == "paddle"]
    colors = [
        "#6C757D",
        "#33658A",
        "#55A630",
        "#BC4749",
        "#FF9F1C",
        "#6A4C93",
        "#1982C4",
    ][: len(ordered_labels)]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name=t(lang, "abbyy_only"),
            x=abbyy_labels,
            y=abbyy_values,
            text=[f"{value:,}" for value in abbyy_values],
            textposition="outside",
            marker_color=colors[:1],
            hovertemplate=f"%{{x}}<br>{t(lang, 'word_count_hover')}: %{{y:,}}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Bar(
            name=t(lang, "paddle_group"),
            x=paddle_labels,
            y=paddle_values,
            text=[f"{value:,}" for value in paddle_values],
            textposition="outside",
            marker_color=colors[1:],
            hovertemplate=f"%{{x}}<br>{t(lang, 'word_count_hover')}: %{{y:,}}<extra></extra>",
        )
    )
    figure.update_layout(
        yaxis_title=t(lang, "word_count_axis"),
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=100),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    figure.update_xaxes(categoryorder="array", categoryarray=ordered_labels)
    figure.update_yaxes(range=[0, max(values) * 1.15])
    return figure


def build_human_word_count_figure(rows: list[HumanSummary], lang: str) -> go.Figure:
    reference_words = rows[0].human_words if rows else 0
    entries = [(t(lang, "human_reference"), reference_words, "human")]
    for row in rows:
        kind = "abbyy" if row.label == "ABBYY" else "paddle"
        entries.append((row.label, row.ocr_words, kind))

    ordered = sorted(entries, key=lambda item: (item[1], item[0]))
    ordered_labels = [label for label, _value, _kind in ordered]
    values = [value for _label, value, _kind in ordered]
    human_labels = [label for label, _value, kind in ordered if kind == "human"]
    human_values = [value for _label, value, kind in ordered if kind == "human"]
    abbyy_labels = [label for label, _value, kind in ordered if kind == "abbyy"]
    abbyy_values = [value for _label, value, kind in ordered if kind == "abbyy"]
    paddle_labels = [label for label, _value, kind in ordered if kind == "paddle"]
    paddle_values = [value for _label, value, kind in ordered if kind == "paddle"]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name=t(lang, "human_reference"),
            x=human_labels,
            y=human_values,
            text=[f"{value:,}" for value in human_values],
            textposition="outside",
            marker_color="#2F4858",
            hovertemplate=f"%{{x}}<br>{t(lang, 'article_word_count_hover')}: %{{y:,}}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Bar(
            name=t(lang, "abbyy_only"),
            x=abbyy_labels,
            y=abbyy_values,
            text=[f"{value:,}" for value in abbyy_values],
            textposition="outside",
            marker_color="#6C757D",
            hovertemplate=f"%{{x}}<br>{t(lang, 'article_word_count_hover')}: %{{y:,}}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Bar(
            name=t(lang, "paddle_group"),
            x=paddle_labels,
            y=paddle_values,
            text=[f"{value:,}" for value in paddle_values],
            textposition="outside",
            marker_color=["#33658A", "#55A630", "#BC4749", "#FF9F1C", "#6A4C93", "#1982C4"][: len(paddle_values)],
            hovertemplate=f"%{{x}}<br>{t(lang, 'article_word_count_hover')}: %{{y:,}}<extra></extra>",
        )
    )
    figure.update_layout(
        yaxis_title=t(lang, "article_word_count_axis"),
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=100),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    figure.update_xaxes(categoryorder="array", categoryarray=ordered_labels)
    figure.update_yaxes(range=[0, max(values) * 1.15])
    return figure


def build_artifact_totals_figure(rows: list[SheetSummary], lang: str) -> go.Figure:
    labels = [row.label for row in rows]
    paddle_values = [row.paddle_artifact_entries for row in rows]
    abbyy_values = [row.abbyy_artifact_entries for row in rows]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name=t(lang, "artifact_bar"),
            x=labels,
            y=paddle_values,
            text=paddle_values,
            textposition="outside",
            marker_color="#BC4749",
        )
    )
    figure.add_trace(
        go.Scatter(
            name=t(lang, "artifact_line"),
            x=labels,
            y=abbyy_values,
            mode="lines+markers+text",
            text=[str(value) for value in abbyy_values],
            textposition="top center",
            line=dict(color="#6C757D", width=3, dash="dash"),
            marker=dict(color="#6C757D", size=9),
        )
    )
    figure.update_layout(
        yaxis_title=t(lang, "artifact_axis"),
        template="plotly_white",
        margin=dict(l=50, r=30, t=50, b=80),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_artifact_category_vs_abbyy_figure(rows: list[SheetSummary], lang: str) -> go.Figure:
    error_types = list(rows[0].paddle_error_counts.keys())
    subplot_titles = [row.label for row in rows]
    figure = make_subplots(
        rows=len(rows),
        cols=1,
        subplot_titles=subplot_titles,
        vertical_spacing=0.045,
    )

    colors = ["#33658A", "#55A630", "#BC4749", "#FF9F1C", "#6A4C93", "#1982C4"]
    max_ratio = 0.0
    for index, row in enumerate(rows):
        ratios = [
            percentage(
                row.paddle_error_counts[error_type],
                max(1, row.abbyy_error_counts[error_type]),
            )
            if row.abbyy_error_counts[error_type] > 0
            else (100.0 if row.paddle_error_counts[error_type] == 0 else row.paddle_error_counts[error_type] * 100.0)
            for error_type in error_types
        ]
        max_ratio = max(max_ratio, max(ratios))
        customdata = [
            [
                row.paddle_error_counts[error_type],
                row.abbyy_error_counts[error_type],
            ]
            for error_type in error_types
        ]
        subplot_row = index + 1
        figure.add_trace(
            go.Bar(
                x=ratios,
                y=error_types,
                orientation="h",
                marker_color=colors[index % len(colors)],
                customdata=customdata,
                hovertemplate=(
                    "%{y}<br>"
                    f"{t(lang, 'artifact_hover_paddle')}: %{{customdata[0]}}<br>"
                    f"{t(lang, 'artifact_hover_abbyy')}: %{{customdata[1]}}<br>"
                    f"{t(lang, 'artifact_hover_pct')}: %{{x:.1f}}%"
                    "<extra></extra>"
                ),
                showlegend=index == 0,
                name=t(lang, "artifact_percent_legend"),
            ),
            row=subplot_row,
            col=1,
        )
    figure.add_trace(
        go.Scatter(
            x=[None],
            y=[None],
            mode="lines",
            line=dict(color="#666666", width=2, dash="dash"),
            name=t(lang, "baseline_100"),
            showlegend=True,
            hoverinfo="skip",
        ),
        row=1,
        col=1,
    )
    for subplot_row in range(1, len(rows) + 1):
        figure.add_vline(
            x=100,
            line_dash="dash",
            line_color="#666666",
            row=subplot_row,
            col=1,
        )

    max_range = max(140.0, min(max_ratio * 1.12, 400.0))
    figure.update_layout(
        template="plotly_white",
        margin=dict(l=120, r=24, t=70, b=80),
        height=max(1500, 235 * len(rows)),
        legend=dict(orientation="h", yanchor="bottom", y=1.01, xanchor="left", x=0),
    )
    for row_index in range(1, len(rows) + 1):
        figure.update_xaxes(
            title_text=t(lang, "artifact_percent_axis"),
            range=[0, max_range],
            row=row_index,
            col=1,
        )
        figure.update_yaxes(autorange="reversed", row=row_index, col=1)
    return figure


def build_human_scatter_figure(rows: list[HumanSummary], lang: str) -> go.Figure:
    figure = go.Figure()
    point_colors = {
        "ABBYY": "#6366F1",
        "PaddleOCR + Chandra": "#EF553B",
        "PaddleOCR + PaddleVL": "#4CC38A",
        "PaddleOCR + DeepSeek": "#9B5DE5",
        "PaddleOCR (No VL)": "#FFA24C",
        "PaddleOCR + olmOCR": "#4DBBD5",
    }
    label_offsets = {
        "ABBYY": dict(ax=0, ay=-38, xanchor="center"),
        "PaddleOCR + Chandra": dict(ax=0, ay=34, xanchor="center", yanchor="top"),
        "PaddleOCR + PaddleVL": dict(ax=0, ay=-42, xanchor="center"),
        "PaddleOCR + DeepSeek": dict(ax=0, ay=-34, xanchor="center"),
        "PaddleOCR (No VL)": dict(ax=18, ay=-34, xanchor="left"),
        "PaddleOCR + olmOCR": dict(ax=0, ay=-38, xanchor="center"),
    }
    for row in rows:
        x_value = row.wer * 100
        y_value = row.cer * 100
        color = point_colors.get(row.label, "#33658A")
        figure.add_trace(
            go.Scatter(
                x=[x_value],
                y=[y_value],
                mode="markers",
                marker=dict(
                    size=22 + row.word_accuracy * 28,
                    color=color,
                    line=dict(width=1, color="#333333"),
                ),
                customdata=[[row.word_accuracy * 100, row.word_edits, row.substitutions, row.insertions, row.deletions]],
                hovertemplate=(
                    f"{row.label}<br>"
                    f"{t(lang, 'word_accuracy')}: %{{customdata[0]:.2f}}%<br>"
                    "WER: %{x:.2f}%<br>"
                    "CER: %{y:.2f}%<br>"
                    f"{t(lang, 'word_edits')}: %{{customdata[1]}}<br>"
                    "S/I/D: %{customdata[2]}/%{customdata[3]}/%{customdata[4]}"
                    "<extra></extra>"
                ),
                name=row.label,
            )
        )
        offset = label_offsets.get(row.label, dict(ax=0, ay=-36, xanchor="center"))
        figure.add_annotation(
            x=x_value,
            y=y_value,
            text=row.label,
            showarrow=False,
            xanchor=offset["xanchor"],
            yanchor=offset.get("yanchor", "bottom"),
            xshift=offset["ax"],
            yshift=-offset["ay"],
            bgcolor=hex_to_rgba(color, 0.16),
            bordercolor=color,
            borderwidth=2,
            borderpad=3,
            font=dict(size=11, color="#16324F"),
        )
    x_values = [row.wer * 100 for row in rows]
    y_values = [row.cer * 100 for row in rows]
    x_min = max(0, min(x_values) - 0.55)
    x_max = max(x_values) + 0.55
    y_min = max(0, min(y_values) - 0.35)
    y_max = max(y_values) + 0.55
    figure.update_layout(
        xaxis_title=t(lang, "wer_axis"),
        yaxis_title=t(lang, "cer_axis"),
        template="plotly_white",
        showlegend=True,
        margin=dict(l=70, r=70, t=60, b=70),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    figure.update_xaxes(range=[x_min, x_max], automargin=True)
    figure.update_yaxes(range=[y_min, y_max], automargin=True)
    return figure


def build_edit_breakdown_figure(rows: list[HumanSummary], lang: str) -> go.Figure:
    ordered = sorted(rows, key=lambda row: row.rank)
    labels = [row.label for row in ordered]
    substitutions = [row.substitutions for row in ordered]
    insertions = [row.insertions for row in ordered]
    deletions = [row.deletions for row in ordered]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name=t(lang, "edit_sub"),
            x=labels,
            y=substitutions,
            marker_color="#3A86FF",
        )
    )
    figure.add_trace(
        go.Bar(
            name=t(lang, "edit_ins"),
            x=labels,
            y=insertions,
            marker_color="#FFBE0B",
        )
    )
    figure.add_trace(
        go.Bar(
            name=t(lang, "edit_del"),
            x=labels,
            y=deletions,
            marker_color="#FB5607",
        )
    )
    figure.update_layout(
        barmode="stack",
        yaxis_title=t(lang, "edit_axis"),
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=90),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_dashboard_html(
    figures_en: list[tuple[str, str, str, go.Figure]],
    figures_fr: list[tuple[str, str, str, go.Figure]],
    paddle_vs_abbyy_path: Path,
    first_article_path: Path,
) -> str:
    generated = datetime.now(timezone.utc).isoformat()
    sections_en, include_plotlyjs = build_dashboard_sections(figures_en, include_plotlyjs=True)
    sections_fr, _ = build_dashboard_sections(figures_fr, include_plotlyjs=include_plotlyjs)

    return "\n".join(
        [
            "<!DOCTYPE html>",
            "<html lang='en'>",
            "<head>",
            "<meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width, initial-scale=1'>",
            f"<title>{html.escape(t('en', 'dashboard_title'))}</title>",
            "<style>",
            "body { font-family: Segoe UI, Arial, sans-serif; margin: 0; background: #f7f7f7; color: #111; }",
            "main { max-width: 1400px; margin: 0 auto; padding: 24px; }",
            "header { margin-bottom: 24px; }",
            "h1 { margin: 0 0 8px; }",
            "h2 { margin-top: 0; }",
            "p { line-height: 1.5; }",
            ".chart-block { background: #fff; border: 1px solid #ddd; padding: 20px; margin: 0 auto 20px; border-radius: 10px; max-width: 1260px; }",
            ".language-section { margin-top: 28px; }",
            ".language-section > h2 { max-width: 1260px; margin-left: auto; margin-right: auto; }",
            ".figure-wrap { width: min(100%, 1120px); margin: 0 auto; }",
            ".figure-wrap.narrow { width: min(100%, 860px); margin: 0 auto; }",
            "code { background: #f0f0f0; padding: 2px 5px; border-radius: 4px; }",
            "</style>",
            "</head>",
            "<body>",
            "<main>",
            "<header>",
            f"<h1>{html.escape(t('en', 'dashboard_title'))}</h1>",
            f"<p>{html.escape(t('en', 'dashboard_desc'))}</p>",
            f"<p><strong>{html.escape(t('en', 'generated'))}:</strong> {html.escape(generated)}</p>",
            f"<p><strong>{html.escape(t('en', 'inputs'))}:</strong> <code>{html.escape(str(paddle_vs_abbyy_path))}</code><br>"
            f"<code>{html.escape(str(first_article_path))}</code></p>",
            "</header>",
            f"<section class='language-section'><h2>{html.escape(t('en', 'english_section'))}</h2>",
            *sections_en,
            "</section>",
            f"<section class='language-section'><h2>{html.escape(t('fr', 'french_section'))}</h2>",
            *sections_fr,
            "</section>",
            "</main>",
            "</body>",
            "</html>",
        ]
    )


def build_dashboard_sections(
    figures: list[tuple[str, str, str, go.Figure]],
    include_plotlyjs: bool,
) -> tuple[list[str], bool]:
    sections: list[str] = []
    for slug, title, description, figure in figures:
        section_html = to_html(
            figure,
            include_plotlyjs=include_plotlyjs,
            full_html=False,
            config={"responsive": True, "displaylogo": False},
        )
        include_plotlyjs = False
        wrap_class = "figure-wrap narrow" if slug == "artifact_categories" else "figure-wrap"
        sections.append(
            "\n".join(
                [
                    "<section class='chart-block'>",
                    f"<h3>{html.escape(title)}</h3>",
                    f"<p>{html.escape(description)}</p>",
                    f"<div class='{wrap_class}'>",
                    section_html,
                    "</div>",
                    "</section>",
                ]
            )
        )
    return sections, include_plotlyjs


def write_png_exports(
    figures: list[tuple[str, str, str, go.Figure]],
    output_dir: Path,
) -> None:
    for slug, _title, _description, figure in figures:
        output_path = output_dir / f"{slug}.png"
        width = 1600
        height = 900
        if slug == "artifact_categories":
            width = 1200
            height = 1400
        write_image(figure, output_path, format="png", width=width, height=height, scale=2)


def friendly_human_label(value: str) -> str:
    lower = value.lower()
    if value == "abbyy":
        return "ABBYY"
    if "chandra" in lower:
        return "PaddleOCR + Chandra"
    if "paddlevl" in lower:
        return "PaddleOCR + PaddleVL"
    if "deepseek" in lower:
        return "PaddleOCR + DeepSeek"
    if "olmocr" in lower:
        return "PaddleOCR + olmOCR"
    if "ppcor-preproccess-3x2-docparse" in lower:
        return "PaddleOCR (No VL)"
    return f"PaddleOCR + {value}"


def friendly_whole_page_label(value: str) -> str:
    lower = value.strip().lower()
    mapping = {
        "preprocess only": "PaddleOCR + Preprocess only",
        "tiling only": "PaddleOCR (No VL)",
        "deepseek": "PaddleOCR + DeepSeek",
        "olmocr": "PaddleOCR + olmOCR",
        "chandra": "PaddleOCR + Chandra",
        "paddlevl": "PaddleOCR + PaddleVL",
    }
    return mapping.get(lower, f"PaddleOCR + {value}")


def parse_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else 0


def parse_float(value: object) -> float:
    if value is None:
        return 0.0
    if isinstance(value, float):
        return value
    if isinstance(value, int):
        return float(value)
    text = str(value).strip().replace("%", "")
    return float(text) if text else 0.0


def percentage(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return (numerator / denominator) * 100.0


def hex_to_rgba(value: str, alpha: float) -> str:
    color = value.lstrip("#")
    if len(color) != 6:
        return f"rgba(255,255,255,{alpha})"
    red = int(color[0:2], 16)
    green = int(color[2:4], 16)
    blue = int(color[4:6], 16)
    return f"rgba({red},{green},{blue},{alpha})"


if __name__ == "__main__":
    main()
