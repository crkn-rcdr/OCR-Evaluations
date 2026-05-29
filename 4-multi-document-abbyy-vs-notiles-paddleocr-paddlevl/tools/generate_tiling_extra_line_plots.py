#!/usr/bin/env python3
"""Generate charts for the 3x2 vs 3x1 extra-line comparison workbook.

Default usage:

  python tools/generate_tiling_extra_line_plots.py

This reads:

- test-results/tiling_3x2_vs_3x1_extra_lines.xlsx

and writes:

- test-results/plots/tiling_3x2_vs_3x1_extra_lines_dashboard.html
- test-results/plots/png/tiling_3x2_vs_3x1_*.png
"""

from __future__ import annotations

import argparse
import html
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from plotly import graph_objects as go
from plotly.io import to_html, write_image


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "test-results" / "tiling_3x2_vs_3x1_extra_lines.xlsx"
DEFAULT_OUTPUT_HTML = REPO_ROOT / "test-results" / "plots" / "tiling_3x2_vs_3x1_extra_lines_dashboard.html"
DEFAULT_OUTPUT_PNG_DIR = REPO_ROOT / "test-results" / "plots" / "png"
LINE_TYPES = ["separator/symbol", "short fragment", "text line"]


@dataclass(frozen=True)
class DocumentSummary:
    document: str
    page_count: int
    extra_lines_3x2: int
    extra_lines_3x1: int
    extra_words_3x2: int
    extra_words_3x1: int


@dataclass(frozen=True)
class PageSummary:
    document: str
    page: str
    extra_lines_3x2: int
    extra_lines_3x1: int

    @property
    def label(self) -> str:
        return f"{self.document} / {self.page}"

    @property
    def delta(self) -> int:
        return self.extra_lines_3x2 - self.extra_lines_3x1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help="Input XLSX workbook path.")
    parser.add_argument("--output-html", type=Path, default=DEFAULT_OUTPUT_HTML, help="Output dashboard HTML path.")
    parser.add_argument("--output-png-dir", type=Path, default=DEFAULT_OUTPUT_PNG_DIR, help="Output PNG directory.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    output_html = args.output_html.resolve()
    output_png_dir = args.output_png_dir.resolve()
    output_html.parent.mkdir(parents=True, exist_ok=True)
    output_png_dir.mkdir(parents=True, exist_ok=True)

    overall_totals, document_rows, page_rows = load_workbook_data(workbook_path)
    figures = [
        (
            "tiling_3x2_vs_3x1_document_extra_lines",
            "Document Extra-Line Totals",
            "Unique-only nonblank line counts by document. Higher bars mean that tiling variant produced more lines that the other variant did not contain.",
            build_document_extra_lines_figure(document_rows),
        ),
        (
            "tiling_3x2_vs_3x1_document_extra_words",
            "Document Extra-Word Totals",
            "Word totals contained inside the unique-only lines for each tiling variant.",
            build_document_extra_words_figure(document_rows),
        ),
        (
            "tiling_3x2_vs_3x1_extra_line_types",
            "Extra-Line Type Mix",
            "Overall mix of separator/symbol lines, short fragments, and longer text lines among the lines unique to each tiling variant.",
            build_type_mix_figure(overall_totals),
        ),
        (
            "tiling_3x2_vs_3x1_page_extra_line_delta",
            "Largest Page-Level Extra-Line Deltas",
            "Pages with the biggest gap in unique-only line counts. Positive bars mean 3x2 produced more unique-only lines; negative bars mean 3x1 did.",
            build_page_delta_figure(page_rows),
        ),
        (
            "tiling_3x2_vs_3x1_page_extra_line_rank",
            "Page Extra-Line Totals",
            "All pages sorted by combined unique-only line count, to show where the tiling variants diverge most strongly.",
            build_page_rank_figure(page_rows),
        ),
    ]

    output_html.write_text(build_dashboard_html(figures, workbook_path, overall_totals), encoding="utf-8")
    write_png_exports(figures, output_png_dir)
    print(f"Wrote {output_html}")
    print(f"Wrote PNG charts to {output_png_dir}")


def load_workbook_data(path: Path) -> tuple[dict[str, tuple[int, int]], list[DocumentSummary], list[PageSummary]]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["Overall Summary"]
    overall_totals: dict[str, tuple[int, int]] = {}
    document_summary_row = None
    page_summary_row = None

    numeric_summary_labels = {
        "document count",
        "page count",
        "nonblank line count",
        "word count",
        "shared matched lines",
        "extra unique-only lines",
        "extra unique-only words",
        "pages with more unique-only lines",
        "pages tied on unique-only lines",
    }
    for row in range(2, worksheet.max_row + 1):
        label = str(worksheet.cell(row, 1).value or "").strip()
        if not label:
            continue
        if label == "document summary":
            document_summary_row = row
            continue
        if label == "page summary":
            page_summary_row = row
            break
        if document_summary_row is None and (
            label in numeric_summary_labels or label.endswith(" extra lines")
        ):
            overall_totals[label] = (
                parse_int(worksheet.cell(row, 2).value),
                parse_int(worksheet.cell(row, 3).value),
            )

    if document_summary_row is None or page_summary_row is None:
        raise ValueError(f"Could not find document/page summary tables in {path}")

    doc_header_row = document_summary_row + 1
    doc_header = [
        str(worksheet.cell(doc_header_row, column).value or "").strip()
        for column in range(1, worksheet.max_column + 1)
    ]
    doc_index = {value: index for index, value in enumerate(doc_header, start=1) if value}
    document_rows: list[DocumentSummary] = []
    for row in range(doc_header_row + 1, worksheet.max_row + 1):
        document = str(worksheet.cell(row, doc_index["document"]).value or "").strip()
        if not document:
            break
        document_rows.append(
            DocumentSummary(
                document=document,
                page_count=parse_int(worksheet.cell(row, doc_index["page count"]).value),
                extra_lines_3x2=parse_int(worksheet.cell(row, doc_index["3x2 only lines"]).value),
                extra_lines_3x1=parse_int(worksheet.cell(row, doc_index["3x1 only lines"]).value),
                extra_words_3x2=parse_int(worksheet.cell(row, doc_index["3x2 only words"]).value),
                extra_words_3x1=parse_int(worksheet.cell(row, doc_index["3x1 only words"]).value),
            )
        )

    page_header_row = page_summary_row + 1
    page_header = [
        str(worksheet.cell(page_header_row, column).value or "").strip()
        for column in range(1, worksheet.max_column + 1)
    ]
    page_index = {value: index for index, value in enumerate(page_header, start=1) if value}
    page_rows: list[PageSummary] = []
    for row in range(page_header_row + 1, worksheet.max_row + 1):
        document = str(worksheet.cell(row, page_index["document"]).value or "").strip()
        if not document:
            break
        page_rows.append(
            PageSummary(
                document=document,
                page=str(worksheet.cell(row, page_index["page"]).value or "").strip(),
                extra_lines_3x2=parse_int(worksheet.cell(row, page_index["3x2 only lines"]).value),
                extra_lines_3x1=parse_int(worksheet.cell(row, page_index["3x1 only lines"]).value),
            )
        )
    return overall_totals, document_rows, page_rows


def build_document_extra_lines_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="3x2 only lines", x=labels, y=[row.extra_lines_3x2 for row in rows], marker_color="#1982C4"))
    figure.add_trace(go.Bar(name="3x1 only lines", x=labels, y=[row.extra_lines_3x1 for row in rows], marker_color="#55A630"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Unique-only line count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_document_extra_words_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="3x2 only words", x=labels, y=[row.extra_words_3x2 for row in rows], marker_color="#BC4749"))
    figure.add_trace(go.Bar(name="3x1 only words", x=labels, y=[row.extra_words_3x1 for row in rows], marker_color="#6A4C93"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Words inside unique-only lines",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_type_mix_figure(overall_totals: dict[str, tuple[int, int]]) -> go.Figure:
    labels = LINE_TYPES
    values_3x2 = [overall_totals.get(f"{label} extra lines", (0, 0))[0] for label in labels]
    values_3x1 = [overall_totals.get(f"{label} extra lines", (0, 0))[1] for label in labels]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="3x2", x=labels, y=values_3x2, marker_color="#1982C4"))
    figure.add_trace(go.Bar(name="3x1", x=labels, y=values_3x1, marker_color="#55A630"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Unique-only line count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=80),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_delta_figure(rows: list[PageSummary]) -> go.Figure:
    ordered = sorted(rows, key=lambda row: (abs(row.delta), row.document, row.page), reverse=True)[:25]
    labels = [row.label for row in ordered][::-1]
    delta_values = [row.delta for row in ordered][::-1]
    customdata = [[row.extra_lines_3x2, row.extra_lines_3x1] for row in ordered][::-1]
    positive_values = [value if value > 0 else None for value in delta_values]
    negative_values = [value if value < 0 else None for value in delta_values]
    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name="3x2 higher unique-only line count",
            x=positive_values,
            y=labels,
            orientation="h",
            marker_color="#1982C4",
            customdata=customdata,
            hovertemplate="%{y}<br>Delta (3x2 - 3x1): %{x}<br>3x2 only lines: %{customdata[0]}<br>3x1 only lines: %{customdata[1]}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Bar(
            name="3x1 higher unique-only line count",
            x=negative_values,
            y=labels,
            orientation="h",
            marker_color="#55A630",
            customdata=customdata,
            hovertemplate="%{y}<br>Delta (3x2 - 3x1): %{x}<br>3x2 only lines: %{customdata[0]}<br>3x1 only lines: %{customdata[1]}<extra></extra>",
        )
    )
    figure.add_vline(x=0, line_dash="dash", line_color="#666666")
    figure.update_layout(
        xaxis_title="Unique-only line delta",
        yaxis_title="Page",
        template="plotly_white",
        margin=dict(l=280, r=30, t=50, b=70),
        height=900,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_rank_figure(rows: list[PageSummary]) -> go.Figure:
    ordered = sorted(rows, key=lambda row: (row.extra_lines_3x2 + row.extra_lines_3x1, row.document, row.page), reverse=True)
    ranks = list(range(1, len(ordered) + 1))
    customdata = [[row.label] for row in ordered]
    figure = go.Figure()
    figure.add_trace(
        go.Scatter(
            name="3x2 only lines",
            x=ranks,
            y=[row.extra_lines_3x2 for row in ordered],
            mode="lines+markers",
            line=dict(color="#1982C4", width=2),
            marker=dict(size=6),
            customdata=customdata,
            hovertemplate="%{customdata[0]}<br>3x2 only lines: %{y}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Scatter(
            name="3x1 only lines",
            x=ranks,
            y=[row.extra_lines_3x1 for row in ordered],
            mode="lines+markers",
            line=dict(color="#55A630", width=2),
            marker=dict(size=6),
            customdata=customdata,
            hovertemplate="%{customdata[0]}<br>3x1 only lines: %{y}<extra></extra>",
        )
    )
    figure.update_layout(
        xaxis_title="Page rank by combined unique-only lines",
        yaxis_title="Unique-only line count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=70),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_dashboard_html(
    figures: list[tuple[str, str, str, go.Figure]],
    workbook_path: Path,
    overall_totals: dict[str, tuple[int, int]],
) -> str:
    generated = datetime.now(timezone.utc).isoformat()
    try:
        workbook_display = str(workbook_path.relative_to(REPO_ROOT))
    except ValueError:
        workbook_display = str(workbook_path)
    page_count = overall_totals.get("page count", (0, 0))[0]
    document_count = overall_totals.get("document count", (0, 0))[0]
    extra_lines_3x2 = overall_totals.get("extra unique-only lines", (0, 0))[0]
    extra_lines_3x1 = overall_totals.get("extra unique-only lines", (0, 0))[1]
    sections, _ = build_dashboard_sections(figures, include_plotlyjs=True)
    return "\n".join(
        [
            "<!DOCTYPE html>",
            "<html lang='en'>",
            "<head>",
            "<meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width, initial-scale=1'>",
            "<title>3x2 vs 3x1 Extra-Line Dashboard</title>",
            "<style>",
            "body { font-family: Segoe UI, Arial, sans-serif; margin: 0; background: #f7f7f7; color: #111; }",
            "main { max-width: 1400px; margin: 0 auto; padding: 24px; }",
            "header { margin-bottom: 24px; }",
            "h1 { margin: 0 0 8px; }",
            "p { line-height: 1.5; }",
            ".chart-block { background: #fff; border: 1px solid #ddd; padding: 20px; margin: 0 auto 20px; border-radius: 10px; max-width: 1260px; }",
            ".figure-wrap { width: min(100%, 1120px); margin: 0 auto; }",
            ".figure-wrap.narrow { width: min(100%, 900px); margin: 0 auto; }",
            "code { background: #f0f0f0; padding: 2px 5px; border-radius: 4px; }",
            "</style>",
            "</head>",
            "<body>",
            "<main>",
            "<header>",
            "<h1>3x2 vs 3x1 Extra-Line Dashboard</h1>",
            "<p>Dashboard for the paired Paddle tiling runs. It compares nonblank lines unique to each variant on the same page, without using ABBYY or human-transcribed ground truth.</p>",
            f"<p><strong>Generated:</strong> {html.escape(generated)}</p>",
            f"<p><strong>Inputs:</strong> <code>{html.escape(workbook_display)}</code></p>",
            f"<p><strong>Scope:</strong> {document_count} documents, {page_count} matched pages, {extra_lines_3x2:,} 3x2-only lines, {extra_lines_3x1:,} 3x1-only lines.</p>",
            "</header>",
            *sections,
            "</main>",
            "</body>",
            "</html>",
        ]
    )


def build_dashboard_sections(
    figures: list[tuple[str, str, str, go.Figure]],
    include_plotlyjs: bool,
) -> tuple[list[str], bool]:
    sections: list[str] = []
    for slug, title, description, figure in figures:
        section_html = to_html(
            figure,
            include_plotlyjs=include_plotlyjs,
            full_html=False,
            config={"responsive": True, "displaylogo": False},
        )
        include_plotlyjs = False
        wrap_class = "figure-wrap narrow" if slug.endswith(("delta", "rank")) else "figure-wrap"
        sections.append(
            "\n".join(
                [
                    "<section class='chart-block'>",
                    f"<h2>{html.escape(title)}</h2>",
                    f"<p>{html.escape(description)}</p>",
                    f"<div class='{wrap_class}'>",
                    section_html,
                    "</div>",
                    "</section>",
                ]
            )
        )
    return sections, include_plotlyjs


def write_png_exports(
    figures: list[tuple[str, str, str, go.Figure]],
    output_dir: Path,
) -> None:
    for slug, _title, _description, figure in figures:
        output_path = output_dir / f"{slug}.png"
        width = 1600
        height = 900
        if slug.endswith("page_extra_line_delta"):
            height = 1250
        write_image(figure, output_path, format="png", width=width, height=height, scale=2)


def parse_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else 0


if __name__ == "__main__":
    main()
