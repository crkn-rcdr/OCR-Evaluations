#!/usr/bin/env python3
"""Generate charts for a Paddle-vs-ABBYY word-coverage workbook."""

from __future__ import annotations

import argparse
import html
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from plotly import graph_objects as go
from plotly.io import to_html, write_image


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "test-results" / "paddle_variants_abbyy_word_coverage.xlsx"
DEFAULT_OUTPUT_HTML = REPO_ROOT / "test-results" / "plots" / "paddle_variants_abbyy_word_coverage_dashboard.html"
DEFAULT_OUTPUT_PNG_DIR = REPO_ROOT / "test-results" / "plots" / "png"
RUN_COLORS = ["#1982C4", "#55A630", "#BC4749", "#6A4C93", "#FF9F1C", "#2F4858"]


@dataclass(frozen=True)
class DocumentSummary:
    document: str
    page_count: int
    abbyy_words: int
    run_values: dict[str, dict[str, float]]


@dataclass(frozen=True)
class PageSummary:
    document: str
    page: str
    abbyy_words: int
    run_values: dict[str, dict[str, float]]

    @property
    def label(self) -> str:
        return f"{self.document} / {self.page}"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-html", type=Path, default=DEFAULT_OUTPUT_HTML)
    parser.add_argument("--output-png-dir", type=Path, default=DEFAULT_OUTPUT_PNG_DIR)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    output_html = args.output_html.resolve()
    output_png_dir = args.output_png_dir.resolve()
    output_html.parent.mkdir(parents=True, exist_ok=True)
    output_png_dir.mkdir(parents=True, exist_ok=True)

    run_labels, overall, documents, pages = load_workbook_data(workbook_path)
    slug_prefix = "paddle_variants_abbyy"
    figures = [
        (
            f"{slug_prefix}_word_coverage",
            "ABBYY Token Coverage by Document",
            "How much of ABBYY's token-occurrence count each Paddle variant captures, aggregated by document.",
            build_document_metric_figure(documents, run_labels, "coverage", "Percent of ABBYY token occurrences"),
        ),
        (
            f"{slug_prefix}_word_precision",
            "Token Precision vs ABBYY by Document",
            "What share of each run's words are also found in ABBYY. Higher values mean the run stays closer to ABBYY's vocabulary instead of overshooting it.",
            build_document_metric_figure(documents, run_labels, "precision", "Percent of run tokens also in ABBYY"),
        ),
        (
            f"{slug_prefix}_word_counts",
            "Document Word Counts vs ABBYY",
            "Absolute document word counts for all Paddle variants compared to ABBYY.",
            build_document_word_count_figure(documents, run_labels),
        ),
        (
            f"{slug_prefix}_coverage_spread",
            "Largest Page-Level Coverage Spreads",
            "Pages where the gap between the best and worst Paddle variant coverage against ABBYY is largest.",
            build_page_coverage_spread_figure(pages, run_labels),
        ),
        (
            f"{slug_prefix}_word_distance",
            "Page Word Count Distance from ABBYY",
            "Absolute page word-count distance from ABBYY for each Paddle variant. Lower lines stay closer to ABBYY page size.",
            build_page_word_distance_figure(pages, run_labels),
        ),
    ]

    output_html.write_text(build_dashboard_html(figures, workbook_path, overall, run_labels), encoding="utf-8")
    write_png_exports(figures, output_png_dir)
    print(f"Wrote {output_html}")
    print(f"Wrote PNG charts to {output_png_dir}")


def load_workbook_data(path: Path) -> tuple[list[str], dict[str, tuple[str, ...]], list[DocumentSummary], list[PageSummary]]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["Overall Summary"]
    run_labels = [
        str(worksheet.cell(1, column).value or "").strip()
        for column in range(2, worksheet.max_column + 1)
        if str(worksheet.cell(1, column).value or "").strip() and str(worksheet.cell(1, column).value or "").strip() != "abbyy baseline"
    ]
    overall: dict[str, tuple[str, ...]] = {}
    doc_start = None
    page_start = None
    for row in range(2, worksheet.max_row + 1):
        label = worksheet.cell(row, 1).value
        if label == "document summary":
            doc_start = row + 1
            continue
        if label == "page summary":
            page_start = row + 1
            break
        if label:
            overall[str(label)] = tuple(str(worksheet.cell(row, column).value or "") for column in range(2, len(run_labels) + 3))
    if doc_start is None or page_start is None:
        raise ValueError(f"Missing summary tables in {path}")

    doc_header = [str(worksheet.cell(doc_start, c).value or "").strip() for c in range(1, worksheet.max_column + 1)]
    doc_index = {value: index for index, value in enumerate(doc_header, start=1) if value}
    documents: list[DocumentSummary] = []
    row = doc_start + 1
    while row <= worksheet.max_row:
        document = str(worksheet.cell(row, doc_index["document"]).value or "").strip()
        if not document:
            break
        run_values: dict[str, dict[str, float]] = {}
        for label in run_labels:
            run_values[label] = {
                "words": parse_int(worksheet.cell(row, doc_index[f"{label} words"]).value),
                "coverage": parse_pct(worksheet.cell(row, doc_index[f"{label} coverage %"]).value),
                "precision": parse_pct(worksheet.cell(row, doc_index[f"{label} precision %"]).value),
            }
        documents.append(
            DocumentSummary(
                document=document,
                page_count=parse_int(worksheet.cell(row, doc_index["page count"]).value),
                abbyy_words=parse_int(worksheet.cell(row, doc_index["abbyy words"]).value),
                run_values=run_values,
            )
        )
        row += 1

    page_header = [str(worksheet.cell(page_start, c).value or "").strip() for c in range(1, worksheet.max_column + 1)]
    page_index = {value: index for index, value in enumerate(page_header, start=1) if value}
    pages: list[PageSummary] = []
    row = page_start + 1
    while row <= worksheet.max_row:
        document = str(worksheet.cell(row, page_index["document"]).value or "").strip()
        if not document:
            break
        run_values: dict[str, dict[str, float]] = {}
        for label in run_labels:
            run_values[label] = {
                "words": parse_int(worksheet.cell(row, page_index[f"{label} words"]).value),
                "coverage": parse_pct(worksheet.cell(row, page_index[f"{label} coverage %"]).value),
                "precision": parse_pct(worksheet.cell(row, page_index[f"{label} precision %"]).value),
                "word_delta": parse_int(worksheet.cell(row, page_index[f"{label} word count delta vs ABBYY"]).value),
            }
        pages.append(
            PageSummary(
                document=document,
                page=str(worksheet.cell(row, page_index["page"]).value or "").strip(),
                abbyy_words=parse_int(worksheet.cell(row, page_index["abbyy words"]).value),
                run_values=run_values,
            )
        )
        row += 1
    return run_labels, overall, documents, pages


def build_document_metric_figure(
    rows: list[DocumentSummary],
    run_labels: list[str],
    metric: str,
    yaxis_title: str,
) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    for index, label in enumerate(run_labels):
        figure.add_trace(
            go.Bar(
                name=label,
                x=labels,
                y=[row.run_values[label][metric] for row in rows],
                marker_color=RUN_COLORS[index % len(RUN_COLORS)],
            )
        )
    figure.update_layout(
        barmode="group",
        yaxis_title=yaxis_title,
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_document_word_count_figure(rows: list[DocumentSummary], run_labels: list[str]) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="ABBYY", x=labels, y=[row.abbyy_words for row in rows], marker_color="#6C757D"))
    for index, label in enumerate(run_labels):
        figure.add_trace(
            go.Bar(
                name=label,
                x=labels,
                y=[row.run_values[label]["words"] for row in rows],
                marker_color=RUN_COLORS[index % len(RUN_COLORS)],
            )
        )
    figure.update_layout(
        barmode="group",
        yaxis_title="Document word count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_coverage_spread_figure(rows: list[PageSummary], run_labels: list[str]) -> go.Figure:
    def spread(page: PageSummary) -> float:
        values = [page.run_values[label]["coverage"] for label in run_labels]
        return max(values) - min(values)

    ordered = sorted(rows, key=lambda page: (spread(page), page.document, page.page), reverse=True)[:20]
    labels = [page.label for page in ordered][::-1]
    figure = go.Figure()
    for index, label in enumerate(run_labels):
        figure.add_trace(
            go.Bar(
                name=label,
                x=[page.run_values[label]["coverage"] for page in ordered][::-1],
                y=labels,
                orientation="h",
                marker_color=RUN_COLORS[index % len(RUN_COLORS)],
                hovertemplate=f"%{{y}}<br>{label} coverage: %{{x:.2f}}%<extra></extra>",
            )
        )
    figure.update_layout(
        barmode="group",
        xaxis_title="Percent of ABBYY token occurrences",
        yaxis_title="Page",
        template="plotly_white",
        margin=dict(l=280, r=30, t=50, b=70),
        height=900,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_word_distance_figure(rows: list[PageSummary], run_labels: list[str]) -> go.Figure:
    ordered = sorted(
        rows,
        key=lambda page: (
            max(abs(page.run_values[label]["word_delta"]) for label in run_labels),
            page.document,
            page.page,
        ),
        reverse=True,
    )
    ranks = list(range(1, len(ordered) + 1))
    figure = go.Figure()
    for index, label in enumerate(run_labels):
        figure.add_trace(
            go.Scatter(
                name=label,
                x=ranks,
                y=[abs(page.run_values[label]["word_delta"]) for page in ordered],
                mode="lines+markers",
                line=dict(color=RUN_COLORS[index % len(RUN_COLORS)], width=2),
                marker=dict(size=6),
                customdata=[[page.label, page.abbyy_words, page.run_values[label]["words"]] for page in ordered],
                hovertemplate=f"%{{customdata[0]}}<br>ABBYY words: %{{customdata[1]}}<br>{label} words: %{{customdata[2]}}<br>Absolute distance: %{{y}}<extra></extra>",
            )
        )
    figure.update_layout(
        xaxis_title="Page rank by largest word-count distance from ABBYY",
        yaxis_title="Absolute word-count distance from ABBYY",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=70),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_dashboard_html(
    figures: list[tuple[str, str, str, go.Figure]],
    workbook_path: Path,
    overall: dict[str, tuple[str, ...]],
    run_labels: list[str],
) -> str:
    generated = datetime.now(timezone.utc).isoformat()
    try:
        workbook_display = str(workbook_path.relative_to(REPO_ROOT))
    except ValueError:
        workbook_display = str(workbook_path)
    page_count = overall.get("page count", ("0",))[0]
    document_count = overall.get("document count", ("0",))[0]
    sections = build_dashboard_sections(figures)
    return "\n".join(
        [
            "<!DOCTYPE html>",
            "<html lang='en'>",
            "<head>",
            "<meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width, initial-scale=1'>",
            "<title>Paddle Variant ABBYY Word Coverage Dashboard</title>",
            "<style>",
            "body { font-family: Segoe UI, Arial, sans-serif; margin: 0; background: #f7f7f7; color: #111; }",
            "main { max-width: 1400px; margin: 0 auto; padding: 24px; }",
            "header { margin-bottom: 24px; }",
            "h1 { margin: 0 0 8px; }",
            "p { line-height: 1.5; }",
            ".chart-block { background: #fff; border: 1px solid #ddd; padding: 20px; margin: 0 auto 20px; border-radius: 10px; max-width: 1260px; }",
            ".figure-wrap { width: min(100%, 1120px); margin: 0 auto; }",
            ".figure-wrap.narrow { width: min(100%, 900px); margin: 0 auto; }",
            "code { background: #f0f0f0; padding: 2px 5px; border-radius: 4px; }",
            "</style>",
            "</head>",
            "<body>",
            "<main>",
            "<header>",
            "<h1>Paddle Variant ABBYY Word Coverage Dashboard</h1>",
            f"<p>Dashboard comparing {', '.join(html.escape(label) for label in run_labels)} against ABBYY token coverage and precision.</p>",
            f"<p><strong>Generated:</strong> {html.escape(generated)}</p>",
            f"<p><strong>Input:</strong> <code>{html.escape(workbook_display)}</code></p>",
            f"<p><strong>Scope:</strong> {html.escape(document_count)} documents, {html.escape(page_count)} matched pages.</p>",
            "</header>",
            *sections,
            "</main>",
            "</body>",
            "</html>",
        ]
    )


def build_dashboard_sections(figures: list[tuple[str, str, str, go.Figure]]) -> list[str]:
    sections: list[str] = []
    include_plotlyjs = True
    for slug, title, description, figure in figures:
        section_html = to_html(
            figure,
            include_plotlyjs=include_plotlyjs,
            full_html=False,
            config={"responsive": True, "displaylogo": False},
        )
        include_plotlyjs = False
        wrap_class = "figure-wrap narrow" if slug.endswith(("spread", "distance")) else "figure-wrap"
        sections.append(
            "\n".join(
                [
                    "<section class='chart-block'>",
                    f"<h2>{html.escape(title)}</h2>",
                    f"<p>{html.escape(description)}</p>",
                    f"<div class='{wrap_class}'>",
                    section_html,
                    "</div>",
                    "</section>",
                ]
            )
        )
    return sections


def write_png_exports(figures: list[tuple[str, str, str, go.Figure]], output_dir: Path) -> None:
    for slug, _title, _description, figure in figures:
        output_path = output_dir / f"{slug}.png"
        width = 1600
        height = 900
        if slug.endswith("coverage_spread"):
            height = 1250
        write_image(figure, output_path, format="png", width=width, height=height, scale=2)


def parse_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else 0


def parse_pct(value: object) -> float:
    if value is None:
        return 0.0
    text = str(value).strip().replace("%", "").replace(",", "")
    return float(text) if text else 0.0


if __name__ == "__main__":
    main()
