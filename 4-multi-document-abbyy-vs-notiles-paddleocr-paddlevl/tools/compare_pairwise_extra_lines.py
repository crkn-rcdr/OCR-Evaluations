#!/usr/bin/env python3
"""Compare extra lines unique to any two OCR runs page by page.

Default usage compares the no-tiling run in this repo against the 3x1 tiled run:

  python tools/compare_pairwise_extra_lines.py

Example 3x2 comparison:

  python tools/compare_pairwise_extra_lines.py ^
    --run-a-label "3x2" ^
    --run-a-dir ..\\2-multi-document-abbyy-vs-tiled3x2+docparse-paddleocr-paddlevl\\test-data\\preprocessing-tiling_3x2_dp+paddleocr+vl+clean ^
    --run-b-label "no tiling" ^
    --run-b-dir test-data\\preprocess-notiles-paddleocr-paddlevl ^
    --output test-results\\tiling_3x2_vs_notiling_extra_lines.xlsx
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment


REPO_ROOT = Path(__file__).resolve().parent.parent
EVAL_ROOT = REPO_ROOT.parent
DEFAULT_RUN_A_LABEL = "no tiling"
DEFAULT_RUN_B_LABEL = "3x1"
DEFAULT_RUN_A_DIR = REPO_ROOT / "test-data" / "preprocess-notiles-paddleocr-paddlevl"
DEFAULT_RUN_B_DIR = (
    EVAL_ROOT
    / "3-multi-document-abbyy-vs-tiled3x1+docparse-paddleocr-paddlevl"
    / "test-data"
    / "preprocessing-tiling_3x1_dp+paddleocr+vl+clean"
)
DEFAULT_OUTPUT = REPO_ROOT / "test-results" / "notiling_vs_3x1_extra_lines.xlsx"

PAGE_NUMBER_RE = re.compile(r"(\d+)$")
WORD_RE = re.compile(r"[A-Za-z0-9]+(?:[â€™'][A-Za-z0-9]+)?")
LINE_TYPE_ORDER = ["separator/symbol", "short fragment", "text line"]
ILLEGAL_EXCEL_CHAR_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")


@dataclass(frozen=True)
class LineRecord:
    line_no: int
    text: str
    normalized: str
    word_count: int
    line_type: str


@dataclass(frozen=True)
class PageSummary:
    document: str
    page: str
    source_a: Path
    source_b: Path
    line_count_a: int
    line_count_b: int
    word_count_a: int
    word_count_b: int
    shared_line_count: int
    extra_line_count_a: int
    extra_line_count_b: int
    extra_word_count_a: int
    extra_word_count_b: int
    extra_type_counts_a: dict[str, int]
    extra_type_counts_b: dict[str, int]

    @property
    def extra_line_delta(self) -> int:
        return self.extra_line_count_a - self.extra_line_count_b


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-a-label", default=DEFAULT_RUN_A_LABEL, help="Display label for run A.")
    parser.add_argument("--run-a-dir", type=Path, default=DEFAULT_RUN_A_DIR, help="Root directory of run A TXT outputs.")
    parser.add_argument("--run-b-label", default=DEFAULT_RUN_B_LABEL, help="Display label for run B.")
    parser.add_argument("--run-b-dir", type=Path, default=DEFAULT_RUN_B_DIR, help="Root directory of run B TXT outputs.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output XLSX workbook path.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    run_a_label = normalize_label(args.run_a_label)
    run_b_label = normalize_label(args.run_b_label)
    summaries, extra_a_rows, extra_b_rows = compare_runs(
        args.run_a_dir.resolve(),
        args.run_b_dir.resolve(),
    )
    write_workbook(
        args.output.resolve(),
        build_overall_sheet_rows(
            summaries,
            run_a_label,
            run_b_label,
            args.run_a_dir.resolve(),
            args.run_b_dir.resolve(),
        ),
        build_extra_sheet_rows(extra_a_rows),
        build_extra_sheet_rows(extra_b_rows),
        run_a_label,
        run_b_label,
    )
    print(f"Wrote {args.output.resolve()}")
    print(f"Compared pages: {len(summaries)}")


def normalize_label(value: str) -> str:
    cleaned = re.sub(r"\s+", " ", value).strip()
    if not cleaned:
        raise ValueError("Run labels must not be empty.")
    return cleaned


def ocr_txt_files_in_dir(base_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(base_dir.glob("*.txt")):
        name = path.name.lower()
        if name.endswith(".post_ocr_vl_progress.txt"):
            continue
        if "b4-cleaned" in name:
            continue
        files.append(path)
    return files


def page_key_for_path(path: Path) -> str:
    match = PAGE_NUMBER_RE.search(path.stem)
    if match:
        return f"page:{int(match.group(1))}"
    return f"name:{path.name.casefold()}"


def normalize_line(text: str) -> str:
    value = unicodedata.normalize("NFKC", text.replace("\ufeff", ""))
    value = re.sub(r"\s+", " ", value).strip().casefold()
    return value


def classify_line(text: str) -> str:
    stripped = text.strip()
    word_count = len(WORD_RE.findall(stripped))
    alnum_count = sum(char.isalnum() for char in stripped)
    alpha_count = sum(char.isalpha() for char in stripped)
    if alnum_count == 0:
        return "separator/symbol"
    if word_count <= 3 or alpha_count < 8:
        return "short fragment"
    return "text line"


def read_line_records(path: Path) -> tuple[list[LineRecord], int]:
    text = path.read_text(encoding="utf-8", errors="replace")
    records: list[LineRecord] = []
    for line_no, raw_line in enumerate(text.splitlines(), 1):
        stripped = raw_line.strip()
        if not stripped:
            continue
        records.append(
            LineRecord(
                line_no=line_no,
                text=stripped,
                normalized=normalize_line(stripped),
                word_count=len(WORD_RE.findall(stripped)),
                line_type=classify_line(stripped),
            )
        )
    return records, len(text.split())


def count_line_types(records: list[LineRecord]) -> Counter[str]:
    counts: Counter[str] = Counter()
    for record in records:
        counts[record.line_type] += 1
    return counts


def select_extra_records(records: list[LineRecord], extra_counts: Counter[str]) -> list[LineRecord]:
    remaining = Counter(extra_counts)
    selected: list[LineRecord] = []
    for record in records:
        if remaining[record.normalized] <= 0:
            continue
        selected.append(record)
        remaining[record.normalized] -= 1
    return selected


def compare_page(run_a_path: Path, run_b_path: Path) -> tuple[PageSummary, list[dict[str, object]], list[dict[str, object]]]:
    records_a, word_count_a = read_line_records(run_a_path)
    records_b, word_count_b = read_line_records(run_b_path)
    counts_a = Counter(record.normalized for record in records_a)
    counts_b = Counter(record.normalized for record in records_b)

    shared_line_count = sum(min(counts_a[key], counts_b[key]) for key in (counts_a.keys() | counts_b.keys()))
    extra_counts_a = counts_a - counts_b
    extra_counts_b = counts_b - counts_a
    extra_records_a = select_extra_records(records_a, extra_counts_a)
    extra_records_b = select_extra_records(records_b, extra_counts_b)
    extra_type_counts_a = count_line_types(extra_records_a)
    extra_type_counts_b = count_line_types(extra_records_b)

    document = run_a_path.parent.name
    page = run_a_path.name
    summary = PageSummary(
        document=document,
        page=page,
        source_a=run_a_path,
        source_b=run_b_path,
        line_count_a=len(records_a),
        line_count_b=len(records_b),
        word_count_a=word_count_a,
        word_count_b=word_count_b,
        shared_line_count=shared_line_count,
        extra_line_count_a=len(extra_records_a),
        extra_line_count_b=len(extra_records_b),
        extra_word_count_a=sum(record.word_count for record in extra_records_a),
        extra_word_count_b=sum(record.word_count for record in extra_records_b),
        extra_type_counts_a={line_type: int(extra_type_counts_a.get(line_type, 0)) for line_type in LINE_TYPE_ORDER},
        extra_type_counts_b={line_type: int(extra_type_counts_b.get(line_type, 0)) for line_type in LINE_TYPE_ORDER},
    )

    extra_a_rows = [build_extra_line_row(document, page, run_a_path, record) for record in extra_records_a]
    extra_b_rows = [build_extra_line_row(document, page, run_b_path, record) for record in extra_records_b]
    return summary, extra_a_rows, extra_b_rows


def build_extra_line_row(document: str, page: str, source_path: Path, record: LineRecord) -> dict[str, object]:
    return {
        "document": document,
        "page": page,
        "source_file": str(source_path),
        "line_no": record.line_no,
        "word_count": record.word_count,
        "line_type": record.line_type,
        "line_text": record.text,
        "normalized": record.normalized,
    }


def compare_runs(run_a_dir: Path, run_b_dir: Path) -> tuple[list[PageSummary], list[dict[str, object]], list[dict[str, object]]]:
    summaries: list[PageSummary] = []
    extra_a_rows: list[dict[str, object]] = []
    extra_b_rows: list[dict[str, object]] = []
    missing_documents: list[str] = []
    missing_pages: list[str] = []

    for run_a_doc_dir in sorted(path for path in run_a_dir.iterdir() if path.is_dir()):
        run_b_doc_dir = run_b_dir / run_a_doc_dir.name
        if not run_b_doc_dir.is_dir():
            missing_documents.append(run_a_doc_dir.name)
            continue

        run_b_by_name = {page_key_for_path(path): path for path in ocr_txt_files_in_dir(run_b_doc_dir)}
        run_a_pages = ocr_txt_files_in_dir(run_a_doc_dir)
        for run_a_path in run_a_pages:
            run_b_path = run_b_by_name.get(page_key_for_path(run_a_path))
            if run_b_path is None:
                missing_pages.append(f"{run_a_doc_dir.name}/{run_a_path.name}")
                continue
            summary, page_extra_a_rows, page_extra_b_rows = compare_page(run_a_path, run_b_path)
            summaries.append(summary)
            extra_a_rows.extend(page_extra_a_rows)
            extra_b_rows.extend(page_extra_b_rows)

        run_a_page_keys = {page_key_for_path(path) for path in run_a_pages}
        for run_b_path in ocr_txt_files_in_dir(run_b_doc_dir):
            if page_key_for_path(run_b_path) not in run_a_page_keys:
                missing_pages.append(f"{run_b_doc_dir.name}/{run_b_path.name}")

    if missing_documents or missing_pages:
        problems: list[str] = []
        if missing_documents:
            problems.append("Missing matching document directories: " + ", ".join(sorted(missing_documents)[:10]))
        if missing_pages:
            problems.append("Missing matching page TXT files: " + ", ".join(sorted(missing_pages)[:10]))
        raise FileNotFoundError("; ".join(problems))

    summaries.sort(key=lambda item: (item.document.casefold(), item.page.casefold()))
    return summaries, extra_a_rows, extra_b_rows


def build_overall_sheet_rows(
    summaries: list[PageSummary],
    run_a_label: str,
    run_b_label: str,
    run_a_dir: Path,
    run_b_dir: Path,
) -> list[list[object]]:
    total_line_count_a = sum(item.line_count_a for item in summaries)
    total_line_count_b = sum(item.line_count_b for item in summaries)
    total_word_count_a = sum(item.word_count_a for item in summaries)
    total_word_count_b = sum(item.word_count_b for item in summaries)
    total_shared_lines = sum(item.shared_line_count for item in summaries)
    total_extra_lines_a = sum(item.extra_line_count_a for item in summaries)
    total_extra_lines_b = sum(item.extra_line_count_b for item in summaries)
    total_extra_words_a = sum(item.extra_word_count_a for item in summaries)
    total_extra_words_b = sum(item.extra_word_count_b for item in summaries)
    pages_a_more = sum(1 for item in summaries if item.extra_line_count_a > item.extra_line_count_b)
    pages_b_more = sum(1 for item in summaries if item.extra_line_count_b > item.extra_line_count_a)
    pages_tied = len(summaries) - pages_a_more - pages_b_more
    total_type_counts_a: Counter[str] = Counter()
    total_type_counts_b: Counter[str] = Counter()
    document_totals: dict[str, dict[str, object]] = defaultdict(
        lambda: {
            "page_count": 0,
            "line_count_a": 0,
            "line_count_b": 0,
            "word_count_a": 0,
            "word_count_b": 0,
            "shared_line_count": 0,
            "extra_line_count_a": 0,
            "extra_line_count_b": 0,
            "extra_word_count_a": 0,
            "extra_word_count_b": 0,
            "type_counts_a": Counter(),
            "type_counts_b": Counter(),
        }
    )

    for summary in summaries:
        total_type_counts_a.update(summary.extra_type_counts_a)
        total_type_counts_b.update(summary.extra_type_counts_b)
        bucket = document_totals[summary.document]
        bucket["page_count"] = int(bucket["page_count"]) + 1
        bucket["line_count_a"] = int(bucket["line_count_a"]) + summary.line_count_a
        bucket["line_count_b"] = int(bucket["line_count_b"]) + summary.line_count_b
        bucket["word_count_a"] = int(bucket["word_count_a"]) + summary.word_count_a
        bucket["word_count_b"] = int(bucket["word_count_b"]) + summary.word_count_b
        bucket["shared_line_count"] = int(bucket["shared_line_count"]) + summary.shared_line_count
        bucket["extra_line_count_a"] = int(bucket["extra_line_count_a"]) + summary.extra_line_count_a
        bucket["extra_line_count_b"] = int(bucket["extra_line_count_b"]) + summary.extra_line_count_b
        bucket["extra_word_count_a"] = int(bucket["extra_word_count_a"]) + summary.extra_word_count_a
        bucket["extra_word_count_b"] = int(bucket["extra_word_count_b"]) + summary.extra_word_count_b
        bucket["type_counts_a"].update(summary.extra_type_counts_a)
        bucket["type_counts_b"].update(summary.extra_type_counts_b)

    rows: list[list[object]] = [
        ["summary metric", run_a_label, run_b_label],
        ["dataset root", str(run_a_dir), str(run_b_dir)],
        ["document count", len(document_totals), len(document_totals)],
        ["page count", len(summaries), len(summaries)],
        ["nonblank line count", total_line_count_a, total_line_count_b],
        ["word count", total_word_count_a, total_word_count_b],
        ["shared matched lines", total_shared_lines, total_shared_lines],
        ["extra unique-only lines", total_extra_lines_a, total_extra_lines_b],
        ["extra unique-only words", total_extra_words_a, total_extra_words_b],
        ["pages with more unique-only lines", pages_a_more, pages_b_more],
        ["pages tied on unique-only lines", pages_tied, pages_tied],
        [
            "counting note",
            f"Extra lines are nonblank normalized lines that appear more times in {run_a_label} than {run_b_label} on the same page.",
            f"Extra lines are nonblank normalized lines that appear more times in {run_b_label} than {run_a_label} on the same page.",
        ],
    ]
    for line_type in LINE_TYPE_ORDER:
        rows.append(
            [
                f"{line_type} extra lines",
                int(total_type_counts_a.get(line_type, 0)),
                int(total_type_counts_b.get(line_type, 0)),
            ]
        )

    rows.append([])
    rows.append(["document summary"])
    document_header: list[object] = [
        "document",
        "page count",
        f"{run_a_label} nonblank lines",
        f"{run_b_label} nonblank lines",
        f"{run_a_label} words",
        f"{run_b_label} words",
        "shared matched lines",
        f"{run_a_label} only lines",
        f"{run_b_label} only lines",
        f"extra line delta ({run_a_label}-{run_b_label})",
        f"{run_a_label} only words",
        f"{run_b_label} only words",
    ]
    for line_type in LINE_TYPE_ORDER:
        document_header.append(f"{run_a_label} {line_type}")
        document_header.append(f"{run_b_label} {line_type}")
    rows.append(document_header)

    for document in sorted(document_totals):
        bucket = document_totals[document]
        row: list[object] = [
            document,
            int(bucket["page_count"]),
            int(bucket["line_count_a"]),
            int(bucket["line_count_b"]),
            int(bucket["word_count_a"]),
            int(bucket["word_count_b"]),
            int(bucket["shared_line_count"]),
            int(bucket["extra_line_count_a"]),
            int(bucket["extra_line_count_b"]),
            int(bucket["extra_line_count_a"]) - int(bucket["extra_line_count_b"]),
            int(bucket["extra_word_count_a"]),
            int(bucket["extra_word_count_b"]),
        ]
        for line_type in LINE_TYPE_ORDER:
            row.append(int(bucket["type_counts_a"].get(line_type, 0)))
            row.append(int(bucket["type_counts_b"].get(line_type, 0)))
        rows.append(row)

    rows.append([])
    rows.append(["page summary"])
    page_header: list[object] = [
        "document",
        "page",
        f"{run_a_label} source file",
        f"{run_b_label} source file",
        f"{run_a_label} nonblank lines",
        f"{run_b_label} nonblank lines",
        f"{run_a_label} words",
        f"{run_b_label} words",
        "shared matched lines",
        f"{run_a_label} only lines",
        f"{run_b_label} only lines",
        f"extra line delta ({run_a_label}-{run_b_label})",
        f"{run_a_label} only words",
        f"{run_b_label} only words",
    ]
    for line_type in LINE_TYPE_ORDER:
        page_header.append(f"{run_a_label} {line_type}")
        page_header.append(f"{run_b_label} {line_type}")
    rows.append(page_header)

    for summary in summaries:
        row: list[object] = [
            summary.document,
            summary.page,
            str(summary.source_a),
            str(summary.source_b),
            summary.line_count_a,
            summary.line_count_b,
            summary.word_count_a,
            summary.word_count_b,
            summary.shared_line_count,
            summary.extra_line_count_a,
            summary.extra_line_count_b,
            summary.extra_line_delta,
            summary.extra_word_count_a,
            summary.extra_word_count_b,
        ]
        for line_type in LINE_TYPE_ORDER:
            row.append(summary.extra_type_counts_a.get(line_type, 0))
            row.append(summary.extra_type_counts_b.get(line_type, 0))
        rows.append(row)
    return rows


def build_extra_sheet_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    output: list[list[object]] = [[
        "document",
        "page",
        "source file",
        "line no",
        "word count",
        "line type",
        "line text",
        "normalized",
    ]]
    for row in sorted(
        rows,
        key=lambda item: (
            str(item["document"]).casefold(),
            str(item["page"]).casefold(),
            int(item["line_no"]),
        ),
    ):
        output.append(
            [
                row["document"],
                row["page"],
                row["source_file"],
                row["line_no"],
                row["word_count"],
                row["line_type"],
                row["line_text"],
                row["normalized"],
            ]
        )
    return output


def safe_sheet_title(title: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", title).strip()
    return cleaned[:31] if len(cleaned) > 31 else cleaned


def write_workbook(
    output_path: Path,
    overall_rows: list[list[object]],
    extra_a_rows: list[list[object]],
    extra_b_rows: list[list[object]],
    run_a_label: str,
    run_b_label: str,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [
        ("Overall Summary", overall_rows),
        (safe_sheet_title(f"{run_a_label} Extra Lines"), extra_a_rows),
        (safe_sheet_title(f"{run_b_label} Extra Lines"), extra_b_rows),
    ]
    for title, rows in sheets:
        worksheet = workbook.create_sheet(title=title)
        for row in rows:
            worksheet.append([sanitize_excel_value(value) for value in row])
        format_sheet(worksheet)
    workbook.save(output_path)


def sanitize_excel_value(value: object) -> object:
    if isinstance(value, str):
        return ILLEGAL_EXCEL_CHAR_RE.sub("", value)
    return value


def format_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    wrap_columns = {"dataset root", "counting note", "source file", "line text", "normalized"}
    for column_cells in worksheet.iter_cols():
        max_length = 0
        header = str(column_cells[0].value or "")
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if header in wrap_columns:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")
            max_length = max(max_length, len(value))
        width = min(max(max_length + 2, 12), 90)
        worksheet.column_dimensions[column_cells[0].column_letter].width = width


if __name__ == "__main__":
    main()
