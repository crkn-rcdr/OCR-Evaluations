#!/usr/bin/env python3
"""Generate charts for a pairwise extra-line comparison workbook."""

from __future__ import annotations

import argparse
import html
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from plotly import graph_objects as go
from plotly.io import to_html, write_image


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "test-results" / "notiling_vs_3x1_extra_lines.xlsx"
DEFAULT_OUTPUT_HTML = REPO_ROOT / "test-results" / "plots" / "notiling_vs_3x1_extra_lines_dashboard.html"
DEFAULT_OUTPUT_PNG_DIR = REPO_ROOT / "test-results" / "plots" / "png"
LINE_TYPES = ["separator/symbol", "short fragment", "text line"]


@dataclass(frozen=True)
class DocumentSummary:
    document: str
    page_count: int
    extra_lines_a: int
    extra_lines_b: int
    extra_words_a: int
    extra_words_b: int


@dataclass(frozen=True)
class PageSummary:
    document: str
    page: str
    extra_lines_a: int
    extra_lines_b: int

    @property
    def label(self) -> str:
        return f"{self.document} / {self.page}"

    @property
    def delta(self) -> int:
        return self.extra_lines_a - self.extra_lines_b


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help="Input XLSX workbook path.")
    parser.add_argument("--output-html", type=Path, default=DEFAULT_OUTPUT_HTML, help="Output dashboard HTML path.")
    parser.add_argument("--output-png-dir", type=Path, default=DEFAULT_OUTPUT_PNG_DIR, help="Output PNG directory.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    output_html = args.output_html.resolve()
    output_png_dir = args.output_png_dir.resolve()
    output_html.parent.mkdir(parents=True, exist_ok=True)
    output_png_dir.mkdir(parents=True, exist_ok=True)

    label_a, label_b, overall_totals, document_rows, page_rows = load_workbook_data(workbook_path)
    prefix = slugify(f"{label_a}_vs_{label_b}")
    figures = [
        (
            f"{prefix}_document_extra_lines",
            "Document Extra-Line Totals",
            f"Unique-only nonblank line counts by document. Higher bars mean that OCR variant produced more lines that the other variant did not contain.",
            build_document_extra_lines_figure(document_rows, label_a, label_b),
        ),
        (
            f"{prefix}_document_extra_words",
            "Document Extra-Word Totals",
            "Word totals contained inside the unique-only lines for each OCR variant.",
            build_document_extra_words_figure(document_rows, label_a, label_b),
        ),
        (
            f"{prefix}_extra_line_types",
            "Extra-Line Type Mix",
            "Overall mix of separator/symbol lines, short fragments, and longer text lines among the lines unique to each OCR variant.",
            build_type_mix_figure(overall_totals, label_a, label_b),
        ),
        (
            f"{prefix}_page_extra_line_delta",
            "Largest Page-Level Extra-Line Deltas",
            f"Pages with the biggest gap in unique-only line counts. Positive bars mean {label_a} produced more unique-only lines; negative bars mean {label_b} did.",
            build_page_delta_figure(page_rows, label_a, label_b),
        ),
        (
            f"{prefix}_page_extra_line_rank",
            "Page Extra-Line Totals",
            "All pages sorted by combined unique-only line count, to show where the OCR variants diverge most strongly.",
            build_page_rank_figure(page_rows, label_a, label_b),
        ),
    ]

    output_html.write_text(build_dashboard_html(figures, workbook_path, overall_totals, label_a, label_b), encoding="utf-8")
    write_png_exports(figures, output_png_dir)
    print(f"Wrote {output_html}")
    print(f"Wrote PNG charts to {output_png_dir}")


def load_workbook_data(path: Path) -> tuple[str, str, dict[str, tuple[int, int]], list[DocumentSummary], list[PageSummary]]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["Overall Summary"]
    label_a = str(worksheet.cell(1, 2).value or "").strip()
    label_b = str(worksheet.cell(1, 3).value or "").strip()
    overall_totals: dict[str, tuple[int, int]] = {}
    document_summary_row = None
    page_summary_row = None

    numeric_summary_labels = {
        "document count",
        "page count",
        "nonblank line count",
        "word count",
        "shared matched lines",
        "extra unique-only lines",
        "extra unique-only words",
        "pages with more unique-only lines",
        "pages tied on unique-only lines",
    }
    for row in range(2, worksheet.max_row + 1):
        label = str(worksheet.cell(row, 1).value or "").strip()
        if not label:
            continue
        if label == "document summary":
            document_summary_row = row
            continue
        if label == "page summary":
            page_summary_row = row
            break
        if document_summary_row is None and (label in numeric_summary_labels or label.endswith(" extra lines")):
            overall_totals[label] = (
                parse_int(worksheet.cell(row, 2).value),
                parse_int(worksheet.cell(row, 3).value),
            )

    if document_summary_row is None or page_summary_row is None:
        raise ValueError(f"Could not find document/page summary tables in {path}")

    doc_header_row = document_summary_row + 1
    doc_header = [str(worksheet.cell(doc_header_row, column).value or "").strip() for column in range(1, worksheet.max_column + 1)]
    doc_index = {value: index for index, value in enumerate(doc_header, start=1) if value}
    document_rows: list[DocumentSummary] = []
    for row in range(doc_header_row + 1, worksheet.max_row + 1):
        document = str(worksheet.cell(row, doc_index["document"]).value or "").strip()
        if not document:
            break
        document_rows.append(
            DocumentSummary(
                document=document,
                page_count=parse_int(worksheet.cell(row, doc_index["page count"]).value),
                extra_lines_a=parse_int(worksheet.cell(row, doc_index[f"{label_a} only lines"]).value),
                extra_lines_b=parse_int(worksheet.cell(row, doc_index[f"{label_b} only lines"]).value),
                extra_words_a=parse_int(worksheet.cell(row, doc_index[f"{label_a} only words"]).value),
                extra_words_b=parse_int(worksheet.cell(row, doc_index[f"{label_b} only words"]).value),
            )
        )

    page_header_row = page_summary_row + 1
    page_header = [str(worksheet.cell(page_header_row, column).value or "").strip() for column in range(1, worksheet.max_column + 1)]
    page_index = {value: index for index, value in enumerate(page_header, start=1) if value}
    page_rows: list[PageSummary] = []
    for row in range(page_header_row + 1, worksheet.max_row + 1):
        document = str(worksheet.cell(row, page_index["document"]).value or "").strip()
        if not document:
            break
        page_rows.append(
            PageSummary(
                document=document,
                page=str(worksheet.cell(row, page_index["page"]).value or "").strip(),
                extra_lines_a=parse_int(worksheet.cell(row, page_index[f"{label_a} only lines"]).value),
                extra_lines_b=parse_int(worksheet.cell(row, page_index[f"{label_b} only lines"]).value),
            )
        )
    return label_a, label_b, overall_totals, document_rows, page_rows


def build_document_extra_lines_figure(rows: list[DocumentSummary], label_a: str, label_b: str) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name=f"{label_a} only lines", x=labels, y=[row.extra_lines_a for row in rows], marker_color="#1982C4"))
    figure.add_trace(go.Bar(name=f"{label_b} only lines", x=labels, y=[row.extra_lines_b for row in rows], marker_color="#55A630"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Unique-only line count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_document_extra_words_figure(rows: list[DocumentSummary], label_a: str, label_b: str) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name=f"{label_a} only words", x=labels, y=[row.extra_words_a for row in rows], marker_color="#BC4749"))
    figure.add_trace(go.Bar(name=f"{label_b} only words", x=labels, y=[row.extra_words_b for row in rows], marker_color="#6A4C93"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Words inside unique-only lines",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_type_mix_figure(overall_totals: dict[str, tuple[int, int]], label_a: str, label_b: str) -> go.Figure:
    labels = LINE_TYPES
    values_a = [overall_totals.get(f"{label} extra lines", (0, 0))[0] for label in labels]
    values_b = [overall_totals.get(f"{label} extra lines", (0, 0))[1] for label in labels]
    figure = go.Figure()
    figure.add_trace(go.Bar(name=label_a, x=labels, y=values_a, marker_color="#1982C4"))
    figure.add_trace(go.Bar(name=label_b, x=labels, y=values_b, marker_color="#55A630"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Unique-only line count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=80),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_delta_figure(rows: list[PageSummary], label_a: str, label_b: str) -> go.Figure:
    ordered = sorted(rows, key=lambda row: (abs(row.delta), row.document, row.page), reverse=True)[:25]
    labels = [row.label for row in ordered][::-1]
    delta_values = [row.delta for row in ordered][::-1]
    customdata = [[row.extra_lines_a, row.extra_lines_b] for row in ordered][::-1]
    positive_values = [value if value > 0 else None for value in delta_values]
    negative_values = [value if value < 0 else None for value in delta_values]
    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name=f"{label_a} higher unique-only line count",
            x=positive_values,
            y=labels,
            orientation="h",
            marker_color="#1982C4",
            customdata=customdata,
            hovertemplate=f"%{{y}}<br>Delta ({label_a} - {label_b}): %{{x}}<br>{label_a} only lines: %{{customdata[0]}}<br>{label_b} only lines: %{{customdata[1]}}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Bar(
            name=f"{label_b} higher unique-only line count",
            x=negative_values,
            y=labels,
            orientation="h",
            marker_color="#55A630",
            customdata=customdata,
            hovertemplate=f"%{{y}}<br>Delta ({label_a} - {label_b}): %{{x}}<br>{label_a} only lines: %{{customdata[0]}}<br>{label_b} only lines: %{{customdata[1]}}<extra></extra>",
        )
    )
    figure.add_vline(x=0, line_dash="dash", line_color="#666666")
    figure.update_layout(
        xaxis_title="Unique-only line delta",
        yaxis_title="Page",
        template="plotly_white",
        margin=dict(l=280, r=30, t=50, b=70),
        height=900,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_rank_figure(rows: list[PageSummary], label_a: str, label_b: str) -> go.Figure:
    ordered = sorted(rows, key=lambda row: (row.extra_lines_a + row.extra_lines_b, row.document, row.page), reverse=True)
    ranks = list(range(1, len(ordered) + 1))
    customdata = [[row.label] for row in ordered]
    figure = go.Figure()
    figure.add_trace(
        go.Scatter(
            name=f"{label_a} only lines",
            x=ranks,
            y=[row.extra_lines_a for row in ordered],
            mode="lines+markers",
            line=dict(color="#1982C4", width=2),
            marker=dict(size=6),
            customdata=customdata,
            hovertemplate=f"%{{customdata[0]}}<br>{label_a} only lines: %{{y}}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Scatter(
            name=f"{label_b} only lines",
            x=ranks,
            y=[row.extra_lines_b for row in ordered],
            mode="lines+markers",
            line=dict(color="#55A630", width=2),
            marker=dict(size=6),
            customdata=customdata,
            hovertemplate=f"%{{customdata[0]}}<br>{label_b} only lines: %{{y}}<extra></extra>",
        )
    )
    figure.update_layout(
        xaxis_title="Page rank by combined unique-only lines",
        yaxis_title="Unique-only line count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=70),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_dashboard_html(
    figures: list[tuple[str, str, str, go.Figure]],
    workbook_path: Path,
    overall_totals: dict[str, tuple[int, int]],
    label_a: str,
    label_b: str,
) -> str:
    generated = datetime.now(timezone.utc).isoformat()
    try:
        workbook_display = str(workbook_path.relative_to(REPO_ROOT))
    except ValueError:
        workbook_display = str(workbook_path)
    page_count = overall_totals.get("page count", (0, 0))[0]
    document_count = overall_totals.get("document count", (0, 0))[0]
    extra_lines_a = overall_totals.get("extra unique-only lines", (0, 0))[0]
    extra_lines_b = overall_totals.get("extra unique-only lines", (0, 0))[1]
    sections = build_dashboard_sections(figures)
    return "\n".join(
        [
            "<!DOCTYPE html>",
            "<html lang='en'>",
            "<head>",
            "<meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width, initial-scale=1'>",
            f"<title>{html.escape(label_a)} vs {html.escape(label_b)} Extra-Line Dashboard</title>",
            "<style>",
            "body { font-family: Segoe UI, Arial, sans-serif; margin: 0; background: #f7f7f7; color: #111; }",
            "main { max-width: 1400px; margin: 0 auto; padding: 24px; }",
            "header { margin-bottom: 24px; }",
            "h1 { margin: 0 0 8px; }",
            "p { line-height: 1.5; }",
            ".chart-block { background: #fff; border: 1px solid #ddd; padding: 20px; margin: 0 auto 20px; border-radius: 10px; max-width: 1260px; }",
            ".figure-wrap { width: min(100%, 1120px); margin: 0 auto; }",
            ".figure-wrap.narrow { width: min(100%, 900px); margin: 0 auto; }",
            "code { background: #f0f0f0; padding: 2px 5px; border-radius: 4px; }",
            "</style>",
            "</head>",
            "<body>",
            "<main>",
            "<header>",
            f"<h1>{html.escape(label_a)} vs {html.escape(label_b)} Extra-Line Dashboard</h1>",
            f"<p>Dashboard for two OCR variants. It compares nonblank lines unique to each variant on the same page, without using ABBYY or human-transcribed ground truth.</p>",
            f"<p><strong>Generated:</strong> {html.escape(generated)}</p>",
            f"<p><strong>Input:</strong> <code>{html.escape(workbook_display)}</code></p>",
            f"<p><strong>Scope:</strong> {document_count} documents, {page_count} matched pages, {extra_lines_a:,} {html.escape(label_a)}-only lines, {extra_lines_b:,} {html.escape(label_b)}-only lines.</p>",
            "</header>",
            *sections,
            "</main>",
            "</body>",
            "</html>",
        ]
    )


def build_dashboard_sections(figures: list[tuple[str, str, str, go.Figure]]) -> list[str]:
    sections: list[str] = []
    include_plotlyjs = True
    for slug, title, description, figure in figures:
        section_html = to_html(
            figure,
            include_plotlyjs=include_plotlyjs,
            full_html=False,
            config={"responsive": True, "displaylogo": False},
        )
        include_plotlyjs = False
        wrap_class = "figure-wrap narrow" if slug.endswith(("delta", "rank")) else "figure-wrap"
        sections.append(
            "\n".join(
                [
                    "<section class='chart-block'>",
                    f"<h2>{html.escape(title)}</h2>",
                    f"<p>{html.escape(description)}</p>",
                    f"<div class='{wrap_class}'>",
                    section_html,
                    "</div>",
                    "</section>",
                ]
            )
        )
    return sections


def write_png_exports(figures: list[tuple[str, str, str, go.Figure]], output_dir: Path) -> None:
    for slug, _title, _description, figure in figures:
        output_path = output_dir / f"{slug}.png"
        width = 1600
        height = 900
        if slug.endswith("page_extra_line_delta"):
            height = 1250
        write_image(figure, output_path, format="png", width=width, height=height, scale=2)


def parse_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else 0


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")


if __name__ == "__main__":
    main()
