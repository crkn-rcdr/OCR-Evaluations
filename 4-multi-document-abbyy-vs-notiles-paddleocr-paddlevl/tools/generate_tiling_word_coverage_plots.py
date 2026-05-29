#!/usr/bin/env python3
"""Generate charts for the 3x2 vs 3x1 ABBYY word-coverage workbook."""

from __future__ import annotations

import argparse
import html
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from plotly import graph_objects as go
from plotly.io import to_html, write_image


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "test-results" / "tiling_3x2_vs_3x1_abbyy_word_coverage.xlsx"
DEFAULT_OUTPUT_HTML = REPO_ROOT / "test-results" / "plots" / "tiling_3x2_vs_3x1_abbyy_word_coverage_dashboard.html"
DEFAULT_OUTPUT_PNG_DIR = REPO_ROOT / "test-results" / "plots" / "png"


@dataclass(frozen=True)
class DocumentSummary:
    document: str
    page_count: int
    words_3x2: int
    words_3x1: int
    words_abbyy: int
    matched_3x2: int
    matched_3x1: int
    coverage_3x2: float
    coverage_3x1: float
    precision_3x2: float
    precision_3x1: float
    missing_3x2: int
    missing_3x1: int
    extra_3x2: int
    extra_3x1: int


@dataclass(frozen=True)
class PageSummary:
    document: str
    page: str
    coverage_3x2: float
    coverage_3x1: float
    precision_3x2: float
    precision_3x1: float
    words_3x2: int
    words_3x1: int
    words_abbyy: int

    @property
    def label(self) -> str:
        return f"{self.document} / {self.page}"

    @property
    def coverage_delta(self) -> float:
        return self.coverage_3x2 - self.coverage_3x1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-html", type=Path, default=DEFAULT_OUTPUT_HTML)
    parser.add_argument("--output-png-dir", type=Path, default=DEFAULT_OUTPUT_PNG_DIR)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    output_html = args.output_html.resolve()
    output_png_dir = args.output_png_dir.resolve()
    output_html.parent.mkdir(parents=True, exist_ok=True)
    output_png_dir.mkdir(parents=True, exist_ok=True)

    overall, documents, pages = load_workbook_data(workbook_path)
    figures = [
        (
            "tiling_3x2_vs_3x1_abbyy_word_coverage",
            "ABBYY Token Coverage by Document",
            "How much of ABBYY's token-occurrence count each tiling run captures, aggregated by document.",
            build_document_coverage_figure(documents),
        ),
        (
            "tiling_3x2_vs_3x1_abbyy_word_precision",
            "Token Precision vs ABBYY by Document",
            "What share of each tiling run's words are also found in ABBYY. Higher values mean the run stays closer to ABBYY's vocabulary instead of overshooting it.",
            build_document_precision_figure(documents),
        ),
        (
            "tiling_3x2_vs_3x1_abbyy_word_counts",
            "Document Word Counts vs ABBYY",
            "Absolute document word counts for both tiling runs compared to ABBYY.",
            build_document_word_count_figure(documents),
        ),
        (
            "tiling_3x2_vs_3x1_abbyy_coverage_delta",
            "Largest Page-Level ABBYY Coverage Deltas",
            "Pages where one tiling run captures substantially more of ABBYY's token-occurrence count than the other.",
            build_page_coverage_delta_figure(pages),
        ),
        (
            "tiling_3x2_vs_3x1_abbyy_word_count_delta",
            "Page Word Count Distance from ABBYY",
            "How far each tiling run's page word count is from ABBYY's page word count after taking absolute distance.",
            build_page_word_distance_figure(pages),
        ),
    ]

    output_html.write_text(build_dashboard_html(figures, workbook_path, overall), encoding="utf-8")
    write_png_exports(figures, output_png_dir)
    print(f"Wrote {output_html}")
    print(f"Wrote PNG charts to {output_png_dir}")


def load_workbook_data(path: Path) -> tuple[dict[str, tuple[str, str, str]], list[DocumentSummary], list[PageSummary]]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["Overall Summary"]
    overall: dict[str, tuple[str, str, str]] = {}
    doc_start = None
    page_start = None
    for row in range(2, worksheet.max_row + 1):
        label = worksheet.cell(row, 1).value
        if label == "document summary":
            doc_start = row + 1
            continue
        if label == "page summary":
            page_start = row + 1
            break
        if label:
            overall[str(label)] = (
                str(worksheet.cell(row, 2).value or ""),
                str(worksheet.cell(row, 3).value or ""),
                str(worksheet.cell(row, 4).value or ""),
            )
    if doc_start is None or page_start is None:
        raise ValueError(f"Missing summary tables in {path}")

    doc_header = [str(worksheet.cell(doc_start, c).value or "").strip() for c in range(1, worksheet.max_column + 1)]
    doc_index = {value: index for index, value in enumerate(doc_header, start=1) if value}
    documents: list[DocumentSummary] = []
    row = doc_start + 1
    while row <= worksheet.max_row:
        document = str(worksheet.cell(row, doc_index["document"]).value or "").strip()
        if not document:
            break
        documents.append(
            DocumentSummary(
                document=document,
                page_count=parse_int(worksheet.cell(row, doc_index["page count"]).value),
                words_3x2=parse_int(worksheet.cell(row, doc_index["3x2 words"]).value),
                words_3x1=parse_int(worksheet.cell(row, doc_index["3x1 words"]).value),
                words_abbyy=parse_int(worksheet.cell(row, doc_index["abbyy words"]).value),
                matched_3x2=parse_int(worksheet.cell(row, doc_index["3x2 matched ABBYY tokens"]).value),
                matched_3x1=parse_int(worksheet.cell(row, doc_index["3x1 matched ABBYY tokens"]).value),
                coverage_3x2=parse_pct(worksheet.cell(row, doc_index["3x2 coverage %"]).value),
                coverage_3x1=parse_pct(worksheet.cell(row, doc_index["3x1 coverage %"]).value),
                precision_3x2=parse_pct(worksheet.cell(row, doc_index["3x2 precision %"]).value),
                precision_3x1=parse_pct(worksheet.cell(row, doc_index["3x1 precision %"]).value),
                missing_3x2=parse_int(worksheet.cell(row, doc_index["3x2 missing ABBYY tokens"]).value),
                missing_3x1=parse_int(worksheet.cell(row, doc_index["3x1 missing ABBYY tokens"]).value),
                extra_3x2=parse_int(worksheet.cell(row, doc_index["3x2 extra tokens"]).value),
                extra_3x1=parse_int(worksheet.cell(row, doc_index["3x1 extra tokens"]).value),
            )
        )
        row += 1

    page_header = [str(worksheet.cell(page_start, c).value or "").strip() for c in range(1, worksheet.max_column + 1)]
    page_index = {value: index for index, value in enumerate(page_header, start=1) if value}
    pages: list[PageSummary] = []
    row = page_start + 1
    while row <= worksheet.max_row:
        document = str(worksheet.cell(row, page_index["document"]).value or "").strip()
        if not document:
            break
        pages.append(
            PageSummary(
                document=document,
                page=str(worksheet.cell(row, page_index["page"]).value or "").strip(),
                coverage_3x2=parse_pct(worksheet.cell(row, page_index["3x2 coverage %"]).value),
                coverage_3x1=parse_pct(worksheet.cell(row, page_index["3x1 coverage %"]).value),
                precision_3x2=parse_pct(worksheet.cell(row, page_index["3x2 precision %"]).value),
                precision_3x1=parse_pct(worksheet.cell(row, page_index["3x1 precision %"]).value),
                words_3x2=parse_int(worksheet.cell(row, page_index["3x2 words"]).value),
                words_3x1=parse_int(worksheet.cell(row, page_index["3x1 words"]).value),
                words_abbyy=parse_int(worksheet.cell(row, page_index["abbyy words"]).value),
            )
        )
        row += 1
    return overall, documents, pages


def build_document_coverage_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="3x2 coverage", x=labels, y=[row.coverage_3x2 for row in rows], marker_color="#1982C4"))
    figure.add_trace(go.Bar(name="3x1 coverage", x=labels, y=[row.coverage_3x1 for row in rows], marker_color="#55A630"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Percent of ABBYY token occurrences",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_document_precision_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="3x2 precision", x=labels, y=[row.precision_3x2 for row in rows], marker_color="#BC4749"))
    figure.add_trace(go.Bar(name="3x1 precision", x=labels, y=[row.precision_3x1 for row in rows], marker_color="#6A4C93"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Percent of run tokens also in ABBYY",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_document_word_count_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="ABBYY", x=labels, y=[row.words_abbyy for row in rows], marker_color="#6C757D"))
    figure.add_trace(go.Bar(name="3x2", x=labels, y=[row.words_3x2 for row in rows], marker_color="#1982C4"))
    figure.add_trace(go.Bar(name="3x1", x=labels, y=[row.words_3x1 for row in rows], marker_color="#55A630"))
    figure.update_layout(
        barmode="group",
        yaxis_title="Document word count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_coverage_delta_figure(rows: list[PageSummary]) -> go.Figure:
    ordered = sorted(rows, key=lambda row: (abs(row.coverage_delta), row.document, row.page), reverse=True)[:25]
    labels = [row.label for row in ordered][::-1]
    deltas = [row.coverage_delta for row in ordered][::-1]
    customdata = [[row.coverage_3x2, row.coverage_3x1] for row in ordered][::-1]
    positive = [value if value > 0 else None for value in deltas]
    negative = [value if value < 0 else None for value in deltas]
    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name="3x2 higher ABBYY coverage",
            x=positive,
            y=labels,
            orientation="h",
            marker_color="#1982C4",
            customdata=customdata,
            hovertemplate="%{y}<br>Coverage delta (3x2 - 3x1): %{x:.2f} pts<br>3x2: %{customdata[0]:.2f}%<br>3x1: %{customdata[1]:.2f}%<extra></extra>",
        )
    )
    figure.add_trace(
        go.Bar(
            name="3x1 higher ABBYY coverage",
            x=negative,
            y=labels,
            orientation="h",
            marker_color="#55A630",
            customdata=customdata,
            hovertemplate="%{y}<br>Coverage delta (3x2 - 3x1): %{x:.2f} pts<br>3x2: %{customdata[0]:.2f}%<br>3x1: %{customdata[1]:.2f}%<extra></extra>",
        )
    )
    figure.add_vline(x=0, line_dash="dash", line_color="#666666")
    figure.update_layout(
        xaxis_title="ABBYY coverage percentage-point delta",
        yaxis_title="Page",
        template="plotly_white",
        margin=dict(l=280, r=30, t=50, b=70),
        height=900,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_word_distance_figure(rows: list[PageSummary]) -> go.Figure:
    ordered = sorted(
        rows,
        key=lambda row: (abs(row.words_3x2 - row.words_abbyy) + abs(row.words_3x1 - row.words_abbyy), row.document, row.page),
        reverse=True,
    )[:25]
    labels = [row.label for row in ordered][::-1]
    dist_3x2 = [abs(row.words_3x2 - row.words_abbyy) for row in ordered][::-1]
    dist_3x1 = [abs(row.words_3x1 - row.words_abbyy) for row in ordered][::-1]
    figure = go.Figure()
    figure.add_trace(go.Bar(name="3x2 absolute distance from ABBYY", x=dist_3x2, y=labels, orientation="h", marker_color="#1982C4"))
    figure.add_trace(go.Bar(name="3x1 absolute distance from ABBYY", x=dist_3x1, y=labels, orientation="h", marker_color="#55A630"))
    figure.update_layout(
        barmode="group",
        xaxis_title="Absolute word-count distance from ABBYY",
        yaxis_title="Page",
        template="plotly_white",
        margin=dict(l=280, r=30, t=50, b=70),
        height=900,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_dashboard_html(figures, workbook_path: Path, overall) -> str:
    generated = datetime.now(timezone.utc).isoformat()
    try:
        workbook_display = str(workbook_path.relative_to(REPO_ROOT))
    except ValueError:
        workbook_display = str(workbook_path)
    page_count = overall.get("page count", ("0", "0", "0"))[0]
    coverage_3x2 = overall.get("coverage of ABBYY token occurrences", ("0%", "0%", "100%"))[0]
    coverage_3x1 = overall.get("coverage of ABBYY token occurrences", ("0%", "0%", "100%"))[1]
    sections, _ = build_dashboard_sections(figures, True)
    return "\n".join(
        [
            "<!DOCTYPE html>",
            "<html lang='en'>",
            "<head>",
            "<meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width, initial-scale=1'>",
            "<title>3x2 vs 3x1 ABBYY Word Coverage Dashboard</title>",
            "<style>",
            "body { font-family: Segoe UI, Arial, sans-serif; margin: 0; background: #f7f7f7; color: #111; }",
            "main { max-width: 1400px; margin: 0 auto; padding: 24px; }",
            "header { margin-bottom: 24px; }",
            "h1 { margin: 0 0 8px; }",
            "p { line-height: 1.5; }",
            ".chart-block { background: #fff; border: 1px solid #ddd; padding: 20px; margin: 0 auto 20px; border-radius: 10px; max-width: 1260px; }",
            ".figure-wrap { width: min(100%, 1120px); margin: 0 auto; }",
            ".figure-wrap.narrow { width: min(100%, 900px); margin: 0 auto; }",
            "code { background: #f0f0f0; padding: 2px 5px; border-radius: 4px; }",
            "</style>",
            "</head>",
            "<body>",
            "<main>",
            "<header>",
            "<h1>3x2 vs 3x1 ABBYY Word Coverage Dashboard</h1>",
            "<p>Dashboard for the paired tiling runs against ABBYY token overlap. It measures both how much of ABBYY each run captures and how much each run overshoots ABBYY.</p>",
            f"<p><strong>Generated:</strong> {html.escape(generated)}</p>",
            f"<p><strong>Inputs:</strong> <code>{html.escape(workbook_display)}</code></p>",
            f"<p><strong>Scope:</strong> {html.escape(page_count)} matched pages, 3x2 ABBYY-token coverage {html.escape(coverage_3x2)}, 3x1 ABBYY-token coverage {html.escape(coverage_3x1)}.</p>",
            "</header>",
            *sections,
            "</main>",
            "</body>",
            "</html>",
        ]
    )


def build_dashboard_sections(figures, include_plotlyjs: bool):
    sections = []
    for slug, title, description, figure in figures:
        section_html = to_html(
            figure,
            include_plotlyjs=include_plotlyjs,
            full_html=False,
            config={"responsive": True, "displaylogo": False},
        )
        include_plotlyjs = False
        wrap_class = "figure-wrap narrow" if slug.endswith(("delta", "word_count_delta")) else "figure-wrap"
        sections.append(
            "\n".join(
                [
                    "<section class='chart-block'>",
                    f"<h2>{html.escape(title)}</h2>",
                    f"<p>{html.escape(description)}</p>",
                    f"<div class='{wrap_class}'>",
                    section_html,
                    "</div>",
                    "</section>",
                ]
            )
        )
    return sections, include_plotlyjs


def write_png_exports(figures, output_dir: Path) -> None:
    for slug, _title, _description, figure in figures:
        output_path = output_dir / f"{slug}.png"
        width = 1600
        height = 900
        if slug.endswith(("delta", "word_count_delta")):
            height = 1250
        write_image(figure, output_path, format="png", width=width, height=height, scale=2)


def parse_int(value) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else 0


def parse_pct(value) -> float:
    text = str(value or "").strip().replace("%", "")
    return float(text) if text else 0.0


if __name__ == "__main__":
    main()
