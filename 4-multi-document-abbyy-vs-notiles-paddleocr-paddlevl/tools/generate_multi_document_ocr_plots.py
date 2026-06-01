#!/usr/bin/env python3
"""Generate Plotly dashboards for the multi-document ABBYY vs Paddle workbook.

Default usage:

  python tools/generate_multi_document_ocr_plots.py

This reads:

- test-results/multi_document_paddle_vs_abbyy_errors_artifacts.xlsx

and writes:

- test-results/plots/multi_document_ocr_dashboard.html

It also writes PNG files for reuse under:

- test-results/plots/png/
"""

from __future__ import annotations

import argparse
import html
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from plotly import graph_objects as go
from plotly.io import to_html, write_image
from plotly.subplots import make_subplots


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "test-results" / "multi_document_paddle_vs_abbyy_errors_artifacts.xlsx"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "test-results" / "plots"
DEFAULT_OUTPUT_HTML = DEFAULT_OUTPUT_DIR / "multi_document_ocr_dashboard.html"
DEFAULT_OUTPUT_PNG_DIR = DEFAULT_OUTPUT_DIR / "png"
PAGE_NUMBER_RE = re.compile(r"(\d+)(?=\.txt$)", re.IGNORECASE)


@dataclass(frozen=True)
class PageSummary:
    document: str
    page: str
    paddle_source: Path
    abbyy_source: Path
    paddle_entries: int
    abbyy_entries: int
    paddle_error_counts: dict[str, int]
    abbyy_error_counts: dict[str, int]
    paddle_line_count: int
    abbyy_line_count: int
    paddle_word_count: int
    abbyy_word_count: int

    @property
    def label(self) -> str:
        return f"{self.document} / {self.page}"


@dataclass(frozen=True)
class DocumentSummary:
    document: str
    page_count: int
    paddle_entries: int
    abbyy_entries: int
    paddle_error_counts: dict[str, int]
    abbyy_error_counts: dict[str, int]
    paddle_line_count: int
    abbyy_line_count: int
    paddle_word_count: int
    abbyy_word_count: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate Plotly visualizations for the multi-document OCR workbook."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Path to multi_document_paddle_vs_abbyy_errors_artifacts.xlsx",
    )
    parser.add_argument(
        "--output-html",
        type=Path,
        default=DEFAULT_OUTPUT_HTML,
        help="Output HTML dashboard path.",
    )
    parser.add_argument(
        "--output-png-dir",
        type=Path,
        default=DEFAULT_OUTPUT_PNG_DIR,
        help="Output directory for PNG chart exports.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    output_html = args.output_html.resolve()
    output_png_dir = args.output_png_dir.resolve()
    output_html.parent.mkdir(parents=True, exist_ok=True)
    output_png_dir.mkdir(parents=True, exist_ok=True)

    error_types, page_rows, overall_totals = load_page_summaries(workbook_path)
    document_rows = build_document_summaries(page_rows, error_types)

    figures = [
        (
            "document_line_recovery",
            "Document Line Recovery",
            (
                "How many lines the Paddle workflow recovered relative to ABBYY, "
                "aggregated by document. Values above 100% mean Paddle produced "
                "more line segments than ABBYY."
            ),
            build_line_recovery_figure(document_rows),
        ),
        (
            "document_word_recovery",
            "Document Word Recovery",
            (
                "How many words the Paddle workflow recovered relative to ABBYY, "
                "aggregated by document. Values above 100% mean Paddle produced "
                "more OCR tokens than ABBYY."
            ),
            build_word_recovery_figure(document_rows),
        ),
        (
            "document_word_counts",
            "Document Word Counts",
            (
                "Absolute whole-document word counts for ABBYY and the "
                "Paddle workflow."
            ),
            build_absolute_word_count_figure(document_rows),
        ),
        (
            "document_artifact_totals",
            "Document Artifact Totals",
            (
                "Total artifact-entry counts by document. Lower is cleaner, but "
                "the counts should be read together with the recovery charts."
            ),
            build_artifact_totals_figure(document_rows),
        ),
        (
            "document_artifact_categories",
            "Artifact Categories vs ABBYY",
            (
                "Each subplot shows one document. Category values are expressed "
                "as a percent of ABBYY's count for the same category."
            ),
            build_artifact_category_vs_abbyy_figure(document_rows, error_types),
        ),
        (
            "page_artifact_rank",
            "Page Artifact Totals",
            (
                "All pages sorted from highest to lowest combined artifact count. "
                "This makes the hardest pages and the Paddle-vs-ABBYY spread easy "
                "to spot without a human reference set."
            ),
            build_page_artifact_rank_figure(page_rows),
        ),
        (
            "page_artifact_delta",
            "Largest Page-Level Artifact Deltas",
            (
                "Pages with the biggest gap between Paddle and ABBYY artifact "
                "totals. Positive bars mean Paddle had more tagged artifacts; "
                "negative bars mean ABBYY had more."
            ),
            build_page_delta_figure(page_rows),
        ),
    ]

    html_text = build_dashboard_html(figures, workbook_path, overall_totals)
    output_html.write_text(html_text, encoding="utf-8")
    write_png_exports(figures, output_png_dir)
    print(f"Wrote {output_html}")
    print(f"Wrote PNG charts to {output_png_dir}")


def load_page_summaries(path: Path) -> tuple[list[str], list[PageSummary], dict[str, tuple[int, int]]]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["Overall Summary"]
    overall_totals: dict[str, tuple[int, int]] = {}

    page_summary_row = None
    for row in range(2, worksheet.max_row + 1):
        label = str(worksheet.cell(row, 1).value or "").strip()
        if not label:
            continue
        if label == "page summary":
            page_summary_row = row
            break
        if label.endswith(" issue count"):
            overall_totals[label[: -len(" issue count")]] = (
                parse_int(worksheet.cell(row, 2).value),
                parse_int(worksheet.cell(row, 3).value),
            )
        elif label in {
            "document count",
            "page count",
            "line count",
            "word count",
            "error/artifact entries",
        }:
            overall_totals[label] = (
                parse_int(worksheet.cell(row, 2).value),
                parse_int(worksheet.cell(row, 3).value),
            )

    if page_summary_row is None:
        raise ValueError(f"Could not find page summary table in {path}")

    header_row = page_summary_row + 1
    header = [
        str(worksheet.cell(header_row, column).value or "").strip()
        for column in range(1, worksheet.max_column + 1)
    ]
    header_map = {value: index for index, value in enumerate(header, start=1) if value}
    error_types = [
        value.replace("paddle ", "", 1)
        for value in header
        if value.startswith("paddle ")
        and value not in {"paddle source file", "paddle entries"}
    ]

    text_stats_cache: dict[Path, tuple[int, int]] = {}
    page_rows: list[PageSummary] = []
    for row in range(header_row + 1, worksheet.max_row + 1):
        document = str(worksheet.cell(row, header_map["document"]).value or "").strip()
        if not document:
            break
        page = str(worksheet.cell(row, header_map["page"]).value or "").strip()
        paddle_source = Path(str(worksheet.cell(row, header_map["paddle source file"]).value or "").strip())
        abbyy_source = Path(str(worksheet.cell(row, header_map["abbyy source file"]).value or "").strip())
        paddle_lines, paddle_words = read_text_stats(paddle_source, text_stats_cache)
        abbyy_lines, abbyy_words = read_text_stats(abbyy_source, text_stats_cache)
        paddle_error_counts = {
            error_type: parse_int(worksheet.cell(row, header_map[f"paddle {error_type}"]).value)
            for error_type in error_types
        }
        abbyy_error_counts = {
            error_type: parse_int(worksheet.cell(row, header_map[f"abbyy {error_type}"]).value)
            for error_type in error_types
        }
        page_rows.append(
            PageSummary(
                document=document,
                page=page,
                paddle_source=paddle_source,
                abbyy_source=abbyy_source,
                paddle_entries=parse_int(worksheet.cell(row, header_map["paddle entries"]).value),
                abbyy_entries=parse_int(worksheet.cell(row, header_map["abbyy entries"]).value),
                paddle_error_counts=paddle_error_counts,
                abbyy_error_counts=abbyy_error_counts,
                paddle_line_count=paddle_lines,
                abbyy_line_count=abbyy_lines,
                paddle_word_count=paddle_words,
                abbyy_word_count=abbyy_words,
            )
        )

    page_rows.sort(key=page_sort_key)
    return error_types, page_rows, overall_totals


def build_document_summaries(
    page_rows: list[PageSummary],
    error_types: list[str],
) -> list[DocumentSummary]:
    buckets: dict[str, dict[str, object]] = defaultdict(
        lambda: {
            "page_count": 0,
            "paddle_entries": 0,
            "abbyy_entries": 0,
            "paddle_line_count": 0,
            "abbyy_line_count": 0,
            "paddle_word_count": 0,
            "abbyy_word_count": 0,
            "paddle_error_counts": Counter(),
            "abbyy_error_counts": Counter(),
        }
    )

    for row in page_rows:
        bucket = buckets[row.document]
        bucket["page_count"] = int(bucket["page_count"]) + 1
        bucket["paddle_entries"] = int(bucket["paddle_entries"]) + row.paddle_entries
        bucket["abbyy_entries"] = int(bucket["abbyy_entries"]) + row.abbyy_entries
        bucket["paddle_line_count"] = int(bucket["paddle_line_count"]) + row.paddle_line_count
        bucket["abbyy_line_count"] = int(bucket["abbyy_line_count"]) + row.abbyy_line_count
        bucket["paddle_word_count"] = int(bucket["paddle_word_count"]) + row.paddle_word_count
        bucket["abbyy_word_count"] = int(bucket["abbyy_word_count"]) + row.abbyy_word_count
        bucket["paddle_error_counts"].update(row.paddle_error_counts)
        bucket["abbyy_error_counts"].update(row.abbyy_error_counts)

    rows: list[DocumentSummary] = []
    for document in sorted(buckets):
        bucket = buckets[document]
        rows.append(
            DocumentSummary(
                document=document,
                page_count=int(bucket["page_count"]),
                paddle_entries=int(bucket["paddle_entries"]),
                abbyy_entries=int(bucket["abbyy_entries"]),
                paddle_error_counts={
                    error_type: int(bucket["paddle_error_counts"].get(error_type, 0))
                    for error_type in error_types
                },
                abbyy_error_counts={
                    error_type: int(bucket["abbyy_error_counts"].get(error_type, 0))
                    for error_type in error_types
                },
                paddle_line_count=int(bucket["paddle_line_count"]),
                abbyy_line_count=int(bucket["abbyy_line_count"]),
                paddle_word_count=int(bucket["paddle_word_count"]),
                abbyy_word_count=int(bucket["abbyy_word_count"]),
            )
        )
    return rows


def build_line_recovery_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    values = [percentage(row.paddle_line_count, row.abbyy_line_count) for row in rows]
    customdata = [
        [row.paddle_line_count, row.abbyy_line_count, row.page_count]
        for row in rows
    ]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name="Paddle line count vs ABBYY",
            x=labels,
            y=values,
            text=[f"{value:.1f}%" for value in values],
            textposition="outside",
            customdata=customdata,
            marker_color="#33658A",
            hovertemplate=(
                "%{x}<br>Paddle lines: %{customdata[0]:,}"
                "<br>ABBYY lines: %{customdata[1]:,}"
                "<br>Pages: %{customdata[2]}"
                "<br>Recovery: %{y:.1f}%<extra></extra>"
            ),
        )
    )
    figure.add_hline(line_dash="dash", line_color="#666666", y=100)
    figure.update_layout(
        yaxis_title="Percent of ABBYY line count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
    )
    figure.update_yaxes(range=[0, max(max(values) * 1.12, 108.0)])
    return figure


def build_word_recovery_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    values = [percentage(row.paddle_word_count, row.abbyy_word_count) for row in rows]
    customdata = [
        [row.paddle_word_count, row.abbyy_word_count, row.page_count]
        for row in rows
    ]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name="Paddle word count vs ABBYY",
            x=labels,
            y=values,
            text=[f"{value:.1f}%" for value in values],
            textposition="outside",
            customdata=customdata,
            marker_color="#55A630",
            hovertemplate=(
                "%{x}<br>Paddle words: %{customdata[0]:,}"
                "<br>ABBYY words: %{customdata[1]:,}"
                "<br>Pages: %{customdata[2]}"
                "<br>Recovery: %{y:.1f}%<extra></extra>"
            ),
        )
    )
    figure.add_hline(line_dash="dash", line_color="#666666", y=100)
    figure.update_layout(
        yaxis_title="Percent of ABBYY word count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
    )
    figure.update_yaxes(range=[0, max(max(values) * 1.12, 108.0)])
    return figure


def build_absolute_word_count_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    paddle_values = [row.paddle_word_count for row in rows]
    abbyy_values = [row.abbyy_word_count for row in rows]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name="ABBYY",
            x=labels,
            y=abbyy_values,
            text=[f"{value:,}" for value in abbyy_values],
            textposition="outside",
            marker_color="#6C757D",
        )
    )
    figure.add_trace(
        go.Bar(
            name="Paddle",
            x=labels,
            y=paddle_values,
            text=[f"{value:,}" for value in paddle_values],
            textposition="outside",
            marker_color="#1982C4",
        )
    )
    figure.update_layout(
        barmode="group",
        yaxis_title="Whole-document word count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    figure.update_yaxes(range=[0, max(max(paddle_values), max(abbyy_values)) * 1.18])
    return figure


def build_artifact_totals_figure(rows: list[DocumentSummary]) -> go.Figure:
    labels = [row.document for row in rows]
    paddle_values = [row.paddle_entries for row in rows]
    abbyy_values = [row.abbyy_entries for row in rows]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name="Paddle",
            x=labels,
            y=paddle_values,
            text=[f"{value:,}" for value in paddle_values],
            textposition="outside",
            marker_color="#BC4749",
        )
    )
    figure.add_trace(
        go.Scatter(
            name="ABBYY",
            x=labels,
            y=abbyy_values,
            mode="lines+markers+text",
            text=[f"{value:,}" for value in abbyy_values],
            textposition="top center",
            line=dict(color="#6C757D", width=3, dash="dash"),
            marker=dict(color="#6C757D", size=9),
        )
    )
    figure.update_layout(
        yaxis_title="Artifact-entry count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=110),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_artifact_category_vs_abbyy_figure(
    rows: list[DocumentSummary],
    error_types: list[str],
) -> go.Figure:
    figure = make_subplots(
        rows=len(rows),
        cols=1,
        subplot_titles=[row.document for row in rows],
        vertical_spacing=0.03,
    )

    colors = ["#33658A", "#55A630", "#BC4749", "#FF9F1C", "#6A4C93", "#1982C4", "#2F4858"]
    max_ratio = 0.0
    for index, row in enumerate(rows, start=1):
        ratios = []
        customdata = []
        for error_type in error_types:
            paddle_count = row.paddle_error_counts[error_type]
            abbyy_count = row.abbyy_error_counts[error_type]
            if abbyy_count > 0:
                ratio = percentage(paddle_count, abbyy_count)
            else:
                ratio = 100.0 if paddle_count == 0 else paddle_count * 100.0
            ratios.append(ratio)
            customdata.append([paddle_count, abbyy_count])
        max_ratio = max(max_ratio, max(ratios) if ratios else 0.0)
        figure.add_trace(
            go.Bar(
                x=ratios,
                y=error_types,
                orientation="h",
                marker_color=colors[(index - 1) % len(colors)],
                customdata=customdata,
                hovertemplate=(
                    "%{y}<br>Paddle: %{customdata[0]}"
                    "<br>ABBYY: %{customdata[1]}"
                    "<br>Paddle as % of ABBYY: %{x:.1f}%<extra></extra>"
                ),
                showlegend=index == 1,
                name="Paddle as % of ABBYY",
            ),
            row=index,
            col=1,
        )

    figure.add_trace(
        go.Scatter(
            x=[None],
            y=[None],
            mode="lines",
            line=dict(color="#666666", width=2, dash="dash"),
            name="ABBYY baseline = 100%",
            showlegend=True,
            hoverinfo="skip",
        ),
        row=1,
        col=1,
    )

    max_range = max(140.0, min(max_ratio * 1.12, 420.0))
    for row_index in range(1, len(rows) + 1):
        figure.add_vline(x=100, line_dash="dash", line_color="#666666", row=row_index, col=1)
        figure.update_xaxes(
            title_text="Percent of ABBYY",
            range=[0, max_range],
            row=row_index,
            col=1,
        )
        figure.update_yaxes(
            autorange="reversed",
            automargin=True,
            ticklabelposition="outside",
            tickmode="array",
            tickvals=error_types,
            ticktext=error_types,
            row=row_index,
            col=1,
        )

    figure.update_layout(
        template="plotly_white",
        margin=dict(l=260, r=30, t=70, b=80),
        height=max(2200, 310 * len(rows)),
        legend=dict(orientation="h", yanchor="bottom", y=1.01, xanchor="left", x=0),
    )
    return figure


def build_page_artifact_rank_figure(rows: list[PageSummary]) -> go.Figure:
    ordered = sorted(
        rows,
        key=lambda row: (row.paddle_entries + row.abbyy_entries, row.document, row.page),
        reverse=True,
    )
    ranks = list(range(1, len(ordered) + 1))
    paddle_values = [row.paddle_entries for row in ordered]
    abbyy_values = [row.abbyy_entries for row in ordered]
    customdata = [[row.label] for row in ordered]

    figure = go.Figure()
    figure.add_trace(
        go.Scatter(
            name="Paddle",
            x=ranks,
            y=paddle_values,
            mode="lines+markers",
            line=dict(color="#BC4749", width=2),
            marker=dict(size=6),
            customdata=customdata,
            hovertemplate="%{customdata[0]}<br>Paddle artifacts: %{y}<extra></extra>",
        )
    )
    figure.add_trace(
        go.Scatter(
            name="ABBYY",
            x=ranks,
            y=abbyy_values,
            mode="lines+markers",
            line=dict(color="#6C757D", width=2),
            marker=dict(size=6),
            customdata=customdata,
            hovertemplate="%{customdata[0]}<br>ABBYY artifacts: %{y}<extra></extra>",
        )
    )
    figure.update_layout(
        xaxis_title="Page rank by combined artifact count",
        yaxis_title="Artifact-entry count",
        template="plotly_white",
        margin=dict(l=60, r=30, t=50, b=70),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_page_delta_figure(rows: list[PageSummary]) -> go.Figure:
    ordered = sorted(
        rows,
        key=lambda row: (abs(row.paddle_entries - row.abbyy_entries), row.document, row.page),
        reverse=True,
    )[:25]
    labels = [row.label for row in ordered][::-1]
    delta_values = [(row.paddle_entries - row.abbyy_entries) for row in ordered][::-1]
    customdata = [[row.paddle_entries, row.abbyy_entries] for row in ordered][::-1]
    positive_values = [value if value > 0 else None for value in delta_values]
    negative_values = [value if value < 0 else None for value in delta_values]

    figure = go.Figure()
    figure.add_trace(
        go.Bar(
            name="Paddle higher artifact count",
            x=positive_values,
            y=labels,
            orientation="h",
            marker_color="#BC4749",
            customdata=customdata,
            hovertemplate=(
                "%{y}<br>Delta (Paddle - ABBYY): %{x}"
                "<br>Paddle: %{customdata[0]}"
                "<br>ABBYY: %{customdata[1]}<extra></extra>"
            ),
        )
    )
    figure.add_trace(
        go.Bar(
            name="ABBYY higher artifact count",
            x=negative_values,
            y=labels,
            orientation="h",
            marker_color="#6C757D",
            customdata=customdata,
            hovertemplate=(
                "%{y}<br>Delta (Paddle - ABBYY): %{x}"
                "<br>Paddle: %{customdata[0]}"
                "<br>ABBYY: %{customdata[1]}<extra></extra>"
            ),
        )
    )
    figure.add_trace(
        go.Scatter(
            x=[None],
            y=[None],
            mode="lines",
            line=dict(color="#666666", width=2, dash="dash"),
            name="Zero baseline",
            showlegend=True,
            hoverinfo="skip",
        )
    )
    figure.add_vline(x=0, line_dash="dash", line_color="#666666")
    figure.update_layout(
        xaxis_title="Artifact-entry delta",
        yaxis_title="Page",
        template="plotly_white",
        margin=dict(l=280, r=30, t=50, b=70),
        height=900,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    return figure


def build_dashboard_html(
    figures: list[tuple[str, str, str, go.Figure]],
    workbook_path: Path,
    overall_totals: dict[str, tuple[int, int]],
) -> str:
    generated = datetime.now(timezone.utc).isoformat()
    workbook_display = format_repo_relative(workbook_path)
    page_count = overall_totals.get("page count", (0, 0))[0]
    document_count = overall_totals.get("document count", (0, 0))[0]
    paddle_entries = overall_totals.get("error/artifact entries", (0, 0))[0]
    abbyy_entries = overall_totals.get("error/artifact entries", (0, 0))[1]
    sections, _ = build_dashboard_sections(figures, include_plotlyjs=True)

    return "\n".join(
        [
            "<!DOCTYPE html>",
            "<html lang='en'>",
            "<head>",
            "<meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width, initial-scale=1'>",
            "<title>Multi-Document OCR Dashboard</title>",
            "<style>",
            "body { font-family: Segoe UI, Arial, sans-serif; margin: 0; background: #f7f7f7; color: #111; }",
            "main { max-width: 1400px; margin: 0 auto; padding: 24px; }",
            "header { margin-bottom: 24px; }",
            "h1 { margin: 0 0 8px; }",
            "p { line-height: 1.5; }",
            ".chart-block { background: #fff; border: 1px solid #ddd; padding: 20px; margin: 0 auto 20px; border-radius: 10px; max-width: 1260px; }",
            ".figure-wrap { width: min(100%, 1120px); margin: 0 auto; }",
            ".figure-wrap.narrow { width: min(100%, 900px); margin: 0 auto; }",
            "code { background: #f0f0f0; padding: 2px 5px; border-radius: 4px; }",
            "</style>",
            "</head>",
            "<body>",
            "<main>",
            "<header>",
            "<h1>Multi-Document OCR Dashboard</h1>",
            (
                "<p>Dashboard for the multi-document ABBYY vs Paddle comparison. "
                "No human transcription workbook is available yet, so these charts focus on "
                "coverage and artifact patterns between the two OCR outputs.</p>"
            ),
            f"<p><strong>Generated:</strong> {html.escape(generated)}</p>",
            f"<p><strong>Inputs:</strong> <code>{html.escape(workbook_display)}</code></p>",
            (
                f"<p><strong>Scope:</strong> {document_count} documents, {page_count} matched pages, "
                f"{paddle_entries:,} Paddle artifact entries, {abbyy_entries:,} ABBYY artifact entries.</p>"
            ),
            "</header>",
            *sections,
            "</main>",
            "</body>",
            "</html>",
        ]
    )


def build_dashboard_sections(
    figures: list[tuple[str, str, str, go.Figure]],
    include_plotlyjs: bool,
) -> tuple[list[str], bool]:
    sections: list[str] = []
    for slug, title, description, figure in figures:
        section_html = to_html(
            figure,
            include_plotlyjs=include_plotlyjs,
            full_html=False,
            config={"responsive": True, "displaylogo": False},
        )
        include_plotlyjs = False
        wrap_class = "figure-wrap narrow" if slug in {"document_artifact_categories", "page_artifact_delta"} else "figure-wrap"
        sections.append(
            "\n".join(
                [
                    "<section class='chart-block'>",
                    f"<h2>{html.escape(title)}</h2>",
                    f"<p>{html.escape(description)}</p>",
                    f"<div class='{wrap_class}'>",
                    section_html,
                    "</div>",
                    "</section>",
                ]
            )
        )
    return sections, include_plotlyjs


def write_png_exports(
    figures: list[tuple[str, str, str, go.Figure]],
    output_dir: Path,
) -> None:
    for slug, _title, _description, figure in figures:
        output_path = output_dir / f"{slug}.png"
        width = 1600
        height = 900
        if slug == "document_artifact_categories":
            width = 1300
            height = 2600
        elif slug == "page_artifact_delta":
            width = 1600
            height = 1250
        write_image(figure, output_path, format="png", width=width, height=height, scale=2)


def read_text_stats(path: Path, cache: dict[Path, tuple[int, int]]) -> tuple[int, int]:
    resolved = path.resolve()
    if resolved in cache:
        return cache[resolved]
    text = resolved.read_text(encoding="utf-8", errors="replace")
    stats = (sum(1 for line in text.splitlines() if line.strip()), len(text.split()))
    cache[resolved] = stats
    return stats


def page_sort_key(row: PageSummary) -> tuple[str, int, str]:
    return (
        row.document.casefold(),
        extract_page_number(row.page),
        row.page.casefold(),
    )


def extract_page_number(value: str) -> int:
    match = PAGE_NUMBER_RE.search(value)
    if match:
        return int(match.group(1))
    return 10**9


def format_repo_relative(path: Path) -> str:
    try:
        return str(path.relative_to(REPO_ROOT))
    except ValueError:
        return str(path)


def parse_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else 0


def percentage(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return (numerator / denominator) * 100.0


if __name__ == "__main__":
    main()
