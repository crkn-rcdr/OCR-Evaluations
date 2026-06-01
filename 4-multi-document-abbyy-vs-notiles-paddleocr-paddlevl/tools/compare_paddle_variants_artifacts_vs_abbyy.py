#!/usr/bin/env python3
"""Merge one or more ABBYY artifact workbooks into a combined comparison workbook."""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment


REPO_ROOT = Path(__file__).resolve().parent.parent
EVAL_ROOT = REPO_ROOT.parent
RUN_WORKBOOKS = [
    (
        "3x2",
        REPO_ROOT / "test-results" / "3x2_vs_abbyy_errors_artifacts.xlsx",
    ),
    (
        "3x1",
        REPO_ROOT / "test-results" / "3x1_vs_abbyy_errors_artifacts.xlsx",
    ),
    (
        "no tiling",
        REPO_ROOT / "test-results" / "multi_document_paddle_vs_abbyy_errors_artifacts.xlsx",
    ),
]
DEFAULT_OUTPUT = REPO_ROOT / "test-results" / "paddle_variants_abbyy_artifacts.xlsx"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--run-workbook",
        dest="run_workbooks",
        action="append",
        metavar="LABEL=PATH",
        help=(
            "Custom labeled workbook input. Repeat for each run, for example "
            "--run-workbook \"3x1 PaddleVL=C:\\path\\to\\workbook.xlsx\". "
            "If omitted, the default 3x2 / 3x1 / no-tiling set is used."
        ),
    )
    parser.add_argument("--workbook-3x2", type=Path, default=RUN_WORKBOOKS[0][1])
    parser.add_argument("--workbook-3x1", type=Path, default=RUN_WORKBOOKS[1][1])
    parser.add_argument("--workbook-notiling", type=Path, default=RUN_WORKBOOKS[2][1])
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def parse_run_workbook_specs(values: list[str] | None) -> list[tuple[str, Path]]:
    if not values:
        return []
    run_inputs: list[tuple[str, Path]] = []
    for value in values:
        if "=" not in value:
            raise ValueError(f"Invalid --run-workbook value '{value}'. Expected LABEL=PATH.")
        label, raw_path = value.split("=", 1)
        label = re.sub(r"\s+", " ", label).strip()
        if not label:
            raise ValueError(f"Invalid --run-workbook value '{value}'. Label must not be empty.")
        path = Path(raw_path.strip())
        if not str(path):
            raise ValueError(f"Invalid --run-workbook value '{value}'. Path must not be empty.")
        run_inputs.append((label, path))
    return run_inputs


def main() -> None:
    args = parse_args()
    custom_inputs = parse_run_workbook_specs(args.run_workbooks)
    if custom_inputs:
        run_inputs = [(label, path.resolve()) for label, path in custom_inputs]
    else:
        run_inputs = [
            ("3x2", args.workbook_3x2.resolve()),
            ("3x1", args.workbook_3x1.resolve()),
            ("no tiling", args.workbook_notiling.resolve()),
        ]
    run_pages = {label: load_run_pages(path) for label, path in run_inputs}
    combined = combine_pages(run_pages)
    write_workbook(args.output.resolve(), build_overall_rows(combined, run_inputs))
    print(f"Wrote {args.output.resolve()}")
    print(f"Compared pages: {len(combined)}")


def parse_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else 0


def parse_pct(value: object) -> float:
    if value is None:
        return 0.0
    text = str(value).strip().replace("%", "").replace(",", "")
    return float(text) if text else 0.0


def load_run_pages(path: Path) -> dict[tuple[str, str], dict[str, object]]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["Overall Summary"]
    page_summary_row = None
    for row in range(2, worksheet.max_row + 1):
        label = str(worksheet.cell(row, 1).value or "").strip()
        if label == "page summary":
            page_summary_row = row + 1
            break
    if page_summary_row is None:
        raise ValueError(f"Could not find page summary in {path}")

    header = [str(worksheet.cell(page_summary_row, c).value or "").strip() for c in range(1, worksheet.max_column + 1)]
    index = {value: idx for idx, value in enumerate(header, start=1) if value}
    required = [
        "document",
        "page",
        "paddle source file",
        "abbyy source file",
        "paddle entries",
        "abbyy entries",
    ]
    for key in required:
        if key not in index:
            raise KeyError(f"Missing expected column '{key}' in {path}")

    pages: dict[tuple[str, str], dict[str, object]] = {}
    row = page_summary_row + 1
    while row <= worksheet.max_row:
        document = str(worksheet.cell(row, index["document"]).value or "").strip()
        if not document:
            break
        page = str(worksheet.cell(row, index["page"]).value or "").strip()
        paddle_source = Path(str(worksheet.cell(row, index["paddle source file"]).value or ""))
        abbyy_source = Path(str(worksheet.cell(row, index["abbyy source file"]).value or ""))
        paddle_line_count, paddle_word_count = read_text_metrics(paddle_source)
        abbyy_line_count, abbyy_word_count = read_text_metrics(abbyy_source)
        pages[(document, page)] = {
            "paddle_source_file": str(paddle_source),
            "abbyy_source_file": str(abbyy_source),
            "paddle_entry_count": parse_int(worksheet.cell(row, index["paddle entries"]).value),
            "abbyy_entry_count": parse_int(worksheet.cell(row, index["abbyy entries"]).value),
            "paddle_line_count": paddle_line_count,
            "abbyy_line_count": abbyy_line_count,
            "paddle_word_count": paddle_word_count,
            "abbyy_word_count": abbyy_word_count,
        }
        row += 1
    return pages


WORD_RE = re.compile(r"\S+")


def read_text_metrics(path: Path) -> tuple[int, int]:
    text = path.read_text(encoding="utf-8", errors="replace")
    line_count = sum(1 for line in text.splitlines() if line.strip())
    word_count = len(WORD_RE.findall(text))
    return line_count, word_count


def combine_pages(run_pages: dict[str, dict[tuple[str, str], dict[str, object]]]) -> list[dict[str, object]]:
    labels = list(run_pages)
    canonical_abbyy_label = "no tiling" if "no tiling" in labels else labels[0]
    page_keys = set(run_pages[labels[0]].keys())
    for label in labels[1:]:
        if set(run_pages[label].keys()) != page_keys:
            missing = sorted(page_keys ^ set(run_pages[label].keys()))[:10]
            raise FileNotFoundError(f"Page mismatch between runs; sample keys: {missing}")

    combined: list[dict[str, object]] = []
    for key in sorted(page_keys, key=lambda item: (item[0].casefold(), item[1].casefold())):
        document, page = key
        first = run_pages[canonical_abbyy_label][key]
        page_row: dict[str, object] = {
            "document": document,
            "page": page,
            "abbyy_source_file": first["abbyy_source_file"],
            "abbyy_entry_count": first["abbyy_entry_count"],
            "abbyy_line_count": first["abbyy_line_count"],
            "abbyy_word_count": first["abbyy_word_count"],
        }
        for label in labels:
            if label == canonical_abbyy_label:
                continue
            current = run_pages[label][key]
            for field in ("abbyy_line_count", "abbyy_word_count"):
                if current[field] != page_row[field]:
                    raise ValueError(f"ABBYY mismatch for {document}/{page} between runs")
        for label in labels:
            current = run_pages[label][key]
            page_row[f"{label}_source_file"] = current["paddle_source_file"]
            page_row[f"{label}_entry_count"] = current["paddle_entry_count"]
            page_row[f"{label}_line_count"] = current["paddle_line_count"]
            page_row[f"{label}_word_count"] = current["paddle_word_count"]
            page_row[f"{label}_entry_delta_vs_abbyy"] = current["paddle_entry_count"] - page_row["abbyy_entry_count"]
            page_row[f"{label}_line_recovery_pct"] = (
                100.0 * current["paddle_line_count"] / page_row["abbyy_line_count"]
                if int(page_row["abbyy_line_count"]) > 0
                else 0.0
            )
            page_row[f"{label}_word_recovery_pct"] = (
                100.0 * current["paddle_word_count"] / page_row["abbyy_word_count"]
                if int(page_row["abbyy_word_count"]) > 0
                else 0.0
            )
        combined.append(page_row)
    return combined


def build_overall_rows(
    pages: list[dict[str, object]],
    run_inputs: list[tuple[str, Path]],
) -> list[list[object]]:
    labels = [label for label, _path in run_inputs]
    canonical_abbyy_label = "no tiling" if "no tiling" in labels else labels[0]
    doc_totals: dict[str, dict[str, object]] = defaultdict(lambda: {"page_count": 0, "abbyy_entry_count": 0, "abbyy_line_count": 0, "abbyy_word_count": 0})

    totals = {
        "abbyy_entry_count": 0,
        "abbyy_line_count": 0,
        "abbyy_word_count": 0,
    }
    for label in labels:
        totals[f"{label}_entry_count"] = 0
        totals[f"{label}_line_count"] = 0
        totals[f"{label}_word_count"] = 0

    pages_lower_than_abbyy = {label: 0 for label in labels}
    pages_tied_abbyy = {label: 0 for label in labels}
    best_run_counts = {label: 0 for label in labels}
    best_run_ties = 0

    for page in pages:
        document = str(page["document"])
        bucket = doc_totals[document]
        bucket["page_count"] = int(bucket["page_count"]) + 1
        bucket["abbyy_entry_count"] = int(bucket["abbyy_entry_count"]) + int(page["abbyy_entry_count"])
        bucket["abbyy_line_count"] = int(bucket["abbyy_line_count"]) + int(page["abbyy_line_count"])
        bucket["abbyy_word_count"] = int(bucket["abbyy_word_count"]) + int(page["abbyy_word_count"])
        totals["abbyy_entry_count"] += int(page["abbyy_entry_count"])
        totals["abbyy_line_count"] += int(page["abbyy_line_count"])
        totals["abbyy_word_count"] += int(page["abbyy_word_count"])

        run_entries = {}
        for label in labels:
            bucket.setdefault(f"{label}_entry_count", 0)
            bucket.setdefault(f"{label}_line_count", 0)
            bucket.setdefault(f"{label}_word_count", 0)
            bucket[f"{label}_entry_count"] = int(bucket[f"{label}_entry_count"]) + int(page[f"{label}_entry_count"])
            bucket[f"{label}_line_count"] = int(bucket[f"{label}_line_count"]) + int(page[f"{label}_line_count"])
            bucket[f"{label}_word_count"] = int(bucket[f"{label}_word_count"]) + int(page[f"{label}_word_count"])
            totals[f"{label}_entry_count"] += int(page[f"{label}_entry_count"])
            totals[f"{label}_line_count"] += int(page[f"{label}_line_count"])
            totals[f"{label}_word_count"] += int(page[f"{label}_word_count"])
            run_entries[label] = int(page[f"{label}_entry_count"])
            if run_entries[label] < int(page["abbyy_entry_count"]):
                pages_lower_than_abbyy[label] += 1
            elif run_entries[label] == int(page["abbyy_entry_count"]):
                pages_tied_abbyy[label] += 1

        lowest = min(run_entries.values())
        winners = [label for label, value in run_entries.items() if value == lowest]
        if len(winners) == 1:
            best_run_counts[winners[0]] += 1
        else:
            best_run_ties += 1

    rows: list[list[object]] = [
        ["summary metric", *labels, "abbyy baseline"],
        ["workbook source", *[str(path) for _label, path in run_inputs], ""],
        ["baseline note", *["" for _label in labels], f"ABBYY artifact counts in this merged workbook use the {canonical_abbyy_label} pairing as the canonical baseline because ABBYY mismatch tagging is pair-dependent."],
        ["document count", *([len(doc_totals)] * len(labels)), len(doc_totals)],
        ["page count", *([len(pages)] * len(labels)), len(pages)],
        ["error/artifact entries", *[totals[f"{label}_entry_count"] for label in labels], totals["abbyy_entry_count"]],
        ["line count", *[totals[f'{label}_line_count'] for label in labels], totals["abbyy_line_count"]],
        ["word count", *[totals[f'{label}_word_count'] for label in labels], totals["abbyy_word_count"]],
        ["pages lower than ABBYY artifact count", *[pages_lower_than_abbyy[label] for label in labels], ""],
        ["pages tied with ABBYY artifact count", *[pages_tied_abbyy[label] for label in labels], ""],
        ["pages with lowest artifact count among Paddle variants", *[best_run_counts[label] for label in labels], best_run_ties],
    ]

    rows.append([])
    rows.append(["document summary"])
    doc_header = ["document", "page count"]
    for label in labels:
        doc_header.extend(
            [
                f"{label} artifact entries",
                f"{label} line count",
                f"{label} word count",
                f"{label} line recovery %",
                f"{label} word recovery %",
            ]
        )
    doc_header.extend(["abbyy artifact entries", "abbyy line count", "abbyy word count"])
    rows.append(doc_header)

    for document in sorted(doc_totals):
        bucket = doc_totals[document]
        row: list[object] = [document, int(bucket["page_count"])]
        abbyy_lines = int(bucket["abbyy_line_count"])
        abbyy_words = int(bucket["abbyy_word_count"])
        for label in labels:
            line_count = int(bucket[f"{label}_line_count"])
            word_count = int(bucket[f"{label}_word_count"])
            row.extend(
                [
                    int(bucket[f"{label}_entry_count"]),
                    line_count,
                    word_count,
                    f"{(100.0 * line_count / abbyy_lines) if abbyy_lines else 0.0:.2f}%",
                    f"{(100.0 * word_count / abbyy_words) if abbyy_words else 0.0:.2f}%",
                ]
            )
        row.extend([int(bucket["abbyy_entry_count"]), abbyy_lines, abbyy_words])
        rows.append(row)

    rows.append([])
    rows.append(["page summary"])
    page_header = ["document", "page"]
    for label in labels:
        page_header.extend(
            [
                f"{label} source file",
                f"{label} artifact entries",
                f"{label} entry delta vs ABBYY",
                f"{label} line count",
                f"{label} word count",
                f"{label} line recovery %",
                f"{label} word recovery %",
            ]
        )
    page_header.extend(["abbyy source file", "abbyy artifact entries", "abbyy line count", "abbyy word count"])
    rows.append(page_header)

    for page in pages:
        row: list[object] = [page["document"], page["page"]]
        for label in labels:
            row.extend(
                [
                    page[f"{label}_source_file"],
                    page[f"{label}_entry_count"],
                    page[f"{label}_entry_delta_vs_abbyy"],
                    page[f"{label}_line_count"],
                    page[f"{label}_word_count"],
                    f"{float(page[f'{label}_line_recovery_pct']):.2f}%",
                    f"{float(page[f'{label}_word_recovery_pct']):.2f}%",
                ]
            )
        row.extend([page["abbyy_source_file"], page["abbyy_entry_count"], page["abbyy_line_count"], page["abbyy_word_count"]])
        rows.append(row)
    return rows


def write_workbook(output_path: Path, overall_rows: list[list[object]]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Overall Summary"
    for row in overall_rows:
        worksheet.append(row)
    format_sheet(worksheet)
    workbook.save(output_path)


def format_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    wrap_columns = {"workbook source", "source file"}
    for column_cells in worksheet.iter_cols():
        max_length = 0
        header = str(column_cells[0].value or "")
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if header in wrap_columns:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 90)


if __name__ == "__main__":
    main()
