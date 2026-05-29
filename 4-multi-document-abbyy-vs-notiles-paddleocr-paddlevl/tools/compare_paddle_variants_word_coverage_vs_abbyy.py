#!/usr/bin/env python3
"""Compare 3x2, 3x1, and no-tiling word coverage against ABBYY page by page."""

from __future__ import annotations

import argparse
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment


REPO_ROOT = Path(__file__).resolve().parent.parent
EVAL_ROOT = REPO_ROOT.parent
RUN_CONFIGS = [
    (
        "3x2",
        EVAL_ROOT
        / "2-multi-document-abbyy-vs-tiled3x2+docparse-paddleocr-paddlevl"
        / "test-data"
        / "preprocessing-tiling_3x2_dp+paddleocr+vl+clean",
    ),
    (
        "3x1",
        EVAL_ROOT
        / "3-multi-document-abbyy-vs-tiled3x1+docparse-paddleocr-paddlevl"
        / "test-data"
        / "preprocessing-tiling_3x1_dp+paddleocr+vl+clean",
    ),
    (
        "no tiling",
        REPO_ROOT / "test-data" / "preprocess-notiles-paddleocr-paddlevl",
    ),
]
DEFAULT_ABBYY_DIR = REPO_ROOT / "test-data" / "abbyy"
DEFAULT_OUTPUT = REPO_ROOT / "test-results" / "paddle_variants_abbyy_word_coverage.xlsx"

PAGE_NUMBER_RE = re.compile(r"(\d+)$")
WORD_RE = re.compile(r"[0-9A-Za-zÀ-ÖØ-öø-ÿ]+(?:['â€™][0-9A-Za-zÀ-ÖØ-öø-ÿ]+)?")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-3x2-dir", type=Path, default=RUN_CONFIGS[0][1])
    parser.add_argument("--run-3x1-dir", type=Path, default=RUN_CONFIGS[1][1])
    parser.add_argument("--run-notiling-dir", type=Path, default=RUN_CONFIGS[2][1])
    parser.add_argument("--abbyy-dir", type=Path, default=DEFAULT_ABBYY_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    runs = [
        ("3x2", args.run_3x2_dir.resolve()),
        ("3x1", args.run_3x1_dir.resolve()),
        ("no tiling", args.run_notiling_dir.resolve()),
    ]
    summaries, token_rows = compare_runs(runs, args.abbyy_dir.resolve())
    write_workbook(
        args.output.resolve(),
        build_overall_rows(summaries, runs, args.abbyy_dir.resolve()),
        token_rows,
    )
    print(f"Wrote {args.output.resolve()}")
    print(f"Compared pages: {len(summaries)}")


def ocr_txt_files_in_dir(base_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(base_dir.glob("*.txt")):
        name = path.name.lower()
        if name.endswith(".post_ocr_vl_progress.txt"):
            continue
        if "b4-cleaned" in name:
            continue
        files.append(path)
    return files


def page_key_for_path(path: Path) -> str:
    match = PAGE_NUMBER_RE.search(path.stem)
    if match:
        return f"page:{int(match.group(1))}"
    return f"name:{path.name.casefold()}"


def normalize_token(token: str) -> str:
    return unicodedata.normalize("NFKC", token).replace("â€™", "'").casefold()


def read_token_counter(path: Path) -> tuple[Counter[str], int]:
    text = path.read_text(encoding="utf-8", errors="replace")
    tokens = [normalize_token(token) for token in WORD_RE.findall(text)]
    return Counter(tokens), len(tokens)


def matched_token_count(left: Counter[str], right: Counter[str]) -> int:
    return sum((left & right).values())


def matched_unique_count(left: Counter[str], right: Counter[str]) -> int:
    return len(set(left) & set(right))


def percentage(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return (100.0 * numerator) / denominator


def compare_runs(
    runs: list[tuple[str, Path]],
    abbyy_dir: Path,
) -> tuple[list[dict[str, object]], dict[str, list[dict[str, object]]]]:
    run_map = {label: run_dir for label, run_dir in runs}
    token_rows: dict[str, list[dict[str, object]]] = {label: [] for label, _dir in runs}
    summaries: list[dict[str, object]] = []
    missing_documents: list[str] = []
    missing_pages: list[str] = []

    anchor_label, anchor_dir = runs[0]
    for anchor_doc_dir in sorted(path for path in anchor_dir.iterdir() if path.is_dir()):
        abbyy_doc_dir = abbyy_dir / anchor_doc_dir.name
        if not abbyy_doc_dir.is_dir():
            missing_documents.append(anchor_doc_dir.name)
            continue

        run_pages_by_label: dict[str, dict[str, Path]] = {}
        for label, run_dir in runs:
            doc_dir = run_dir / anchor_doc_dir.name
            if not doc_dir.is_dir():
                missing_documents.append(anchor_doc_dir.name)
                continue
            run_pages_by_label[label] = {page_key_for_path(path): path for path in ocr_txt_files_in_dir(doc_dir)}
        if len(run_pages_by_label) != len(runs):
            continue

        abbyy_pages = {page_key_for_path(path): path for path in ocr_txt_files_in_dir(abbyy_doc_dir)}
        for key, anchor_path in sorted(run_pages_by_label[anchor_label].items()):
            abbyy_path = abbyy_pages.get(key)
            if abbyy_path is None:
                missing_pages.append(f"{anchor_doc_dir.name}/{anchor_path.name}")
                continue

            page_paths: dict[str, Path] = {}
            missing_this_page = False
            for label, _run_dir in runs:
                path = run_pages_by_label[label].get(key)
                if path is None:
                    missing_pages.append(f"{anchor_doc_dir.name}/{anchor_path.name}")
                    missing_this_page = True
                    break
                page_paths[label] = path
            if missing_this_page:
                continue

            counter_abbyy, word_count_abbyy = read_token_counter(abbyy_path)
            page_summary: dict[str, object] = {
                "document": anchor_doc_dir.name,
                "page": anchor_path.name,
                "abbyy_source_file": str(abbyy_path),
                "abbyy_words": word_count_abbyy,
                "abbyy_unique_tokens": len(counter_abbyy),
            }

            for label, _run_dir in runs:
                counter_run, word_count_run = read_token_counter(page_paths[label])
                matched_tokens = matched_token_count(counter_run, counter_abbyy)
                matched_unique = matched_unique_count(counter_run, counter_abbyy)
                missing_counter = counter_abbyy - counter_run
                extra_counter = counter_run - counter_abbyy
                page_summary[f"{label}_source_file"] = str(page_paths[label])
                page_summary[f"{label}_words"] = word_count_run
                page_summary[f"{label}_unique_tokens"] = len(counter_run)
                page_summary[f"{label}_matched_tokens"] = matched_tokens
                page_summary[f"{label}_matched_unique"] = matched_unique
                page_summary[f"{label}_coverage_pct"] = percentage(matched_tokens, word_count_abbyy)
                page_summary[f"{label}_unique_coverage_pct"] = percentage(matched_unique, len(counter_abbyy))
                page_summary[f"{label}_precision_pct"] = percentage(matched_tokens, word_count_run)
                page_summary[f"{label}_missing_tokens"] = word_count_abbyy - matched_tokens
                page_summary[f"{label}_extra_tokens"] = word_count_run - matched_tokens
                page_summary[f"{label}_word_delta"] = word_count_run - word_count_abbyy
                token_rows[label].extend(
                    build_token_rows(anchor_doc_dir.name, anchor_path.name, f"{label} missing vs ABBYY", missing_counter)
                )
                token_rows[label].extend(
                    build_token_rows(anchor_doc_dir.name, anchor_path.name, f"{label} extra vs ABBYY", extra_counter)
                )
            summaries.append(page_summary)

    if missing_documents or missing_pages:
        parts: list[str] = []
        if missing_documents:
            parts.append("Missing matching document directories: " + ", ".join(sorted(set(missing_documents))[:10]))
        if missing_pages:
            parts.append("Missing matching pages: " + ", ".join(sorted(set(missing_pages))[:10]))
        raise FileNotFoundError("; ".join(parts))

    summaries.sort(key=lambda item: (str(item["document"]).casefold(), str(item["page"]).casefold()))
    return summaries, token_rows


def build_token_rows(document: str, page: str, category: str, counter: Counter[str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for token, count in sorted(counter.items(), key=lambda item: (-item[1], item[0])):
        rows.append(
            {
                "document": document,
                "page": page,
                "category": category,
                "token": token,
                "count": int(count),
            }
        )
    return rows


def build_overall_rows(
    summaries: list[dict[str, object]],
    runs: list[tuple[str, Path]],
    abbyy_dir: Path,
) -> list[list[object]]:
    labels = [label for label, _run_dir in runs]
    totals: dict[str, dict[str, int]] = {
        label: {
            "words": 0,
            "unique_tokens": 0,
            "matched_tokens": 0,
            "matched_unique": 0,
            "missing_tokens": 0,
            "extra_tokens": 0,
        }
        for label in labels
    }
    total_abbyy_words = 0
    total_abbyy_unique = 0
    doc_totals: dict[str, dict[str, object]] = defaultdict(lambda: {"page_count": 0, "abbyy_words": 0, "abbyy_unique_tokens": 0})

    for summary in summaries:
        document = str(summary["document"])
        doc_bucket = doc_totals[document]
        doc_bucket["page_count"] = int(doc_bucket["page_count"]) + 1
        doc_bucket["abbyy_words"] = int(doc_bucket["abbyy_words"]) + int(summary["abbyy_words"])
        doc_bucket["abbyy_unique_tokens"] = int(doc_bucket["abbyy_unique_tokens"]) + int(summary["abbyy_unique_tokens"])
        total_abbyy_words += int(summary["abbyy_words"])
        total_abbyy_unique += int(summary["abbyy_unique_tokens"])
        for label in labels:
            totals[label]["words"] += int(summary[f"{label}_words"])
            totals[label]["unique_tokens"] += int(summary[f"{label}_unique_tokens"])
            totals[label]["matched_tokens"] += int(summary[f"{label}_matched_tokens"])
            totals[label]["matched_unique"] += int(summary[f"{label}_matched_unique"])
            totals[label]["missing_tokens"] += int(summary[f"{label}_missing_tokens"])
            totals[label]["extra_tokens"] += int(summary[f"{label}_extra_tokens"])
            doc_bucket.setdefault(f"{label}_words", 0)
            doc_bucket.setdefault(f"{label}_matched_tokens", 0)
            doc_bucket.setdefault(f"{label}_missing_tokens", 0)
            doc_bucket.setdefault(f"{label}_extra_tokens", 0)
            doc_bucket[f"{label}_words"] = int(doc_bucket[f"{label}_words"]) + int(summary[f"{label}_words"])
            doc_bucket[f"{label}_matched_tokens"] = int(doc_bucket[f"{label}_matched_tokens"]) + int(summary[f"{label}_matched_tokens"])
            doc_bucket[f"{label}_missing_tokens"] = int(doc_bucket[f"{label}_missing_tokens"]) + int(summary[f"{label}_missing_tokens"])
            doc_bucket[f"{label}_extra_tokens"] = int(doc_bucket[f"{label}_extra_tokens"]) + int(summary[f"{label}_extra_tokens"])

    pages_best_coverage = {label: 0 for label in labels}
    pages_closest_word_count = {label: 0 for label in labels}
    pages_coverage_tied = 0
    pages_word_count_tied = 0
    for summary in summaries:
        coverage_values = {label: float(summary[f"{label}_coverage_pct"]) for label in labels}
        max_coverage = max(coverage_values.values())
        coverage_winners = [label for label, value in coverage_values.items() if value == max_coverage]
        if len(coverage_winners) == 1:
            pages_best_coverage[coverage_winners[0]] += 1
        else:
            pages_coverage_tied += 1

        distance_values = {label: abs(int(summary[f"{label}_word_delta"])) for label in labels}
        min_distance = min(distance_values.values())
        distance_winners = [label for label, value in distance_values.items() if value == min_distance]
        if len(distance_winners) == 1:
            pages_closest_word_count[distance_winners[0]] += 1
        else:
            pages_word_count_tied += 1

    header = ["summary metric", *labels, "abbyy baseline"]
    dataset_row = ["dataset root", *[str(run_dir) for _label, run_dir in runs], str(abbyy_dir)]
    rows: list[list[object]] = [
        header,
        dataset_row,
        ["document count", *([len(doc_totals)] * len(labels)), len(doc_totals)],
        ["page count", *([len(summaries)] * len(labels)), len(summaries)],
        ["word count", *[totals[label]["words"] for label in labels], total_abbyy_words],
        ["unique token count", *[totals[label]["unique_tokens"] for label in labels], total_abbyy_unique],
        ["matched ABBYY token occurrences", *[totals[label]["matched_tokens"] for label in labels], total_abbyy_words],
        ["matched ABBYY unique tokens", *[totals[label]["matched_unique"] for label in labels], total_abbyy_unique],
        ["coverage of ABBYY token occurrences", *[f"{percentage(totals[label]['matched_tokens'], total_abbyy_words):.2f}%" for label in labels], "100.00%"],
        ["coverage of ABBYY unique tokens", *[f"{percentage(totals[label]['matched_unique'], total_abbyy_unique):.2f}%" for label in labels], "100.00%"],
        ["token precision vs ABBYY", *[f"{percentage(totals[label]['matched_tokens'], totals[label]['words']):.2f}%" for label in labels], ""],
        ["extra tokens vs ABBYY", *[totals[label]["extra_tokens"] for label in labels], 0],
        ["missing ABBYY tokens", *[totals[label]["missing_tokens"] for label in labels], 0],
        ["pages with best ABBYY token coverage", *[pages_best_coverage[label] for label in labels], pages_coverage_tied],
        ["pages closest to ABBYY word count", *[pages_closest_word_count[label] for label in labels], pages_word_count_tied],
    ]

    rows.append([])
    rows.append(["document summary"])
    doc_header = ["document", "page count"]
    for label in labels:
        doc_header.extend(
            [
                f"{label} words",
                f"{label} matched ABBYY tokens",
                f"{label} coverage %",
                f"{label} precision %",
                f"{label} missing ABBYY tokens",
                f"{label} extra tokens",
            ]
        )
    doc_header.append("abbyy words")
    rows.append(doc_header)

    for document in sorted(doc_totals):
        bucket = doc_totals[document]
        row: list[object] = [document, int(bucket["page_count"])]
        abbyy_words = int(bucket["abbyy_words"])
        for label in labels:
            words = int(bucket[f"{label}_words"])
            matched_tokens = int(bucket[f"{label}_matched_tokens"])
            row.extend(
                [
                    words,
                    matched_tokens,
                    f"{percentage(matched_tokens, abbyy_words):.2f}%",
                    f"{percentage(matched_tokens, words):.2f}%",
                    int(bucket[f"{label}_missing_tokens"]),
                    int(bucket[f"{label}_extra_tokens"]),
                ]
            )
        row.append(abbyy_words)
        rows.append(row)

    rows.append([])
    rows.append(["page summary"])
    page_header = ["document", "page"]
    for label in labels:
        page_header.extend(
            [
                f"{label} source file",
                f"{label} words",
                f"{label} unique tokens",
                f"{label} matched ABBYY tokens",
                f"{label} matched ABBYY unique tokens",
                f"{label} coverage %",
                f"{label} unique coverage %",
                f"{label} precision %",
                f"{label} missing ABBYY tokens",
                f"{label} extra tokens",
                f"{label} word count delta vs ABBYY",
            ]
        )
    page_header.extend(["abbyy source file", "abbyy words", "abbyy unique tokens"])
    rows.append(page_header)

    for summary in summaries:
        row = [summary["document"], summary["page"]]
        for label in labels:
            row.extend(
                [
                    summary[f"{label}_source_file"],
                    summary[f"{label}_words"],
                    summary[f"{label}_unique_tokens"],
                    summary[f"{label}_matched_tokens"],
                    summary[f"{label}_matched_unique"],
                    f"{float(summary[f'{label}_coverage_pct']):.2f}%",
                    f"{float(summary[f'{label}_unique_coverage_pct']):.2f}%",
                    f"{float(summary[f'{label}_precision_pct']):.2f}%",
                    summary[f"{label}_missing_tokens"],
                    summary[f"{label}_extra_tokens"],
                    summary[f"{label}_word_delta"],
                ]
            )
        row.extend([summary["abbyy_source_file"], summary["abbyy_words"], summary["abbyy_unique_tokens"]])
        rows.append(row)
    return rows


def build_token_sheet_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    output: list[list[object]] = [["document", "page", "category", "token", "count"]]
    for row in sorted(
        rows,
        key=lambda item: (
            str(item["document"]).casefold(),
            str(item["page"]).casefold(),
            str(item["category"]).casefold(),
            -int(item["count"]),
            str(item["token"]),
        ),
    ):
        output.append([row["document"], row["page"], row["category"], row["token"], row["count"]])
    return output


def safe_sheet_title(title: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", title).strip()
    return cleaned[:31] if len(cleaned) > 31 else cleaned


def write_workbook(
    output_path: Path,
    overall_rows: list[list[object]],
    token_rows: dict[str, list[dict[str, object]]],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Overall Summary")
    for row in overall_rows:
        workbook["Overall Summary"].append(row)
    for label, rows in token_rows.items():
        for suffix, predicate in [("Missing vs ABBYY", "missing"), ("Extra vs ABBYY", "extra")]:
            filtered = [row for row in rows if predicate in str(row["category"]).casefold()]
            sheet = workbook.create_sheet(safe_sheet_title(f"{label} {suffix}"))
            for row in build_token_sheet_rows(filtered):
                sheet.append(row)
    for sheet in workbook.worksheets:
        format_sheet(sheet)
    workbook.save(output_path)


def format_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    wrap_columns = {"dataset root", "source file", "counting note", "token"}
    for column_cells in worksheet.iter_cols():
        max_length = 0
        header = str(column_cells[0].value or "")
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if header in wrap_columns:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 90)


if __name__ == "__main__":
    main()
