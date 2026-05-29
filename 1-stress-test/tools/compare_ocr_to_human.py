#!/usr/bin/env python3
from __future__ import annotations

import argparse
import difflib
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_HUMAN = REPO_ROOT / "test-results/large - human/first-article.txt"
DEFAULT_ABBYY = REPO_ROOT / "test-results/large - abby/first-article.txt"
DEFAULT_PADDLECOMP_DIR = REPO_ROOT / "test-results/paddlecomp"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "test-results"
DEFAULT_PREFIX = "first_article_ocr_comparison"

TOKEN_RE = re.compile(r"[A-Za-z0-9]+(?:['-][A-Za-z0-9]+)?")
SMART_TRANSLATION = str.maketrans(
    {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\ufeff": "",
    }
)


@dataclass(frozen=True)
class TextVersion:
    label: str
    path: Path
    raw_text: str
    normalized_text: str
    tokens: list[str]
    nonblank_lines: list[str]
    normalized_lines: list[str]


@dataclass(frozen=True)
class EditCounts:
    distance: int
    substitutions: int
    insertions: int
    deletions: int


@dataclass(frozen=True)
class Comparison:
    label: str
    path: Path
    human_word_count: int
    ocr_word_count: int
    human_char_count: int
    ocr_char_count: int
    human_line_count: int
    ocr_line_count: int
    word_counts: EditCounts
    char_distance: int
    diff_rows: list[dict[str, str | int]]
    line_diff_rows: list[dict[str, str | int]]

    @property
    def word_error_rate(self) -> float:
        return ratio(self.word_counts.distance, self.human_word_count)

    @property
    def word_accuracy(self) -> float:
        return max(0.0, 1.0 - self.word_error_rate)

    @property
    def char_error_rate(self) -> float:
        return ratio(self.char_distance, self.human_char_count)

    @property
    def char_accuracy(self) -> float:
        return max(0.0, 1.0 - self.char_error_rate)


def main() -> None:
    args = parse_args()
    human_path = args.human.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    human = load_text_version("human", human_path)
    ocr_versions = discover_ocr_versions(
        abbyy_path=args.abbyy,
        paddlecomp_dir=args.paddlecomp_dir,
        explicit_ocr=args.ocr,
    )
    if not ocr_versions:
        raise SystemExit("No OCR first-article.txt files found.")

    comparisons = [
        compare_versions(human, load_text_version(label, path), args.diff_limit)
        for label, path in ocr_versions
    ]
    comparisons.sort(key=lambda item: (item.word_error_rate, item.char_error_rate))

    workbook_path = output_dir / f"{args.prefix}.xlsx"
    write_excel_workbook(workbook_path, human, comparisons)

    print(f"Wrote {workbook_path}")
    print()
    print("Ranked by normalized word error rate:")
    for rank, comparison in enumerate(comparisons, start=1):
        print(
            f"{rank}. {comparison.label}: "
            f"word accuracy {format_percent(comparison.word_accuracy)}, "
            f"WER {format_percent(comparison.word_error_rate)}, "
            f"CER {format_percent(comparison.char_error_rate)}"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compare first-article OCR transcriptions against a human transcription "
            "and write a shareable Excel report."
        )
    )
    parser.add_argument("--human", type=Path, default=DEFAULT_HUMAN)
    parser.add_argument("--abbyy", type=Path, default=DEFAULT_ABBYY)
    parser.add_argument("--paddlecomp-dir", type=Path, default=DEFAULT_PADDLECOMP_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--prefix", default=DEFAULT_PREFIX)
    parser.add_argument(
        "--diff-limit",
        type=int,
        default=12,
        help="Maximum representative word-level diffs to include per OCR.",
    )
    parser.add_argument(
        "--ocr",
        action="append",
        default=[],
        metavar="LABEL=PATH",
        help="Additional OCR transcript to compare. Can be repeated.",
    )
    return parser.parse_args()


def discover_ocr_versions(
    abbyy_path: Path,
    paddlecomp_dir: Path,
    explicit_ocr: Sequence[str],
) -> list[tuple[str, Path]]:
    versions: list[tuple[str, Path]] = []
    if abbyy_path.exists():
        versions.append(("abbyy", abbyy_path.resolve()))

    if paddlecomp_dir.exists():
        for path in sorted(paddlecomp_dir.glob("*/first-article.txt")):
            versions.append((f"paddlecomp/{path.parent.name}", path.resolve()))

    for raw in explicit_ocr:
        if "=" not in raw:
            raise SystemExit(f"--ocr must be LABEL=PATH, got: {raw}")
        label, path_text = raw.split("=", 1)
        path = Path(path_text).expanduser()
        if not path.exists():
            raise SystemExit(f"OCR path does not exist: {path}")
        versions.append((label.strip() or path.stem, path.resolve()))

    seen: set[Path] = set()
    unique: list[tuple[str, Path]] = []
    for label, path in versions:
        if path in seen:
            continue
        seen.add(path)
        unique.append((label, path))
    return unique


def load_text_version(label: str, path: Path) -> TextVersion:
    raw_text = path.read_text(encoding="utf-8-sig")
    normalized_text = normalize_text(raw_text)
    nonblank_lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    return TextVersion(
        label=label,
        path=path,
        raw_text=raw_text,
        normalized_text=normalized_text,
        tokens=tokenize(normalized_text),
        nonblank_lines=nonblank_lines,
        normalized_lines=[normalize_line(line) for line in nonblank_lines],
    )


def normalize_text(text: str) -> str:
    value = unicodedata.normalize("NFKC", text).translate(SMART_TRANSLATION)
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "", value)
    value = value.lower()
    value = re.sub(r"[ \t\f\v]+", " ", value)
    value = re.sub(r"\s*\n\s*", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def tokenize(text: str) -> list[str]:
    return TOKEN_RE.findall(text)


def normalize_line(text: str) -> str:
    value = unicodedata.normalize("NFKC", text).translate(SMART_TRANSLATION)
    value = value.lower()
    value = re.sub(r"[ \t\f\v]+", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def compare_versions(
    human: TextVersion,
    ocr: TextVersion,
    diff_limit: int,
) -> Comparison:
    word_counts = levenshtein_counts(human.tokens, ocr.tokens)
    char_distance = levenshtein_distance(human.normalized_text, ocr.normalized_text)
    return Comparison(
        label=ocr.label,
        path=ocr.path,
        human_word_count=len(human.tokens),
        ocr_word_count=len(ocr.tokens),
        human_char_count=len(human.normalized_text),
        ocr_char_count=len(ocr.normalized_text),
        human_line_count=count_nonblank_lines(human.raw_text),
        ocr_line_count=count_nonblank_lines(ocr.raw_text),
        word_counts=word_counts,
        char_distance=char_distance,
        diff_rows=build_diff_rows(human.tokens, ocr.tokens, diff_limit),
        line_diff_rows=build_line_diff_rows(human, ocr),
    )


def levenshtein_distance(left: Sequence[str] | str, right: Sequence[str] | str) -> int:
    if len(left) < len(right):
        left, right = right, left

    previous = list(range(len(right) + 1))
    for left_index, left_value in enumerate(left, start=1):
        current = [left_index]
        for right_index, right_value in enumerate(right, start=1):
            current.append(
                min(
                    current[-1] + 1,
                    previous[right_index] + 1,
                    previous[right_index - 1] + (left_value != right_value),
                )
            )
        previous = current
    return previous[-1]


def levenshtein_counts(
    human_tokens: Sequence[str],
    ocr_tokens: Sequence[str],
) -> EditCounts:
    rows = len(human_tokens) + 1
    cols = len(ocr_tokens) + 1
    costs = [[0] * cols for _ in range(rows)]
    ops = [[""] * cols for _ in range(rows)]

    for row in range(1, rows):
        costs[row][0] = row
        ops[row][0] = "delete"
    for col in range(1, cols):
        costs[0][col] = col
        ops[0][col] = "insert"

    for row in range(1, rows):
        for col in range(1, cols):
            if human_tokens[row - 1] == ocr_tokens[col - 1]:
                costs[row][col] = costs[row - 1][col - 1]
                ops[row][col] = "equal"
                continue

            choices = [
                (costs[row - 1][col - 1] + 1, "substitute"),
                (costs[row - 1][col] + 1, "delete"),
                (costs[row][col - 1] + 1, "insert"),
            ]
            costs[row][col], ops[row][col] = min(choices, key=lambda item: item[0])

    row = len(human_tokens)
    col = len(ocr_tokens)
    substitutions = insertions = deletions = 0
    while row > 0 or col > 0:
        op = ops[row][col]
        if op == "equal":
            row -= 1
            col -= 1
        elif op == "substitute":
            substitutions += 1
            row -= 1
            col -= 1
        elif op == "delete":
            deletions += 1
            row -= 1
        elif op == "insert":
            insertions += 1
            col -= 1
        else:
            break

    return EditCounts(
        distance=costs[-1][-1],
        substitutions=substitutions,
        insertions=insertions,
        deletions=deletions,
    )


def build_diff_rows(
    human_tokens: Sequence[str],
    ocr_tokens: Sequence[str],
    limit: int,
) -> list[dict[str, str | int]]:
    rows: list[dict[str, str | int]] = []
    matcher = difflib.SequenceMatcher(
        a=list(human_tokens),
        b=list(ocr_tokens),
        autojunk=False,
    )
    for tag, human_start, human_end, ocr_start, ocr_end in matcher.get_opcodes():
        if tag == "equal":
            continue
        rows.append(
            {
                "operation": tag,
                "human_start_word": human_start,
                "human_end_word": human_end,
                "ocr_start_word": ocr_start,
                "ocr_end_word": ocr_end,
                "human_text": snippet(human_tokens, human_start, human_end),
                "ocr_text": snippet(ocr_tokens, ocr_start, ocr_end),
            }
        )
        if len(rows) >= limit:
            break
    return rows


def build_line_diff_rows(
    human: TextVersion,
    ocr: TextVersion,
) -> list[dict[str, str | int]]:
    rows: list[dict[str, str | int]] = []
    matcher = difflib.SequenceMatcher(
        a=human.normalized_lines,
        b=ocr.normalized_lines,
        autojunk=False,
    )
    for tag, human_start, human_end, ocr_start, ocr_end in matcher.get_opcodes():
        if tag == "equal":
            continue
        rows.append(
            {
                "operation": tag,
                "human_start_line": human_start + 1,
                "human_end_line": human_end,
                "ocr_start_line": ocr_start + 1,
                "ocr_end_line": ocr_end,
                "human_line_count": max(0, human_end - human_start),
                "ocr_line_count": max(0, ocr_end - ocr_start),
                "human_text": join_lines(human.nonblank_lines[human_start:human_end]),
                "ocr_text": join_lines(ocr.nonblank_lines[ocr_start:ocr_end]),
            }
        )
    return rows


def join_lines(lines: Sequence[str]) -> str:
    if not lines:
        return ""
    return "\n".join(lines)


def snippet(tokens: Sequence[str], start: int, end: int, context: int = 6) -> str:
    left = max(0, start - context)
    right = min(len(tokens), end + context)
    prefix = "... " if left > 0 else ""
    suffix = " ..." if right < len(tokens) else ""
    return prefix + " ".join(tokens[left:right]) + suffix


def write_excel_workbook(
    path: Path,
    human: TextVersion,
    comparisons: Sequence[Comparison],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    write_summary_sheet(summary_sheet, comparisons)

    diff_sheet = workbook.create_sheet("Representative Diffs")
    write_diff_sheet(diff_sheet, comparisons)

    for comparison in comparisons:
        sheet = workbook.create_sheet(make_line_diff_sheet_name(comparison.label))
        write_line_diff_sheet(sheet, comparison)

    notes_sheet = workbook.create_sheet("Notes")
    write_notes_sheet(notes_sheet, human)

    workbook.save(path)


def write_summary_sheet(sheet, comparisons: Sequence[Comparison]) -> None:
    headers = [
        "Rank",
        "OCR",
        "Source path",
        "Word accuracy",
        "WER",
        "CER",
        "Word edits",
        "Substitutions",
        "Insertions",
        "Deletions",
        "Character edits",
        "Human words",
        "OCR words",
        "Human chars",
        "OCR chars",
        "Human nonblank lines",
        "OCR nonblank lines",
    ]
    sheet.append(headers)
    for rank, comparison in enumerate(comparisons, start=1):
        counts = comparison.word_counts
        sheet.append(
            [
                rank,
                comparison.label,
                str(comparison.path),
                comparison.word_accuracy,
                comparison.word_error_rate,
                comparison.char_error_rate,
                counts.distance,
                counts.substitutions,
                counts.insertions,
                counts.deletions,
                comparison.char_distance,
                comparison.human_word_count,
                comparison.ocr_word_count,
                comparison.human_char_count,
                comparison.ocr_char_count,
                comparison.human_line_count,
                comparison.ocr_line_count,
            ]
        )

    style_sheet(sheet, freeze="A2")
    for column in ("D", "E", "F"):
        for cell in sheet[column][1:]:
            cell.number_format = "0.00%"
    set_column_widths(
        sheet,
        {
            "A": 8,
            "B": 52,
            "C": 90,
            "D": 16,
            "E": 12,
            "F": 12,
            "G": 12,
            "H": 14,
            "I": 12,
            "J": 12,
            "K": 14,
        },
    )


def write_diff_sheet(sheet, comparisons: Sequence[Comparison]) -> None:
    sheet.append(
        [
            "OCR",
            "Diff #",
            "Operation",
            "Human start word",
            "Human end word",
            "OCR start word",
            "OCR end word",
            "Human text",
            "OCR text",
        ]
    )
    for comparison in comparisons:
        for index, row in enumerate(comparison.diff_rows, start=1):
            sheet.append(
                [
                    comparison.label,
                    index,
                    row["operation"],
                    row["human_start_word"],
                    row["human_end_word"],
                    row["ocr_start_word"],
                    row["ocr_end_word"],
                    row["human_text"],
                    row["ocr_text"],
                ]
            )

    style_sheet(sheet, freeze="A2")
    set_column_widths(
        sheet,
        {
            "A": 52,
            "B": 8,
            "C": 12,
            "D": 16,
            "E": 16,
            "F": 14,
            "G": 14,
            "H": 90,
            "I": 90,
        },
    )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_line_diff_sheet(sheet, comparison: Comparison) -> None:
    sheet.append(
        [
            "Diff #",
            "Operation",
            "Human start line",
            "Human end line",
            "OCR start line",
            "OCR end line",
            "Human line count",
            "OCR line count",
            "Human text",
            "OCR text",
        ]
    )
    for index, row in enumerate(comparison.line_diff_rows, start=1):
        sheet.append(
            [
                index,
                row["operation"],
                row["human_start_line"],
                row["human_end_line"],
                row["ocr_start_line"],
                row["ocr_end_line"],
                row["human_line_count"],
                row["ocr_line_count"],
                row["human_text"],
                row["ocr_text"],
            ]
        )

    style_sheet(sheet, freeze="A2")
    set_column_widths(
        sheet,
        {
            "A": 8,
            "B": 12,
            "C": 14,
            "D": 14,
            "E": 14,
            "F": 14,
            "G": 14,
            "H": 14,
            "I": 88,
            "J": 88,
        },
    )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_notes_sheet(sheet, human: TextVersion) -> None:
    rows = [
        ("Generated", datetime.now(timezone.utc).isoformat()),
        ("Human transcription", str(human.path)),
        (
            "Normalization",
            "Unicode is normalized, smart punctuation is standardized, case is ignored, "
            "whitespace is collapsed, and line-break hyphenation is joined before scoring.",
        ),
        (
            "Ranking",
            "Rows are ranked by normalized word error rate, then character error rate.",
        ),
        (
            "WER",
            "Word error rate = word edit distance / human word count.",
        ),
        (
            "S / I / D",
            "Substitutions, insertions, and deletions in the word-level alignment.",
        ),
        (
            "Per-model line sheets",
            "Each OCR has its own sheet listing every non-equal line-level replace/insert/delete block versus the human transcription.",
        ),
    ]
    sheet.append(["Field", "Value"])
    for row in rows:
        sheet.append(row)
    style_sheet(sheet, freeze="A2")
    set_column_widths(sheet, {"A": 24, "B": 120})
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def style_sheet(sheet, freeze: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.auto_filter.ref = sheet.dimensions
    if freeze:
        sheet.freeze_panes = freeze


def set_column_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for index in range(1, sheet.max_column + 1):
        column = get_column_letter(index)
        if column not in widths:
            sheet.column_dimensions[column].width = 14


def count_nonblank_lines(text: str) -> int:
    return sum(1 for line in text.splitlines() if line.strip())


def ratio(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def format_percent(value: float) -> str:
    return f"{value * 100:.2f}%"


def make_line_diff_sheet_name(label: str) -> str:
    short = friendly_sheet_label(label)
    base = f"Lines - {short}"
    return base[:31]


def friendly_sheet_label(label: str) -> str:
    lower = label.lower()
    if label == "abbyy":
        return "ABBYY"
    if "chandra" in lower:
        return "Chandra"
    if "paddlevl" in lower:
        return "PaddleVL"
    if "deepseek" in lower:
        return "DeepSeek"
    if "olmocr" in lower:
        return "olmOCR"
    if "ppcor-preproccess-3x2-docparse" in lower:
        return "PaddleOCR (No VL)"
    return Path(label).name[:20]


if __name__ == "__main__":
    main()
