#!/usr/bin/env python3
"""Generate side-by-side OCR error reports as CSV or Excel.

Default usage, using the current Paddle/ABBYY comparison files:

  python tools/compare_ocr_errors.py

Custom usage:

  python tools/compare_ocr_errors.py \
    --paddle path/to/paddle.txt \
    --abbyy path/to/abbyy.txt \
    --output path/to/errors.csv

Directory-to-directory Excel usage:

  python tools/compare_ocr_errors.py \
    --paddle-dir test-results/paddlecomp \
    --abbyy-dir "test-results/large - abby" \
    --excel-output test-results/paddle_vs_abbyy_errors_artifacts.xlsx

The output CSV is wide: summary columns are on the left, then each error type
gets a Paddle column and an ABBYY column with line-numbered snippets listed
downward.
"""

from __future__ import annotations

import argparse
import csv
import difflib
import math
import re
import sys
import unicodedata
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment


REPO_ROOT = Path(__file__).resolve().parent.parent

DEFAULT_PADDLE = (
    REPO_ROOT
    / "test-results/paddlecomp/large-ppocr-preprocess-3x2-docparse-olm/"
    / "oocihm.N_00155_18750610.4.txt"
)
DEFAULT_ABBYY = REPO_ROOT / "test-results/large - abby/oocihm.N_00155_18750610.4.txt"
DEFAULT_OUTPUT = REPO_ROOT / "test-results/paddle_vs_abbyy_errors_artifacts.csv"
DEFAULT_PADDLE_DIR = REPO_ROOT / "test-results/paddlecomp"
DEFAULT_ABBYY_DIR = REPO_ROOT / "test-results/large - abby"
DEFAULT_EXCEL_OUTPUT = REPO_ROOT / "test-results/paddle_vs_abbyy_errors_artifacts.xlsx"

WORD_RE = re.compile(r"[A-Za-z0-9]+(?:[’'][A-Za-z0-9]+)?")

# Keep common historical-newspaper punctuation and French accents out of the
# "suspicious" bucket. The junk list catches glyphs observed in this comparison.
ALLOWED_NONASCII = set("’‘“”—–£àéèêôûç°«»")
JUNK_CHARS = set("\ufeff•□口子广福是心醒。，；：？−€")

ERROR_TYPE_ORDER = [
    "reading order",
    "Possible duplicate lines",
    "OCR token/phrase mismatch",
    "suspicious token",
    "suspicious glyph",
    "BOM/control character",
    "isolated marker/separator",
    "punctuation artifact",
    "line-start artifact",
    "spacing/joined text",
    "abbreviation/spacing",
    "other",
]
ERROR_TYPE_PRIORITY = {error_type: index for index, error_type in enumerate(ERROR_TYPE_ORDER)}


def clean_snippet(value: str, limit: int = 170) -> str:
    value = value.replace("\ufeff", "[BOM]").replace("\r", "")
    value = re.sub(r"\s+", " ", value).strip()
    if len(value) > limit:
        return value[: limit - 3] + "..."
    return value


def normalize_word(value: str) -> str:
    return unicodedata.normalize("NFKC", value).replace("’", "'").lower()


def normalize_line(value: str) -> str:
    value = unicodedata.normalize("NFKC", value.replace("\ufeff", "")).lower()
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def is_mixed_case_artifact(token: str) -> bool:
    if len(token) < 4:
        return False
    if re.match(r"(?:Mc|Mac)[A-Z][a-z]", token):
        return False
    if re.match(r"[A-Z]\.[A-Z]", token):
        return False
    if token.isupper() or token.istitle() or token.islower():
        return False
    return bool(re.search(r"[a-z][A-Z][a-z]|[A-Z]{2,}[a-z]|[a-z][A-Z]{2,}", token))


def is_letter_digit_artifact(token: str) -> bool:
    lower = token.lower()
    if re.fullmatch(r"\d+(st|nd|rd|th)", lower):
        return False
    if re.fullmatch(r"\d+[a-z]", lower) and lower[-1] in {"d", "s"}:
        return False
    return bool(re.search(r"[A-Za-z][0-9]|[0-9][A-Za-z]", token))


def suspicious_chars(line: str) -> list[str]:
    chars: list[str] = []
    for char in line:
        if ord(char) <= 127:
            continue
        if char in JUNK_CHARS:
            chars.append(char)
            continue
        if char in ALLOWED_NONASCII:
            continue
        name = unicodedata.name(char, "")
        category = unicodedata.category(char)
        if "CJK" in name or "FULLWIDTH" in name or category.startswith(("S", "C")):
            chars.append(char)
    return chars


def bad_line_start(stripped: str) -> bool:
    return bool(
        re.match(r"^[.][A-Za-z]", stripped)
        or re.match(r"^[.][\"“”']{1,3}[A-Za-z]", stripped)
        or re.match(r"^[-_=]{2,}\w", stripped)
        or re.match(r"^[-_=]*[•*]+\s*\w", stripped)
        or re.match(r"^[\"“”][.,;:!?]+\w", stripped)
    )


def line_reasons(line: str) -> list[str]:
    reasons: list[str] = []
    raw = line.rstrip("\n")
    stripped = raw.strip()
    if not stripped:
        return reasons

    if raw.startswith("\ufeff"):
        reasons.append("BOM/control character at start of file")

    chars = suspicious_chars(raw)
    if chars:
        unique = list(dict.fromkeys(chars))
        names = ", ".join(
            f"U+{ord(char):04X} {unicodedata.name(char, 'UNKNOWN')}"
            for char in unique[:8]
        )
        reasons.append(f"stray/suspicious non-ASCII glyph(s): {names}")

    alnum_count = sum(char.isalnum() for char in stripped)
    alpha_count = sum(char.isalpha() for char in stripped)
    if len(stripped) <= 3 and (alnum_count <= 1 or stripped.isdigit()):
        reasons.append("isolated marker/symbol line")
    elif alpha_count == 0 and len(stripped) <= 45:
        reasons.append("symbol-only separator/garbage line")

    punctuation_reasons = punctuation_artifacts(raw)
    if punctuation_reasons:
        reasons.append("punctuation artifact: " + "; ".join(punctuation_reasons))

    if bad_line_start(stripped):
        reasons.append("line starts with punctuation glued to text")
    if re.match(r"^[0-9][A-Za-z]{2,}", stripped) and not re.match(
        r"^\d+(st|nd|rd|th)\b", stripped.lower()
    ):
        reasons.append("line starts with digit/letter garbage token")
    if re.match(r"^[a-z][A-Z][a-z]", stripped):
        reasons.append("line starts with lowercase prefix glued to capitalized word")

    token_reasons = token_artifacts(raw)
    if token_reasons:
        reasons.append("suspicious token(s): " + ", ".join(token_reasons[:10]))

    if re.search(r"\b(?:[A-Za-z]\.){2,}[A-Za-z]?\.?", raw) and not re.search(
        r"\b(?:D\. A\. G\.|R\. B\.|A\. H\.|N\. Y\.)", raw
    ):
        reasons.append("abbreviation/spacing may be garbled")
    if re.search(r"\b[a-z]{2,}\.[A-Z][a-z]", raw):
        reasons.append("missing space after period")
    if re.search(r"\bI\.am\b|\b[a-z]{2,},[a-z]{2,}\b|\b[a-z]{2,};[A-Za-z]", raw):
        reasons.append("missing space around punctuation")
    if re.search(r"\btoexpose\b|\bduringthe\b|\blastfew\b|\bthe[A-Z][a-z]+\b", raw):
        reasons.append("likely joined words/missing space")

    return reasons


def punctuation_artifacts(line: str) -> list[str]:
    dot_leader = bool(re.search(r"\.{4,}\s*-?\s*\d+\s*$", line))
    patterns = [
        (r"\.{2,}", "repeated periods"),
        (r",{2,}", "repeated commas"),
        (r";{2,}", "repeated semicolons"),
        (r":{2,}", "repeated colons"),
        (r"!{2,}", "repeated exclamation marks"),
        (r"\?{2,}", "repeated question marks"),
        (r"--+", "double/multiple hyphen"),
        (r"[-—]\s*[*•]", "hyphen/bullet debris"),
        (r"[\"“”][.,;:!?]{2,}|[.,;:!?]{2,}[\"“”]", "quote/punctuation debris"),
        (r"\)\.\)|\.\)\.", "parenthesis/punctuation debris"),
    ]

    hits: list[str] = []
    for pattern, label in patterns:
        if label == "repeated periods" and dot_leader:
            continue
        if re.search(pattern, line):
            hits.append(label)
    return list(dict.fromkeys(hits))


def token_artifacts(line: str) -> list[str]:
    reasons: list[str] = []
    for token in WORD_RE.findall(line):
        if is_letter_digit_artifact(token):
            reasons.append(f"{token} (letter/digit mix)")
        elif is_mixed_case_artifact(token):
            reasons.append(f"{token} (odd internal casing)")
        elif len(token) >= 18 and token.isalpha() and token.lower() not in {
            "characteristically",
            "unconstitutionally",
        }:
            reasons.append(f"{token} (very long joined token)")
        elif re.search(r"[a-z]{2,}[A-Z][a-z]{2,}", token) and not re.match(
            r"(Mc|Mac)[A-Z]", token
        ):
            reasons.append(f"{token} (words likely joined)")
    return reasons


def add_issue(
    issues: OrderedDict[tuple[int, str], list[str]],
    line_no: int,
    reason: str,
    line: str,
) -> None:
    key = (line_no, clean_snippet(line))
    issues.setdefault(key, [])
    issues[key].append(reason)


def line_level_issues(lines: list[str], *, detect_overlap: bool) -> OrderedDict[tuple[int, str], list[str]]:
    issues: OrderedDict[tuple[int, str], list[str]] = OrderedDict()

    for index, line in enumerate(lines, 1):
        for reason in line_reasons(line):
            add_issue(issues, index, reason, line)

    if detect_overlap:
        add_overlap_issues(lines, issues)

    return issues


def add_overlap_issues(
    lines: list[str],
    issues: OrderedDict[tuple[int, str], list[str]],
) -> None:
    normalized = [normalize_line(line) for line in lines]
    for index in range(1, len(lines)):
        previous = normalized[index - 1]
        current = normalized[index]
        if len(previous) < 24 or len(current) < 24:
            continue
        if previous.startswith("no ") and current.startswith("no "):
            continue

        ratio = difflib.SequenceMatcher(None, previous, current).ratio()
        shared_words = set(previous.split()) & set(current.split())
        if ratio >= 0.62 and len(shared_words) >= 3:
            add_issue(
                issues,
                index + 1,
                f"possible duplicated/overlap line with previous line (similarity {ratio:.0%})",
                lines[index],
            )


def token_stream(text: str) -> list[dict[str, object]]:
    stream: list[dict[str, object]] = []
    for line_no, line in enumerate(text.splitlines(), 1):
        for match in WORD_RE.finditer(line):
            raw = match.group(0)
            stream.append(
                {
                    "norm": normalize_word(raw),
                    "raw": raw,
                    "line": line_no,
                    "line_text": line,
                }
            )
    return stream


def levenshtein(left: str, right: str) -> int:
    if abs(len(left) - len(right)) > 3:
        return 4

    previous = list(range(len(right) + 1))
    for left_index, left_char in enumerate(left, 1):
        current = [left_index]
        for right_index, right_char in enumerate(right, 1):
            current.append(
                min(
                    current[-1] + 1,
                    previous[right_index] + 1,
                    previous[right_index - 1] + (left_char != right_char),
                )
            )
        previous = current
    return previous[-1]


def build_ngram_scorer(tokens: list[str]):
    combined_counts = Counter(tokens)
    base_tokens = [
        token
        for token, count in combined_counts.items()
        if len(token) >= 3 and token.isalpha() and count >= 2
    ]

    ngrams: Counter[str] = Counter()
    for token in base_tokens:
        padded = "^" + token + "$"
        for ngram_size in (2, 3):
            for index in range(len(padded) - ngram_size + 1):
                ngrams[padded[index : index + ngram_size]] += 1

    vocabulary_size = max(len(ngrams), 1)
    total = sum(ngrams.values()) + vocabulary_size

    def score(token: str) -> float:
        if not token.isalpha():
            return -99.0
        padded = "^" + token.lower() + "$"
        values = []
        for ngram_size in (2, 3):
            for index in range(len(padded) - ngram_size + 1):
                ngram = padded[index : index + ngram_size]
                values.append(math.log((ngrams[ngram] + 1) / total))
        return sum(values) / max(len(values), 1)

    return score


def obviously_bad_token(token: str) -> bool:
    return bool(
        is_letter_digit_artifact(token)
        or (len(token) >= 4 and token.isalpha() and not re.search(r"[aeiouy]", token))
        or re.search(r"(.)\1\1", token)
    )


def possessive_variant(left: str, right: str) -> bool:
    def clean(value: str) -> str:
        return value.replace("'", "").replace("’", "")

    return clean(left) == clean(right) or clean(left).rstrip("s") == clean(right).rstrip("s")


def add_token_comparison_issues(
    paddle_stream: list[dict[str, object]],
    abbyy_stream: list[dict[str, object]],
    paddle_issues: OrderedDict[tuple[int, str], list[str]],
    abbyy_issues: OrderedDict[tuple[int, str], list[str]],
) -> None:
    paddle_tokens = [str(item["norm"]) for item in paddle_stream]
    abbyy_tokens = [str(item["norm"]) for item in abbyy_stream]
    combined_counts = Counter(paddle_tokens) + Counter(abbyy_tokens)
    char_score = build_ngram_scorer(paddle_tokens + abbyy_tokens)

    matcher = difflib.SequenceMatcher(None, paddle_tokens, abbyy_tokens, autojunk=False)
    for tag, p_start, p_end, a_start, a_end in matcher.get_opcodes():
        if tag == "equal" or max(p_end - p_start, a_end - a_start) > 8:
            continue

        paddle_slice = paddle_stream[p_start:p_end]
        abbyy_slice = abbyy_stream[a_start:a_end]
        if len(paddle_slice) == 1 and len(abbyy_slice) == 1:
            add_single_token_mismatch(
                paddle_slice[0],
                abbyy_slice[0],
                combined_counts,
                char_score,
                paddle_issues,
                abbyy_issues,
            )
        else:
            add_phrase_mismatch(
                paddle_slice,
                abbyy_slice,
                paddle_issues,
                abbyy_issues,
            )


def add_single_token_mismatch(
    paddle_token: dict[str, object],
    abbyy_token: dict[str, object],
    combined_counts: Counter[str],
    char_score,
    paddle_issues: OrderedDict[tuple[int, str], list[str]],
    abbyy_issues: OrderedDict[tuple[int, str], list[str]],
) -> None:
    paddle_norm = str(paddle_token["norm"])
    abbyy_norm = str(abbyy_token["norm"])

    if possessive_variant(paddle_norm, abbyy_norm):
        return
    if levenshtein(paddle_norm, abbyy_norm) > 2:
        return
    if len(paddle_norm) < 4 or len(abbyy_norm) < 4:
        return

    reason = (
        f'candidate OCR token mismatch vs other OCR: "{paddle_token["raw"]}" '
        f'vs "{abbyy_token["raw"]}"'
    )

    # Treat ABBYY as the cleaner spelling baseline unless its token is visibly bad.
    if obviously_bad_token(abbyy_norm):
        add_issue(
            abbyy_issues,
            int(abbyy_token["line"]),
            reason,
            str(abbyy_token["line_text"]),
        )
    elif obviously_bad_token(paddle_norm):
        add_issue(
            paddle_issues,
            int(paddle_token["line"]),
            reason,
            str(paddle_token["line_text"]),
        )
    elif combined_counts[paddle_norm] < combined_counts[abbyy_norm] and combined_counts[
        abbyy_norm
    ] >= max(3, 3 * combined_counts[paddle_norm]):
        add_issue(
            paddle_issues,
            int(paddle_token["line"]),
            reason,
            str(paddle_token["line_text"]),
        )
    elif char_score(paddle_norm) < char_score(abbyy_norm) - 0.18:
        add_issue(
            paddle_issues,
            int(paddle_token["line"]),
            reason,
            str(paddle_token["line_text"]),
        )


def add_phrase_mismatch(
    paddle_slice: list[dict[str, object]],
    abbyy_slice: list[dict[str, object]],
    paddle_issues: OrderedDict[tuple[int, str], list[str]],
    abbyy_issues: OrderedDict[tuple[int, str], list[str]],
) -> None:
    paddle_bad = any(obviously_bad_token(str(item["norm"])) for item in paddle_slice)
    abbyy_bad = any(obviously_bad_token(str(item["norm"])) for item in abbyy_slice)
    if not paddle_bad and not abbyy_bad:
        return

    paddle_text = " ".join(str(item["raw"]) for item in paddle_slice)
    abbyy_text = " ".join(str(item["raw"]) for item in abbyy_slice)
    reason = f'localized OCR phrase divergence vs other OCR: "{paddle_text}" vs "{abbyy_text}"'

    if paddle_bad and paddle_slice:
        first = paddle_slice[0]
        add_issue(paddle_issues, int(first["line"]), reason, str(first["line_text"]))
    if abbyy_bad and abbyy_slice:
        first = abbyy_slice[0]
        add_issue(abbyy_issues, int(first["line"]), reason, str(first["line_text"]))


def add_known_structural_notes(
    paddle_lines: list[str],
    abbyy_lines: list[str],
    paddle_issues: OrderedDict[tuple[int, str], list[str]],
    abbyy_issues: OrderedDict[tuple[int, str], list[str]],
) -> None:
    """Add document-specific notes when the expected line context is present."""

    known_notes = [
        (
            "abbyy",
            43,
            "The American Riflemen",
            "reading-order artifact: Mackenzie article breaks here and resumes much later near line 718",
        ),
        (
            "abbyy",
            718,
            "only prove an unmitigated curse",
            "reading-order artifact: continuation of Mackenzie article resumes after unrelated column/article material",
        ),
        (
            "paddle",
            23,
            "cons consistencies in h his policy",
            "tile/overlap duplication inside Mackenzie article: phrase from previous line repeated and garbled",
        ),
        (
            "paddle",
            2499,
            "terest anfong the Assizes",
            "tile/overlap duplication: previous line partially repeated and merged with next article text",
        ),
    ]

    for source, line_no, expected, reason in known_notes:
        lines = paddle_lines if source == "paddle" else abbyy_lines
        issues = paddle_issues if source == "paddle" else abbyy_issues
        if line_no - 1 < len(lines) and expected in lines[line_no - 1]:
            add_issue(issues, line_no, reason, lines[line_no - 1])


def error_type_for_reason(reason: str) -> str:
    if reason.startswith("reading-order artifact"):
        return "reading order"
    if reason.startswith("tile/overlap") or reason.startswith("possible duplicated/overlap"):
        return "Possible duplicate lines"
    if reason.startswith("candidate OCR token mismatch") or reason.startswith(
        "localized OCR phrase divergence"
    ):
        return "OCR token/phrase mismatch"
    if reason.startswith("suspicious token"):
        return "suspicious token"
    if reason.startswith("stray/suspicious non-ASCII"):
        return "suspicious glyph"
    if reason.startswith("BOM/control"):
        return "BOM/control character"
    if reason.startswith("isolated marker") or reason.startswith("symbol-only"):
        return "isolated marker/separator"
    if reason.startswith("punctuation artifact"):
        return "punctuation artifact"
    if reason.startswith("line starts"):
        return "line-start artifact"
    if reason.startswith("missing space") or reason.startswith("likely joined"):
        return "spacing/joined text"
    if reason.startswith("abbreviation/spacing"):
        return "abbreviation/spacing"
    return "other"


def write_csv(
    paddle_records: list[dict[str, object]],
    abbyy_records: list[dict[str, object]],
    paddle_stats: dict[str, int | str],
    abbyy_stats: dict[str, int | str],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.writer(output)
        writer.writerows(
            build_wide_rows(paddle_records, abbyy_records, paddle_stats, abbyy_stats)
        )


def write_excel_workbook(
    sheets: list[dict[str, object]],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    template_map = load_excel_template(output_path)
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_names: set[str] = set()
    for sheet in sheets:
        proposed_name = str(sheet["title"])
        rows = list(sheet["rows"])
        source_key = str(sheet["source_key"])
        template = template_map.get(source_key)
        title = unique_sheet_name(
            str(template.get("title")) if template else proposed_name,
            used_names,
        )
        worksheet = workbook.create_sheet(title=title)
        for row in rows:
            worksheet.append(row)

        apply_excel_formatting(worksheet, template=template)

    workbook.save(output_path)


def apply_excel_formatting(
    worksheet,
    template: dict[str, object] | None = None,
) -> None:
    wrapped_top = Alignment(vertical="top", wrap_text=True)
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = wrapped_top

    if template:
        worksheet.freeze_panes = template.get("freeze_panes") or "A2"
        for column_letter, width in dict(template.get("column_widths") or {}).items():
            worksheet.column_dimensions[str(column_letter)].width = float(width)
        for row_index, height in dict(template.get("row_heights") or {}).items():
            worksheet.row_dimensions[int(row_index)].height = float(height)
        return

    worksheet.freeze_panes = "A2"
    for column_cells in worksheet.columns:
        values = [len(str(cell.value or "")) for cell in column_cells]
        column_letter = column_cells[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(
            max(values, default=0) + 2,
            60,
        )


def load_excel_template(output_path: Path) -> dict[str, dict[str, object]]:
    if not output_path.exists():
        return {}

    workbook = load_workbook(output_path)
    template_map: dict[str, dict[str, object]] = {}
    for worksheet in workbook.worksheets:
        source_key = normalize_source_key(worksheet["B2"].value)
        if not source_key:
            continue
        template_map[source_key] = {
            "title": worksheet.title,
            "freeze_panes": worksheet.freeze_panes,
            "column_widths": {
                key: dimension.width
                for key, dimension in worksheet.column_dimensions.items()
                if dimension.width is not None
            },
            "row_heights": {
                key: dimension.height
                for key, dimension in worksheet.row_dimensions.items()
                if dimension.height is not None
            },
        }
    return template_map


def unique_sheet_name(proposed_name: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]+", "_", proposed_name).strip() or "Sheet"
    cleaned = cleaned[:31]
    if cleaned not in used_names:
        used_names.add(cleaned)
        return cleaned

    base = cleaned[:28].rstrip() or "Sheet"
    counter = 2
    while True:
        candidate = f"{base}_{counter}"[:31]
        if candidate not in used_names:
            used_names.add(candidate)
            return candidate
        counter += 1


def build_issue_records(issue_map: OrderedDict[tuple[int, str], list[str]]) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for (line_no, snippet), reasons in issue_map.items():
        unique_reasons = list(dict.fromkeys(reasons))
        error_types = sorted(
            {error_type_for_reason(reason) for reason in unique_reasons},
            key=lambda error_type: ERROR_TYPE_PRIORITY.get(error_type, len(ERROR_TYPE_PRIORITY)),
        )
        primary_type = error_types[0] if error_types else "other"
        text = f"L{line_no}: " + " | ".join(unique_reasons) + " :: " + snippet
        records.append(
            {
                "line_no": line_no,
                "snippet": snippet,
                "reasons": unique_reasons,
                "error_types": error_types,
                "primary_type": primary_type,
                "text": text,
            }
        )
    records.sort(
        key=lambda record: (
            ERROR_TYPE_PRIORITY.get(str(record["primary_type"]), len(ERROR_TYPE_PRIORITY)),
            int(record["line_no"]),
            str(record["snippet"]),
        )
    )
    return records


def build_summary_rows(
    paddle_records: list[dict[str, object]],
    abbyy_records: list[dict[str, object]],
    paddle_stats: dict[str, int | str],
    abbyy_stats: dict[str, int | str],
) -> list[list[str]]:
    paddle_counts = count_error_types(paddle_records)
    abbyy_counts = count_error_types(abbyy_records)

    rows = [
        [
            "source file",
            str(paddle_stats["path"]),
            str(abbyy_stats["path"]),
        ],
        [
            "line count",
            str(paddle_stats["lines"]),
            str(abbyy_stats["lines"]),
        ],
        [
            "word count",
            str(paddle_stats["words"]),
            str(abbyy_stats["words"]),
        ],
        [
            "error/artifact entries",
            str(len(paddle_records)),
            str(len(abbyy_records)),
        ],
        [
            "counting note",
            "Error-type counts include all tags on a row, so counts can exceed total entries.",
            "Error-type counts include all tags on a row, so counts can exceed total entries.",
        ],
    ]

    for error_type in ERROR_TYPE_ORDER:
        rows.append(
            [
                f"{error_type} issue count",
                str(paddle_counts.get(error_type, 0)),
                str(abbyy_counts.get(error_type, 0)),
            ]
        )
    return rows


def count_error_types(records: list[dict[str, object]]) -> Counter[str]:
    counts: Counter[str] = Counter()
    for record in records:
        for error_type in record["error_types"]:
            counts[str(error_type)] += 1
    return counts


def build_wide_rows(
    paddle_records: list[dict[str, object]],
    abbyy_records: list[dict[str, object]],
    paddle_stats: dict[str, int | str],
    abbyy_stats: dict[str, int | str],
) -> list[list[str]]:
    summary_rows = build_summary_rows(
        paddle_records,
        abbyy_records,
        paddle_stats,
        abbyy_stats,
    )
    detail_columns: list[tuple[str, list[str]]] = []
    for error_type in ERROR_TYPE_ORDER:
        paddle_items = detail_items_for_type(paddle_records, error_type)
        abbyy_items = detail_items_for_type(abbyy_records, error_type)
        detail_columns.append((f"{error_type} - paddle", paddle_items))
        detail_columns.append((f"{error_type} - abbyy", abbyy_items))

    header = ["summary metric", "summary paddle", "summary abbyy"] + [
        name for name, _items in detail_columns
    ]
    max_rows = max(
        len(summary_rows),
        *(len(items) for _name, items in detail_columns),
    )

    rows = [header]
    for row_index in range(max_rows):
        if row_index < len(summary_rows):
            row = list(summary_rows[row_index])
        else:
            row = ["", "", ""]
        for _name, items in detail_columns:
            row.append(items[row_index] if row_index < len(items) else "")
        rows.append(row)
    return rows


def detail_items_for_type(records: list[dict[str, object]], error_type: str) -> list[str]:
    matching = [
        record
        for record in records
        if error_type in record["error_types"]
    ]
    matching.sort(
        key=lambda record: (
            int(record["line_no"]),
            str(record["snippet"]),
        )
    )
    return [str(record["text"]) for record in matching]


def document_stats(path: Path, text: str, lines: list[str]) -> dict[str, int | str]:
    return {
        "path": str(path),
        "lines": len(lines),
        "words": len(text.split()),
    }


def generate_report(
    paddle_path: Path,
    abbyy_path: Path,
    output_path: Path,
    *,
    known_notes: bool = True,
) -> tuple[int, int]:
    paddle_text = paddle_path.read_text(errors="replace")
    abbyy_text = abbyy_path.read_text(errors="replace")
    paddle_lines = paddle_text.splitlines()
    abbyy_lines = abbyy_text.splitlines()

    paddle_issues = line_level_issues(paddle_lines, detect_overlap=True)
    abbyy_issues = line_level_issues(abbyy_lines, detect_overlap=False)

    if known_notes:
        add_known_structural_notes(paddle_lines, abbyy_lines, paddle_issues, abbyy_issues)

    add_token_comparison_issues(
        token_stream(paddle_text),
        token_stream(abbyy_text),
        paddle_issues,
        abbyy_issues,
    )

    paddle_records = build_issue_records(paddle_issues)
    abbyy_records = build_issue_records(abbyy_issues)
    write_csv(
        paddle_records,
        abbyy_records,
        document_stats(paddle_path, paddle_text, paddle_lines),
        document_stats(abbyy_path, abbyy_text, abbyy_lines),
        output_path,
    )
    return len(paddle_records), len(abbyy_records)


def generate_report_rows(
    paddle_path: Path,
    abbyy_path: Path,
    *,
    known_notes: bool = True,
) -> tuple[list[list[str]], int, int]:
    paddle_text = paddle_path.read_text(errors="replace")
    abbyy_text = abbyy_path.read_text(errors="replace")
    paddle_lines = paddle_text.splitlines()
    abbyy_lines = abbyy_text.splitlines()

    paddle_issues = line_level_issues(paddle_lines, detect_overlap=True)
    abbyy_issues = line_level_issues(abbyy_lines, detect_overlap=False)

    if known_notes:
        add_known_structural_notes(paddle_lines, abbyy_lines, paddle_issues, abbyy_issues)

    add_token_comparison_issues(
        token_stream(paddle_text),
        token_stream(abbyy_text),
        paddle_issues,
        abbyy_issues,
    )

    paddle_records = build_issue_records(paddle_issues)
    abbyy_records = build_issue_records(abbyy_issues)
    rows = build_wide_rows(
        paddle_records,
        abbyy_records,
        document_stats(paddle_path, paddle_text, paddle_lines),
        document_stats(abbyy_path, abbyy_text, abbyy_lines),
    )
    return rows, len(paddle_records), len(abbyy_records)


def primary_ocr_txt_files(base_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(base_dir.rglob("*.txt")):
        name = path.name.lower()
        if name.endswith(".post_ocr_vl_progress.txt"):
            continue
        if "b4-cleaned" in name:
            continue
        files.append(path)
    return files


def normalize_source_key(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    return text.replace("/", "\\").casefold()


def source_key_for_path(path: Path) -> str:
    try:
        value = str(path.relative_to(REPO_ROOT))
    except ValueError:
        value = str(path)
    return normalize_source_key(value)


def default_sheet_title_for_path(paddle_path: Path) -> str:
    folder = paddle_path.parent.name
    alias_map = {
        "large-ppocr-preproccess3x2-docparse-chandra": "Chandra",
        "large-ppocr-preproccess3x2-docparse-deepseek": "Deepseek",
        "large-ppocr-preproccess3x2-docparse-olmocr": "olmOCR",
        "large-ppocr-preproccess3x2-docparse-paddlevl": "PaddleVL",
        "large-ppocr-preprocess": "PaddleOCR only",
    }
    return alias_map.get(folder.casefold(), folder)


def compare_directory_to_abbyy(
    paddle_dir: Path,
    abbyy_dir: Path,
    output_path: Path,
    *,
    known_notes: bool = True,
) -> list[tuple[Path, Path, int, int]]:
    abbyy_by_name: dict[str, list[Path]] = {}
    for path in primary_ocr_txt_files(abbyy_dir):
        abbyy_by_name.setdefault(path.name.lower(), []).append(path)

    sheets: list[dict[str, object]] = []
    summaries: list[tuple[Path, Path, int, int]] = []
    for paddle_path in primary_ocr_txt_files(paddle_dir):
        matches = abbyy_by_name.get(paddle_path.name.lower(), [])
        if not matches:
            continue

        abbyy_path = matches[0]
        rows, paddle_count, abbyy_count = generate_report_rows(
            paddle_path,
            abbyy_path,
            known_notes=known_notes,
        )
        sheets.append(
            {
                "title": default_sheet_title_for_path(paddle_path),
                "rows": rows,
                "source_key": source_key_for_path(paddle_path),
            }
        )
        summaries.append((paddle_path, abbyy_path, paddle_count, abbyy_count))

    if not sheets:
        raise FileNotFoundError(
            f"No matching OCR TXT files found between {paddle_dir} and {abbyy_dir}"
        )

    write_excel_workbook(sheets, output_path)
    return summaries


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--paddle", type=Path, default=DEFAULT_PADDLE, help="Paddle TXT file")
    parser.add_argument("--abbyy", type=Path, default=DEFAULT_ABBYY, help="ABBYY TXT file")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output CSV file")
    parser.add_argument(
        "--paddle-dir",
        type=Path,
        default=None,
        help="Folder of Paddle result subfolders to compare against ABBYY TXT files",
    )
    parser.add_argument(
        "--abbyy-dir",
        type=Path,
        default=None,
        help="Folder containing ABBYY TXT files for directory comparisons",
    )
    parser.add_argument(
        "--excel-output",
        type=Path,
        default=DEFAULT_EXCEL_OUTPUT,
        help="Output XLSX workbook for directory comparisons",
    )
    parser.add_argument(
        "--no-known-notes",
        action="store_true",
        help="Skip the document-specific reading-order/tile-overlap notes",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    if args.paddle_dir or args.abbyy_dir:
        paddle_dir = args.paddle_dir or DEFAULT_PADDLE_DIR
        abbyy_dir = args.abbyy_dir or DEFAULT_ABBYY_DIR
        for label, path in (("Paddle directory", paddle_dir), ("ABBYY directory", abbyy_dir)):
            if not path.exists():
                print(f"{label} does not exist: {path}", file=sys.stderr)
                return 2

        summaries = compare_directory_to_abbyy(
            paddle_dir,
            abbyy_dir,
            args.excel_output,
            known_notes=not args.no_known_notes,
        )
        print(f"Excel workbook written to: {args.excel_output}")
        print(f"Compared files: {len(summaries)}")
        for paddle_path, abbyy_path, paddle_count, abbyy_count in summaries:
            print(
                f"- {paddle_path.parent.name}/{paddle_path.name} vs {abbyy_path.name}: "
                f"Paddle entries={paddle_count}, ABBYY entries={abbyy_count}"
            )
        return 0

    for label, path in (("Paddle", args.paddle), ("ABBYY", args.abbyy)):
        if not path.exists():
            print(f"{label} file does not exist: {path}", file=sys.stderr)
            return 2

    paddle_count, abbyy_count = generate_report(
        args.paddle,
        args.abbyy,
        args.output,
        known_notes=not args.no_known_notes,
    )
    print(f"CSV written to: {args.output}")
    print(f"Paddle entries: {paddle_count}")
    print(f"ABBYY entries: {abbyy_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
