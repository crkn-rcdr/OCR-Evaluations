#!/usr/bin/env python3
"""Compare extra lines unique to the 3x2 and 3x1 tiled Paddle runs.

Default usage:

  python tools/compare_tiling_extra_lines.py

This pairs matching OCR TXT files from:

- 2-multi-document-abbyy-vs-tiled3x2+docparse-paddleocr-paddlevl
- 3-multi-document-abbyy-vs-tiled3x1+docparse-paddleocr-paddlevl

and writes:

- test-results/tiling_3x2_vs_3x1_extra_lines.xlsx
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment


REPO_ROOT = Path(__file__).resolve().parent.parent
EVAL_ROOT = REPO_ROOT.parent
DEFAULT_3X2_DIR = (
    EVAL_ROOT
    / "2-multi-document-abbyy-vs-tiled3x2+docparse-paddleocr-paddlevl"
    / "test-data/preprocessing-tiling_3x2_dp+paddleocr+vl+clean"
)
DEFAULT_3X1_DIR = REPO_ROOT / "test-data/preprocessing-tiling_3x1_dp+paddleocr+vl+clean"
DEFAULT_OUTPUT = REPO_ROOT / "test-results" / "tiling_3x2_vs_3x1_extra_lines.xlsx"

PAGE_NUMBER_RE = re.compile(r"(\d+)$")
WORD_RE = re.compile(r"[A-Za-z0-9]+(?:[â€™'][A-Za-z0-9]+)?")
LINE_TYPE_ORDER = ["separator/symbol", "short fragment", "text line"]


@dataclass(frozen=True)
class LineRecord:
    line_no: int
    text: str
    normalized: str
    word_count: int
    line_type: str


@dataclass(frozen=True)
class PageSummary:
    document: str
    page: str
    source_3x2: Path
    source_3x1: Path
    line_count_3x2: int
    line_count_3x1: int
    word_count_3x2: int
    word_count_3x1: int
    shared_line_count: int
    extra_line_count_3x2: int
    extra_line_count_3x1: int
    extra_word_count_3x2: int
    extra_word_count_3x1: int
    extra_type_counts_3x2: dict[str, int]
    extra_type_counts_3x1: dict[str, int]

    @property
    def extra_line_delta(self) -> int:
        return self.extra_line_count_3x2 - self.extra_line_count_3x1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--run-3x2-dir",
        type=Path,
        default=DEFAULT_3X2_DIR,
        help="Root directory of the 3x2 Paddle OCR TXT outputs.",
    )
    parser.add_argument(
        "--run-3x1-dir",
        type=Path,
        default=DEFAULT_3X1_DIR,
        help="Root directory of the 3x1 Paddle OCR TXT outputs.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Output XLSX workbook path.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = args.output.resolve()
    summaries, extra_3x2_rows, extra_3x1_rows = compare_runs(
        args.run_3x2_dir.resolve(),
        args.run_3x1_dir.resolve(),
    )
    write_workbook(
        workbook_path,
        build_overall_sheet_rows(summaries, args.run_3x2_dir.resolve(), args.run_3x1_dir.resolve()),
        extra_3x2_rows,
        extra_3x1_rows,
    )
    print(f"Wrote {workbook_path}")
    print(f"Compared pages: {len(summaries)}")


def ocr_txt_files_in_dir(base_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(base_dir.glob("*.txt")):
        name = path.name.lower()
        if name.endswith(".post_ocr_vl_progress.txt"):
            continue
        if "b4-cleaned" in name:
            continue
        files.append(path)
    return files


def page_key_for_path(path: Path) -> str:
    match = PAGE_NUMBER_RE.search(path.stem)
    if match:
        return f"page:{int(match.group(1))}"
    return f"name:{path.name.casefold()}"


def normalize_line(text: str) -> str:
    value = unicodedata.normalize("NFKC", text.replace("\ufeff", ""))
    value = re.sub(r"\s+", " ", value).strip().casefold()
    return value


def classify_line(text: str) -> str:
    stripped = text.strip()
    word_count = len(WORD_RE.findall(stripped))
    alnum_count = sum(char.isalnum() for char in stripped)
    alpha_count = sum(char.isalpha() for char in stripped)
    if alnum_count == 0:
        return "separator/symbol"
    if word_count <= 3 or alpha_count < 8:
        return "short fragment"
    return "text line"


def read_line_records(path: Path) -> tuple[list[LineRecord], int]:
    text = path.read_text(encoding="utf-8", errors="replace")
    records: list[LineRecord] = []
    for line_no, raw_line in enumerate(text.splitlines(), 1):
        stripped = raw_line.strip()
        if not stripped:
            continue
        records.append(
            LineRecord(
                line_no=line_no,
                text=stripped,
                normalized=normalize_line(stripped),
                word_count=len(WORD_RE.findall(stripped)),
                line_type=classify_line(stripped),
            )
        )
    return records, len(text.split())


def count_line_types(records: list[LineRecord]) -> Counter[str]:
    counts: Counter[str] = Counter()
    for record in records:
        counts[record.line_type] += 1
    return counts


def select_extra_records(records: list[LineRecord], extra_counts: Counter[str]) -> list[LineRecord]:
    remaining = Counter(extra_counts)
    selected: list[LineRecord] = []
    for record in records:
        if remaining[record.normalized] <= 0:
            continue
        selected.append(record)
        remaining[record.normalized] -= 1
    return selected


def compare_page(run_3x2_path: Path, run_3x1_path: Path) -> tuple[PageSummary, list[dict[str, object]], list[dict[str, object]]]:
    records_3x2, word_count_3x2 = read_line_records(run_3x2_path)
    records_3x1, word_count_3x1 = read_line_records(run_3x1_path)
    counts_3x2 = Counter(record.normalized for record in records_3x2)
    counts_3x1 = Counter(record.normalized for record in records_3x1)

    shared_line_count = sum(min(counts_3x2[key], counts_3x1[key]) for key in (counts_3x2.keys() | counts_3x1.keys()))
    extra_counts_3x2 = counts_3x2 - counts_3x1
    extra_counts_3x1 = counts_3x1 - counts_3x2
    extra_records_3x2 = select_extra_records(records_3x2, extra_counts_3x2)
    extra_records_3x1 = select_extra_records(records_3x1, extra_counts_3x1)
    extra_type_counts_3x2 = count_line_types(extra_records_3x2)
    extra_type_counts_3x1 = count_line_types(extra_records_3x1)

    document = run_3x1_path.parent.name
    page = run_3x1_path.name
    summary = PageSummary(
        document=document,
        page=page,
        source_3x2=run_3x2_path,
        source_3x1=run_3x1_path,
        line_count_3x2=len(records_3x2),
        line_count_3x1=len(records_3x1),
        word_count_3x2=word_count_3x2,
        word_count_3x1=word_count_3x1,
        shared_line_count=shared_line_count,
        extra_line_count_3x2=len(extra_records_3x2),
        extra_line_count_3x1=len(extra_records_3x1),
        extra_word_count_3x2=sum(record.word_count for record in extra_records_3x2),
        extra_word_count_3x1=sum(record.word_count for record in extra_records_3x1),
        extra_type_counts_3x2={line_type: int(extra_type_counts_3x2.get(line_type, 0)) for line_type in LINE_TYPE_ORDER},
        extra_type_counts_3x1={line_type: int(extra_type_counts_3x1.get(line_type, 0)) for line_type in LINE_TYPE_ORDER},
    )

    extra_3x2_rows = [
        build_extra_line_row(document, page, run_3x2_path, record)
        for record in extra_records_3x2
    ]
    extra_3x1_rows = [
        build_extra_line_row(document, page, run_3x1_path, record)
        for record in extra_records_3x1
    ]
    return summary, extra_3x2_rows, extra_3x1_rows


def build_extra_line_row(
    document: str,
    page: str,
    source_path: Path,
    record: LineRecord,
) -> dict[str, object]:
    return {
        "document": document,
        "page": page,
        "source_file": str(source_path),
        "line_no": record.line_no,
        "word_count": record.word_count,
        "line_type": record.line_type,
        "line_text": record.text,
        "normalized": record.normalized,
    }


def compare_runs(
    run_3x2_dir: Path,
    run_3x1_dir: Path,
) -> tuple[list[PageSummary], list[dict[str, object]], list[dict[str, object]]]:
    summaries: list[PageSummary] = []
    extra_3x2_rows: list[dict[str, object]] = []
    extra_3x1_rows: list[dict[str, object]] = []
    missing_documents: list[str] = []
    missing_pages: list[str] = []

    for run_3x1_doc_dir in sorted(path for path in run_3x1_dir.iterdir() if path.is_dir()):
        run_3x2_doc_dir = run_3x2_dir / run_3x1_doc_dir.name
        if not run_3x2_doc_dir.is_dir():
            missing_documents.append(run_3x1_doc_dir.name)
            continue

        run_3x2_by_name = {
            page_key_for_path(path): path for path in ocr_txt_files_in_dir(run_3x2_doc_dir)
        }
        run_3x1_pages = ocr_txt_files_in_dir(run_3x1_doc_dir)
        for run_3x1_path in run_3x1_pages:
            run_3x2_path = run_3x2_by_name.get(page_key_for_path(run_3x1_path))
            if run_3x2_path is None:
                missing_pages.append(f"{run_3x1_doc_dir.name}/{run_3x1_path.name}")
                continue

            summary, page_extra_3x2_rows, page_extra_3x1_rows = compare_page(run_3x2_path, run_3x1_path)
            summaries.append(summary)
            extra_3x2_rows.extend(page_extra_3x2_rows)
            extra_3x1_rows.extend(page_extra_3x1_rows)

        run_3x1_page_keys = {page_key_for_path(path) for path in run_3x1_pages}
        for run_3x2_path in ocr_txt_files_in_dir(run_3x2_doc_dir):
            if page_key_for_path(run_3x2_path) not in run_3x1_page_keys:
                missing_pages.append(f"{run_3x2_doc_dir.name}/{run_3x2_path.name}")

    if missing_documents or missing_pages:
        problems: list[str] = []
        if missing_documents:
            problems.append("Missing matching document directories: " + ", ".join(sorted(missing_documents)[:10]))
        if missing_pages:
            problems.append("Missing matching page TXT files: " + ", ".join(sorted(missing_pages)[:10]))
        raise FileNotFoundError("; ".join(problems))

    summaries.sort(key=lambda item: (item.document.casefold(), item.page.casefold()))
    return summaries, extra_3x2_rows, extra_3x1_rows


def build_overall_sheet_rows(
    summaries: list[PageSummary],
    run_3x2_dir: Path,
    run_3x1_dir: Path,
) -> list[list[object]]:
    total_line_count_3x2 = sum(item.line_count_3x2 for item in summaries)
    total_line_count_3x1 = sum(item.line_count_3x1 for item in summaries)
    total_word_count_3x2 = sum(item.word_count_3x2 for item in summaries)
    total_word_count_3x1 = sum(item.word_count_3x1 for item in summaries)
    total_shared_lines = sum(item.shared_line_count for item in summaries)
    total_extra_lines_3x2 = sum(item.extra_line_count_3x2 for item in summaries)
    total_extra_lines_3x1 = sum(item.extra_line_count_3x1 for item in summaries)
    total_extra_words_3x2 = sum(item.extra_word_count_3x2 for item in summaries)
    total_extra_words_3x1 = sum(item.extra_word_count_3x1 for item in summaries)
    pages_3x2_more = sum(1 for item in summaries if item.extra_line_count_3x2 > item.extra_line_count_3x1)
    pages_3x1_more = sum(1 for item in summaries if item.extra_line_count_3x1 > item.extra_line_count_3x2)
    pages_tied = len(summaries) - pages_3x2_more - pages_3x1_more
    total_type_counts_3x2: Counter[str] = Counter()
    total_type_counts_3x1: Counter[str] = Counter()
    document_totals: dict[str, dict[str, object]] = defaultdict(
        lambda: {
            "page_count": 0,
            "line_count_3x2": 0,
            "line_count_3x1": 0,
            "word_count_3x2": 0,
            "word_count_3x1": 0,
            "shared_line_count": 0,
            "extra_line_count_3x2": 0,
            "extra_line_count_3x1": 0,
            "extra_word_count_3x2": 0,
            "extra_word_count_3x1": 0,
            "type_counts_3x2": Counter(),
            "type_counts_3x1": Counter(),
        }
    )

    for summary in summaries:
        total_type_counts_3x2.update(summary.extra_type_counts_3x2)
        total_type_counts_3x1.update(summary.extra_type_counts_3x1)
        bucket = document_totals[summary.document]
        bucket["page_count"] = int(bucket["page_count"]) + 1
        bucket["line_count_3x2"] = int(bucket["line_count_3x2"]) + summary.line_count_3x2
        bucket["line_count_3x1"] = int(bucket["line_count_3x1"]) + summary.line_count_3x1
        bucket["word_count_3x2"] = int(bucket["word_count_3x2"]) + summary.word_count_3x2
        bucket["word_count_3x1"] = int(bucket["word_count_3x1"]) + summary.word_count_3x1
        bucket["shared_line_count"] = int(bucket["shared_line_count"]) + summary.shared_line_count
        bucket["extra_line_count_3x2"] = int(bucket["extra_line_count_3x2"]) + summary.extra_line_count_3x2
        bucket["extra_line_count_3x1"] = int(bucket["extra_line_count_3x1"]) + summary.extra_line_count_3x1
        bucket["extra_word_count_3x2"] = int(bucket["extra_word_count_3x2"]) + summary.extra_word_count_3x2
        bucket["extra_word_count_3x1"] = int(bucket["extra_word_count_3x1"]) + summary.extra_word_count_3x1
        bucket["type_counts_3x2"].update(summary.extra_type_counts_3x2)
        bucket["type_counts_3x1"].update(summary.extra_type_counts_3x1)

    rows: list[list[object]] = [
        ["summary metric", "run 3x2", "run 3x1"],
        ["dataset root", str(run_3x2_dir), str(run_3x1_dir)],
        ["document count", len(document_totals), len(document_totals)],
        ["page count", len(summaries), len(summaries)],
        ["nonblank line count", total_line_count_3x2, total_line_count_3x1],
        ["word count", total_word_count_3x2, total_word_count_3x1],
        ["shared matched lines", total_shared_lines, total_shared_lines],
        ["extra unique-only lines", total_extra_lines_3x2, total_extra_lines_3x1],
        ["extra unique-only words", total_extra_words_3x2, total_extra_words_3x1],
        ["pages with more unique-only lines", pages_3x2_more, pages_3x1_more],
        ["pages tied on unique-only lines", pages_tied, pages_tied],
        [
            "counting note",
            "Extra lines are nonblank normalized lines that appear more times in 3x2 than 3x1 on the same page.",
            "Extra lines are nonblank normalized lines that appear more times in 3x1 than 3x2 on the same page.",
        ],
    ]
    for line_type in LINE_TYPE_ORDER:
        rows.append(
            [
                f"{line_type} extra lines",
                int(total_type_counts_3x2.get(line_type, 0)),
                int(total_type_counts_3x1.get(line_type, 0)),
            ]
        )

    rows.append([])
    rows.append(["document summary"])
    document_header: list[object] = [
        "document",
        "page count",
        "3x2 nonblank lines",
        "3x1 nonblank lines",
        "3x2 words",
        "3x1 words",
        "shared matched lines",
        "3x2 only lines",
        "3x1 only lines",
        "extra line delta (3x2-3x1)",
        "3x2 only words",
        "3x1 only words",
    ]
    for line_type in LINE_TYPE_ORDER:
        document_header.append(f"3x2 {line_type}")
        document_header.append(f"3x1 {line_type}")
    rows.append(document_header)

    for document in sorted(document_totals):
        bucket = document_totals[document]
        row: list[object] = [
            document,
            int(bucket["page_count"]),
            int(bucket["line_count_3x2"]),
            int(bucket["line_count_3x1"]),
            int(bucket["word_count_3x2"]),
            int(bucket["word_count_3x1"]),
            int(bucket["shared_line_count"]),
            int(bucket["extra_line_count_3x2"]),
            int(bucket["extra_line_count_3x1"]),
            int(bucket["extra_line_count_3x2"]) - int(bucket["extra_line_count_3x1"]),
            int(bucket["extra_word_count_3x2"]),
            int(bucket["extra_word_count_3x1"]),
        ]
        for line_type in LINE_TYPE_ORDER:
            row.append(int(bucket["type_counts_3x2"].get(line_type, 0)))
            row.append(int(bucket["type_counts_3x1"].get(line_type, 0)))
        rows.append(row)

    rows.append([])
    rows.append(["page summary"])
    page_header: list[object] = [
        "document",
        "page",
        "3x2 source file",
        "3x1 source file",
        "3x2 nonblank lines",
        "3x1 nonblank lines",
        "3x2 words",
        "3x1 words",
        "shared matched lines",
        "3x2 only lines",
        "3x1 only lines",
        "extra line delta (3x2-3x1)",
        "3x2 only words",
        "3x1 only words",
    ]
    for line_type in LINE_TYPE_ORDER:
        page_header.append(f"3x2 {line_type}")
        page_header.append(f"3x1 {line_type}")
    rows.append(page_header)

    for summary in summaries:
        row = [
            summary.document,
            summary.page,
            str(summary.source_3x2),
            str(summary.source_3x1),
            summary.line_count_3x2,
            summary.line_count_3x1,
            summary.word_count_3x2,
            summary.word_count_3x1,
            summary.shared_line_count,
            summary.extra_line_count_3x2,
            summary.extra_line_count_3x1,
            summary.extra_line_delta,
            summary.extra_word_count_3x2,
            summary.extra_word_count_3x1,
        ]
        for line_type in LINE_TYPE_ORDER:
            row.append(summary.extra_type_counts_3x2.get(line_type, 0))
            row.append(summary.extra_type_counts_3x1.get(line_type, 0))
        rows.append(row)

    return rows


def build_extra_sheet_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    output: list[list[object]] = [[
        "document",
        "page",
        "source file",
        "line no",
        "word count",
        "line type",
        "line text",
        "normalized",
    ]]
    for row in sorted(rows, key=lambda item: (str(item["document"]).casefold(), str(item["page"]).casefold(), int(item["line_no"]))):
        output.append(
            [
                row["document"],
                row["page"],
                row["source_file"],
                row["line_no"],
                row["word_count"],
                row["line_type"],
                row["line_text"],
                row["normalized"],
            ]
        )
    return output


def write_workbook(
    output_path: Path,
    overall_rows: list[list[object]],
    extra_3x2_rows: list[dict[str, object]],
    extra_3x1_rows: list[dict[str, object]],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [
        ("Overall Summary", overall_rows, "A2"),
        ("3x2 Extra Lines", build_extra_sheet_rows(extra_3x2_rows), "A2"),
        ("3x1 Extra Lines", build_extra_sheet_rows(extra_3x1_rows), "A2"),
    ]
    for title, rows, freeze_panes in sheets:
        worksheet = workbook.create_sheet(title=title)
        for row in rows:
            worksheet.append(row)
        apply_excel_formatting(worksheet, freeze_panes=freeze_panes)
    workbook.save(output_path)


def apply_excel_formatting(worksheet, *, freeze_panes: str) -> None:
    worksheet.freeze_panes = freeze_panes
    wrapped_top = Alignment(vertical="top", wrap_text=True)
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = wrapped_top
    for column_cells in worksheet.columns:
        values = [len(str(cell.value or "")) for cell in column_cells]
        column_letter = column_cells[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(max(values, default=0) + 2, 70)


if __name__ == "__main__":
    main()
