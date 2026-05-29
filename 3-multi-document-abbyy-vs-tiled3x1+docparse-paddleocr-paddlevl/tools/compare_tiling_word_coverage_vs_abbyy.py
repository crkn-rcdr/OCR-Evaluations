#!/usr/bin/env python3
"""Compare 3x2 and 3x1 word coverage against ABBYY page by page.

Default usage:

  python tools/compare_tiling_word_coverage_vs_abbyy.py

This pairs matching OCR TXT files from:

- 2-multi-document-abbyy-vs-tiled3x2+docparse-paddleocr-paddlevl
- 3-multi-document-abbyy-vs-tiled3x1+docparse-paddleocr-paddlevl
- ABBYY pages under the 3x1 repo

and writes:

- test-results/tiling_3x2_vs_3x1_abbyy_word_coverage.xlsx
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment


REPO_ROOT = Path(__file__).resolve().parent.parent
EVAL_ROOT = REPO_ROOT.parent
DEFAULT_3X2_DIR = (
    EVAL_ROOT
    / "2-multi-document-abbyy-vs-tiled3x2+docparse-paddleocr-paddlevl"
    / "test-data/preprocessing-tiling_3x2_dp+paddleocr+vl+clean"
)
DEFAULT_3X1_DIR = REPO_ROOT / "test-data/preprocessing-tiling_3x1_dp+paddleocr+vl+clean"
DEFAULT_ABBYY_DIR = REPO_ROOT / "test-data/abbyy"
DEFAULT_OUTPUT = REPO_ROOT / "test-results" / "tiling_3x2_vs_3x1_abbyy_word_coverage.xlsx"

PAGE_NUMBER_RE = re.compile(r"(\d+)$")
WORD_RE = re.compile(r"[0-9A-Za-zÀ-ÖØ-öø-ÿ]+(?:['’][0-9A-Za-zÀ-ÖØ-öø-ÿ]+)?")


@dataclass(frozen=True)
class PageCoverage:
    document: str
    page: str
    source_3x2: Path
    source_3x1: Path
    source_abbyy: Path
    word_count_3x2: int
    word_count_3x1: int
    word_count_abbyy: int
    unique_count_3x2: int
    unique_count_3x1: int
    unique_count_abbyy: int
    matched_tokens_3x2: int
    matched_tokens_3x1: int
    matched_unique_3x2: int
    matched_unique_3x1: int
    extra_tokens_3x2: int
    extra_tokens_3x1: int
    missing_tokens_3x2: int
    missing_tokens_3x1: int

    @property
    def coverage_pct_3x2(self) -> float:
        return percentage(self.matched_tokens_3x2, self.word_count_abbyy)

    @property
    def coverage_pct_3x1(self) -> float:
        return percentage(self.matched_tokens_3x1, self.word_count_abbyy)

    @property
    def precision_pct_3x2(self) -> float:
        return percentage(self.matched_tokens_3x2, self.word_count_3x2)

    @property
    def precision_pct_3x1(self) -> float:
        return percentage(self.matched_tokens_3x1, self.word_count_3x1)

    @property
    def unique_coverage_pct_3x2(self) -> float:
        return percentage(self.matched_unique_3x2, self.unique_count_abbyy)

    @property
    def unique_coverage_pct_3x1(self) -> float:
        return percentage(self.matched_unique_3x1, self.unique_count_abbyy)

    @property
    def word_count_delta_3x2(self) -> int:
        return self.word_count_3x2 - self.word_count_abbyy

    @property
    def word_count_delta_3x1(self) -> int:
        return self.word_count_3x1 - self.word_count_abbyy


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-3x2-dir", type=Path, default=DEFAULT_3X2_DIR)
    parser.add_argument("--run-3x1-dir", type=Path, default=DEFAULT_3X1_DIR)
    parser.add_argument("--abbyy-dir", type=Path, default=DEFAULT_ABBYY_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    summaries, missing_3x2_rows, missing_3x1_rows, extra_3x2_rows, extra_3x1_rows = compare_runs(
        args.run_3x2_dir.resolve(),
        args.run_3x1_dir.resolve(),
        args.abbyy_dir.resolve(),
    )
    write_workbook(
        args.output.resolve(),
        build_overall_rows(summaries, args.run_3x2_dir.resolve(), args.run_3x1_dir.resolve(), args.abbyy_dir.resolve()),
        build_token_sheet_rows(missing_3x2_rows),
        build_token_sheet_rows(missing_3x1_rows),
        build_token_sheet_rows(extra_3x2_rows),
        build_token_sheet_rows(extra_3x1_rows),
    )
    print(f"Wrote {args.output.resolve()}")
    print(f"Compared pages: {len(summaries)}")


def ocr_txt_files_in_dir(base_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(base_dir.glob("*.txt")):
        name = path.name.lower()
        if name.endswith(".post_ocr_vl_progress.txt"):
            continue
        if "b4-cleaned" in name:
            continue
        files.append(path)
    return files


def page_key_for_path(path: Path) -> str:
    match = PAGE_NUMBER_RE.search(path.stem)
    if match:
        return f"page:{int(match.group(1))}"
    return f"name:{path.name.casefold()}"


def normalize_token(token: str) -> str:
    return unicodedata.normalize("NFKC", token).replace("’", "'").casefold()


def read_token_counter(path: Path) -> tuple[Counter[str], int]:
    text = path.read_text(encoding="utf-8", errors="replace")
    tokens = [normalize_token(token) for token in WORD_RE.findall(text)]
    return Counter(tokens), len(tokens)


def matched_token_count(left: Counter[str], right: Counter[str]) -> int:
    return sum((left & right).values())


def matched_unique_count(left: Counter[str], right: Counter[str]) -> int:
    return len(set(left) & set(right))


def compare_page(
    path_3x2: Path,
    path_3x1: Path,
    abbyy_path: Path,
) -> tuple[PageCoverage, list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    counter_3x2, word_count_3x2 = read_token_counter(path_3x2)
    counter_3x1, word_count_3x1 = read_token_counter(path_3x1)
    counter_abbyy, word_count_abbyy = read_token_counter(abbyy_path)

    matched_tokens_3x2 = matched_token_count(counter_3x2, counter_abbyy)
    matched_tokens_3x1 = matched_token_count(counter_3x1, counter_abbyy)
    matched_unique_3x2 = matched_unique_count(counter_3x2, counter_abbyy)
    matched_unique_3x1 = matched_unique_count(counter_3x1, counter_abbyy)

    missing_3x2 = counter_abbyy - counter_3x2
    missing_3x1 = counter_abbyy - counter_3x1
    extra_3x2 = counter_3x2 - counter_abbyy
    extra_3x1 = counter_3x1 - counter_abbyy

    summary = PageCoverage(
        document=path_3x1.parent.name,
        page=path_3x1.name,
        source_3x2=path_3x2,
        source_3x1=path_3x1,
        source_abbyy=abbyy_path,
        word_count_3x2=word_count_3x2,
        word_count_3x1=word_count_3x1,
        word_count_abbyy=word_count_abbyy,
        unique_count_3x2=len(counter_3x2),
        unique_count_3x1=len(counter_3x1),
        unique_count_abbyy=len(counter_abbyy),
        matched_tokens_3x2=matched_tokens_3x2,
        matched_tokens_3x1=matched_tokens_3x1,
        matched_unique_3x2=matched_unique_3x2,
        matched_unique_3x1=matched_unique_3x1,
        extra_tokens_3x2=word_count_3x2 - matched_tokens_3x2,
        extra_tokens_3x1=word_count_3x1 - matched_tokens_3x1,
        missing_tokens_3x2=word_count_abbyy - matched_tokens_3x2,
        missing_tokens_3x1=word_count_abbyy - matched_tokens_3x1,
    )

    missing_3x2_rows = build_token_rows(summary.document, summary.page, "3x2 missing vs ABBYY", missing_3x2)
    missing_3x1_rows = build_token_rows(summary.document, summary.page, "3x1 missing vs ABBYY", missing_3x1)
    extra_3x2_rows = build_token_rows(summary.document, summary.page, "3x2 extra vs ABBYY", extra_3x2)
    extra_3x1_rows = build_token_rows(summary.document, summary.page, "3x1 extra vs ABBYY", extra_3x1)
    return summary, missing_3x2_rows, missing_3x1_rows, extra_3x2_rows, extra_3x1_rows


def build_token_rows(document: str, page: str, category: str, counter: Counter[str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for token, count in sorted(counter.items(), key=lambda item: (-item[1], item[0])):
        rows.append(
            {
                "document": document,
                "page": page,
                "category": category,
                "token": token,
                "count": int(count),
            }
        )
    return rows


def compare_runs(
    run_3x2_dir: Path,
    run_3x1_dir: Path,
    abbyy_dir: Path,
) -> tuple[list[PageCoverage], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    summaries: list[PageCoverage] = []
    missing_3x2_rows: list[dict[str, object]] = []
    missing_3x1_rows: list[dict[str, object]] = []
    extra_3x2_rows: list[dict[str, object]] = []
    extra_3x1_rows: list[dict[str, object]] = []
    missing_documents: list[str] = []
    missing_pages: list[str] = []

    for doc_dir_3x1 in sorted(path for path in run_3x1_dir.iterdir() if path.is_dir()):
        doc_dir_3x2 = run_3x2_dir / doc_dir_3x1.name
        doc_dir_abbyy = abbyy_dir / doc_dir_3x1.name
        if not doc_dir_3x2.is_dir() or not doc_dir_abbyy.is_dir():
            missing_documents.append(doc_dir_3x1.name)
            continue

        pages_3x2 = {page_key_for_path(path): path for path in ocr_txt_files_in_dir(doc_dir_3x2)}
        pages_abbyy = {page_key_for_path(path): path for path in ocr_txt_files_in_dir(doc_dir_abbyy)}
        pages_3x1 = ocr_txt_files_in_dir(doc_dir_3x1)
        for path_3x1 in pages_3x1:
            key = page_key_for_path(path_3x1)
            path_3x2 = pages_3x2.get(key)
            abbyy_path = pages_abbyy.get(key)
            if path_3x2 is None or abbyy_path is None:
                missing_pages.append(f"{doc_dir_3x1.name}/{path_3x1.name}")
                continue

            summary, page_missing_3x2, page_missing_3x1, page_extra_3x2, page_extra_3x1 = compare_page(
                path_3x2, path_3x1, abbyy_path
            )
            summaries.append(summary)
            missing_3x2_rows.extend(page_missing_3x2)
            missing_3x1_rows.extend(page_missing_3x1)
            extra_3x2_rows.extend(page_extra_3x2)
            extra_3x1_rows.extend(page_extra_3x1)

    if missing_documents or missing_pages:
        parts: list[str] = []
        if missing_documents:
            parts.append("Missing matching document directories: " + ", ".join(sorted(missing_documents)[:10]))
        if missing_pages:
            parts.append("Missing matching pages: " + ", ".join(sorted(missing_pages)[:10]))
        raise FileNotFoundError("; ".join(parts))

    summaries.sort(key=lambda item: (item.document.casefold(), item.page.casefold()))
    return summaries, missing_3x2_rows, missing_3x1_rows, extra_3x2_rows, extra_3x1_rows


def build_overall_rows(
    summaries: list[PageCoverage],
    run_3x2_dir: Path,
    run_3x1_dir: Path,
    abbyy_dir: Path,
) -> list[list[object]]:
    total_word_count_3x2 = sum(item.word_count_3x2 for item in summaries)
    total_word_count_3x1 = sum(item.word_count_3x1 for item in summaries)
    total_word_count_abbyy = sum(item.word_count_abbyy for item in summaries)
    total_unique_3x2 = sum(item.unique_count_3x2 for item in summaries)
    total_unique_3x1 = sum(item.unique_count_3x1 for item in summaries)
    total_unique_abbyy = sum(item.unique_count_abbyy for item in summaries)
    total_matched_tokens_3x2 = sum(item.matched_tokens_3x2 for item in summaries)
    total_matched_tokens_3x1 = sum(item.matched_tokens_3x1 for item in summaries)
    total_matched_unique_3x2 = sum(item.matched_unique_3x2 for item in summaries)
    total_matched_unique_3x1 = sum(item.matched_unique_3x1 for item in summaries)
    total_extra_tokens_3x2 = sum(item.extra_tokens_3x2 for item in summaries)
    total_extra_tokens_3x1 = sum(item.extra_tokens_3x1 for item in summaries)
    total_missing_tokens_3x2 = sum(item.missing_tokens_3x2 for item in summaries)
    total_missing_tokens_3x1 = sum(item.missing_tokens_3x1 for item in summaries)

    pages_better_coverage_3x2 = sum(1 for item in summaries if item.coverage_pct_3x2 > item.coverage_pct_3x1)
    pages_better_coverage_3x1 = sum(1 for item in summaries if item.coverage_pct_3x1 > item.coverage_pct_3x2)
    pages_tied_coverage = len(summaries) - pages_better_coverage_3x2 - pages_better_coverage_3x1
    pages_closer_count_3x2 = sum(
        1 for item in summaries if abs(item.word_count_delta_3x2) < abs(item.word_count_delta_3x1)
    )
    pages_closer_count_3x1 = sum(
        1 for item in summaries if abs(item.word_count_delta_3x1) < abs(item.word_count_delta_3x2)
    )
    pages_tied_count = len(summaries) - pages_closer_count_3x2 - pages_closer_count_3x1

    document_totals: dict[str, dict[str, object]] = defaultdict(
        lambda: {
            "page_count": 0,
            "word_count_3x2": 0,
            "word_count_3x1": 0,
            "word_count_abbyy": 0,
            "matched_tokens_3x2": 0,
            "matched_tokens_3x1": 0,
            "matched_unique_3x2": 0,
            "matched_unique_3x1": 0,
            "extra_tokens_3x2": 0,
            "extra_tokens_3x1": 0,
            "missing_tokens_3x2": 0,
            "missing_tokens_3x1": 0,
        }
    )
    for item in summaries:
        bucket = document_totals[item.document]
        bucket["page_count"] = int(bucket["page_count"]) + 1
        bucket["word_count_3x2"] = int(bucket["word_count_3x2"]) + item.word_count_3x2
        bucket["word_count_3x1"] = int(bucket["word_count_3x1"]) + item.word_count_3x1
        bucket["word_count_abbyy"] = int(bucket["word_count_abbyy"]) + item.word_count_abbyy
        bucket["matched_tokens_3x2"] = int(bucket["matched_tokens_3x2"]) + item.matched_tokens_3x2
        bucket["matched_tokens_3x1"] = int(bucket["matched_tokens_3x1"]) + item.matched_tokens_3x1
        bucket["matched_unique_3x2"] = int(bucket["matched_unique_3x2"]) + item.matched_unique_3x2
        bucket["matched_unique_3x1"] = int(bucket["matched_unique_3x1"]) + item.matched_unique_3x1
        bucket["extra_tokens_3x2"] = int(bucket["extra_tokens_3x2"]) + item.extra_tokens_3x2
        bucket["extra_tokens_3x1"] = int(bucket["extra_tokens_3x1"]) + item.extra_tokens_3x1
        bucket["missing_tokens_3x2"] = int(bucket["missing_tokens_3x2"]) + item.missing_tokens_3x2
        bucket["missing_tokens_3x1"] = int(bucket["missing_tokens_3x1"]) + item.missing_tokens_3x1

    rows: list[list[object]] = [
        ["summary metric", "run 3x2", "run 3x1", "abbyy baseline"],
        ["dataset root", str(run_3x2_dir), str(run_3x1_dir), str(abbyy_dir)],
        ["document count", len(document_totals), len(document_totals), len(document_totals)],
        ["page count", len(summaries), len(summaries), len(summaries)],
        ["word count", total_word_count_3x2, total_word_count_3x1, total_word_count_abbyy],
        ["unique token count", total_unique_3x2, total_unique_3x1, total_unique_abbyy],
        ["matched ABBYY token occurrences", total_matched_tokens_3x2, total_matched_tokens_3x1, total_word_count_abbyy],
        ["matched ABBYY unique tokens", total_matched_unique_3x2, total_matched_unique_3x1, total_unique_abbyy],
        ["coverage of ABBYY token occurrences", f"{percentage(total_matched_tokens_3x2, total_word_count_abbyy):.2f}%", f"{percentage(total_matched_tokens_3x1, total_word_count_abbyy):.2f}%", "100.00%"],
        ["coverage of ABBYY unique tokens", f"{percentage(total_matched_unique_3x2, total_unique_abbyy):.2f}%", f"{percentage(total_matched_unique_3x1, total_unique_abbyy):.2f}%", "100.00%"],
        ["token precision vs ABBYY", f"{percentage(total_matched_tokens_3x2, total_word_count_3x2):.2f}%", f"{percentage(total_matched_tokens_3x1, total_word_count_3x1):.2f}%", ""],
        ["extra tokens vs ABBYY", total_extra_tokens_3x2, total_extra_tokens_3x1, 0],
        ["missing ABBYY tokens", total_missing_tokens_3x2, total_missing_tokens_3x1, 0],
        ["pages with better ABBYY token coverage", pages_better_coverage_3x2, pages_better_coverage_3x1, pages_tied_coverage],
        ["pages closer to ABBYY word count", pages_closer_count_3x2, pages_closer_count_3x1, pages_tied_count],
    ]

    rows.append([])
    rows.append(["document summary"])
    rows.append(
        [
            "document",
            "page count",
            "3x2 words",
            "3x1 words",
            "abbyy words",
            "3x2 matched ABBYY tokens",
            "3x1 matched ABBYY tokens",
            "3x2 coverage %",
            "3x1 coverage %",
            "3x2 precision %",
            "3x1 precision %",
            "3x2 missing ABBYY tokens",
            "3x1 missing ABBYY tokens",
            "3x2 extra tokens",
            "3x1 extra tokens",
        ]
    )
    for document in sorted(document_totals):
        bucket = document_totals[document]
        word_count_abbyy = int(bucket["word_count_abbyy"])
        word_count_3x2 = int(bucket["word_count_3x2"])
        word_count_3x1 = int(bucket["word_count_3x1"])
        matched_tokens_3x2 = int(bucket["matched_tokens_3x2"])
        matched_tokens_3x1 = int(bucket["matched_tokens_3x1"])
        rows.append(
            [
                document,
                int(bucket["page_count"]),
                word_count_3x2,
                word_count_3x1,
                word_count_abbyy,
                matched_tokens_3x2,
                matched_tokens_3x1,
                f"{percentage(matched_tokens_3x2, word_count_abbyy):.2f}%",
                f"{percentage(matched_tokens_3x1, word_count_abbyy):.2f}%",
                f"{percentage(matched_tokens_3x2, word_count_3x2):.2f}%",
                f"{percentage(matched_tokens_3x1, word_count_3x1):.2f}%",
                int(bucket["missing_tokens_3x2"]),
                int(bucket["missing_tokens_3x1"]),
                int(bucket["extra_tokens_3x2"]),
                int(bucket["extra_tokens_3x1"]),
            ]
        )

    rows.append([])
    rows.append(["page summary"])
    rows.append(
        [
            "document",
            "page",
            "3x2 source file",
            "3x1 source file",
            "abbyy source file",
            "3x2 words",
            "3x1 words",
            "abbyy words",
            "3x2 unique tokens",
            "3x1 unique tokens",
            "abbyy unique tokens",
            "3x2 matched ABBYY tokens",
            "3x1 matched ABBYY tokens",
            "3x2 matched ABBYY unique tokens",
            "3x1 matched ABBYY unique tokens",
            "3x2 coverage %",
            "3x1 coverage %",
            "3x2 unique coverage %",
            "3x1 unique coverage %",
            "3x2 precision %",
            "3x1 precision %",
            "3x2 missing ABBYY tokens",
            "3x1 missing ABBYY tokens",
            "3x2 extra tokens",
            "3x1 extra tokens",
            "3x2 word count delta vs ABBYY",
            "3x1 word count delta vs ABBYY",
        ]
    )
    for item in summaries:
        rows.append(
            [
                item.document,
                item.page,
                str(item.source_3x2),
                str(item.source_3x1),
                str(item.source_abbyy),
                item.word_count_3x2,
                item.word_count_3x1,
                item.word_count_abbyy,
                item.unique_count_3x2,
                item.unique_count_3x1,
                item.unique_count_abbyy,
                item.matched_tokens_3x2,
                item.matched_tokens_3x1,
                item.matched_unique_3x2,
                item.matched_unique_3x1,
                f"{item.coverage_pct_3x2:.2f}%",
                f"{item.coverage_pct_3x1:.2f}%",
                f"{item.unique_coverage_pct_3x2:.2f}%",
                f"{item.unique_coverage_pct_3x1:.2f}%",
                f"{item.precision_pct_3x2:.2f}%",
                f"{item.precision_pct_3x1:.2f}%",
                item.missing_tokens_3x2,
                item.missing_tokens_3x1,
                item.extra_tokens_3x2,
                item.extra_tokens_3x1,
                item.word_count_delta_3x2,
                item.word_count_delta_3x1,
            ]
        )
    return rows


def build_token_sheet_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    output: list[list[object]] = [["document", "page", "category", "token", "count"]]
    for row in sorted(rows, key=lambda item: (str(item["document"]).casefold(), str(item["page"]).casefold(), str(item["category"]).casefold(), -int(item["count"]), str(item["token"]))):
        output.append([row["document"], row["page"], row["category"], row["token"], row["count"]])
    return output


def write_workbook(
    output_path: Path,
    overall_rows: list[list[object]],
    missing_3x2_rows: list[list[object]],
    missing_3x1_rows: list[list[object]],
    extra_3x2_rows: list[list[object]],
    extra_3x1_rows: list[list[object]],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in [
        ("Overall Summary", overall_rows),
        ("3x2 Missing vs ABBYY", missing_3x2_rows),
        ("3x1 Missing vs ABBYY", missing_3x1_rows),
        ("3x2 Extra vs ABBYY", extra_3x2_rows),
        ("3x1 Extra vs ABBYY", extra_3x1_rows),
    ]:
        worksheet = workbook.create_sheet(title=title)
        for row in rows:
            worksheet.append(row)
        apply_excel_formatting(worksheet)
    workbook.save(output_path)


def apply_excel_formatting(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    wrapped_top = Alignment(vertical="top", wrap_text=True)
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = wrapped_top
    for column_cells in worksheet.columns:
        values = [len(str(cell.value or "")) for cell in column_cells]
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(values, default=0) + 2, 70)


def percentage(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return (numerator / denominator) * 100.0


if __name__ == "__main__":
    main()
